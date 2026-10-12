"""商品 Excel 导入：只读预览、数据库查重及商品/规则的整批提交。"""

import base64
import hashlib
import hmac
import json
import math
import re
import time
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import String, func, literal, select, text, union_all
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from models import Product, ProductImpairmentRule
from schemas import ProductImportError, ProductImportResponse, ProductImportRow


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_SHEET_ROWS = 10001
MAX_SAFE_STOCK = 2147483647
PREVIEW_TTL_SECONDS = 15 * 60
IMPORT_SHEET = "商品导入"
IMPORT_HEADERS = ("SKU", "商品名称", "商品类型", "安全库存")
HEADER_FIELDS = {
    "sku": "sku", "商品名称": "name", "产品名称": "name",
    "商品类型": "product_type", "安全库存": "safe_stock_quantity",
}
TYPE_NAMES = {"new": "新品", "stable": "稳健商品"}
TYPE_VALUES = {"new": "new", "stable": "stable", "新品": "new", "稳健商品": "stable"}
STATUS_NAMES = {"new": "可新增", "skipped": "已存在，跳过", "error": "错误", "created": "已新增"}


@dataclass
class ParsedProductImport:
    file_hash: str
    rows: list[ProductImportRow]


def _set_text(cell, value) -> None:
    # 导出的用户文本即使以 "=" 开头，也不能成为 Excel 公式。
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.number_format = "@"


def _style_header(sheet) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28


def build_product_import_template() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = IMPORT_SHEET
    sheet.append(IMPORT_HEADERS)
    _style_header(sheet)
    for column, width in {"A": 24, "B": 36, "C": 18, "D": 18}.items():
        sheet.column_dimensions[column].width = width
    for row in range(2, MAX_IMPORT_ROWS + 2):
        sheet.cell(row, 1).number_format = "@"
        sheet.cell(row, 2).number_format = "@"
        sheet.cell(row, 4).number_format = "0"
    types = DataValidation(type="list", formula1='"新品,稳健商品"', allow_blank=False)
    types.errorTitle = "商品类型无效"
    types.error = "请选择新品或稳健商品"
    types.showErrorMessage = True
    types.errorStyle = "stop"
    sheet.add_data_validation(types)
    types.add(f"C2:C{MAX_IMPORT_ROWS + 1}")
    stock = DataValidation(
        type="whole", operator="between", formula1=0, formula2=MAX_SAFE_STOCK,
        allow_blank=True,
    )
    stock.errorTitle = "安全库存无效"
    stock.error = "请输入非负整数；新品填写 0"
    stock.showErrorMessage = True
    stock.errorStyle = "stop"
    sheet.add_data_validation(stock)
    stock.add(f"D2:D{MAX_IMPORT_ROWS + 1}")
    instructions = workbook.create_sheet("填写说明")
    for note in (
        "在“商品导入”工作表填写数据；一行一个商品，最多 2000 条，文件不超过 5 MB。",
        "SKU 必须按文本填写，最多 64 字符；数字 SKU 丢失的前导零无法自动恢复。",
        "商品名称必填，最多 255 字符；商品类型必填，选择新品或稳健商品。",
        "安全库存留空按 0 处理；新品必须为 0；关键列不接受公式。",
        "数据库已有 SKU 默认跳过；文件内重复 SKU 为错误，需修正后再导入。",
        "预览包含隐藏行；有错误时整批不写入。确认时会再次校验数据库状态。",
        "这里只新增商品基础资料及当天规则，不导入图片、库存数量、成本或历史日期。",
        "兼容商品导出的四个同名字段；商品 ID、图片、库存、金额和日期等列不参与导入。",
        "下方仅为填写示例，不会作为导入数据。",
    ):
        instructions.append([note])
    instructions.append([])
    instructions.append(list(IMPORT_HEADERS))
    instructions.append(["DEMO001", "示例商品 A", "新品", 0])
    instructions.append(["DEMO002", "示例商品 B", "稳健商品", 100])
    instructions.column_dimensions["A"].width = 105
    instructions.column_dimensions["B"].width = 30
    instructions.column_dimensions["C"].width = 18
    instructions.column_dimensions["D"].width = 18
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _check_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise HTTPException(400, "Excel 解压后内容过大，请减少数据或移除图片")
            if len({entry.filename for entry in entries}) != len(entries):
                raise HTTPException(400, "Excel 文件包含重复内容，请重新保存")
            for entry in entries:
                name = entry.filename.lower()
                if name.endswith("vbaproject.bin"):
                    raise HTTPException(400, "仅支持不含宏的 .xlsx 文件")
                if name.endswith((".xml", ".rels")):
                    xml = archive.read(entry).replace(b"\x00", b"").upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise HTTPException(400, "Excel 包含不支持的 XML 声明，请重新保存")
    except (BadZipFile, OSError, RuntimeError, ValueError):
        raise HTTPException(400, "Excel 文件损坏或已加密，请重新保存为 .xlsx") from None


def _display(value, limit: int) -> str:
    value = "" if value is None else str(value).strip()
    return value if len(value) <= limit else value[:limit - 1] + "…"


def parse_product_import(filename: str, content: bytes) -> ParsedProductImport:
    if Path(filename).suffix.lower() != ".xlsx":
        raise HTTPException(400, "仅支持 .xlsx 商品文件，请先下载模板")
    if not content:
        raise HTTPException(400, "导入文件为空")
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "商品导入文件不能超过 5 MB")
    _check_archive(content)
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = (
            workbook[IMPORT_SHEET] if IMPORT_SHEET in workbook.sheetnames
            else workbook["商品数据"] if "商品数据" in workbook.sheetnames
            else workbook.worksheets[0]
        )
        # 不信任第三方文件声明的维度，避免 A1 维度导致有效数据被忽略。
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=32)
        header = next(iterator, ())
        indexes = {}
        for index, cell in enumerate(header):
            key = str(cell.value or "").strip().casefold()
            field = HEADER_FIELDS.get(key)
            if field:
                if field in indexes:
                    raise HTTPException(400, f"Excel 重复列：{cell.value}")
                indexes[field] = index
        missing = [label for label in IMPORT_HEADERS if HEADER_FIELDS[label.casefold()] not in indexes]
        if missing:
            raise HTTPException(400, "Excel 缺少必要列：" + "、".join(missing))
        rows = []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number > MAX_SHEET_ROWS:
                raise HTTPException(400, "工作表空白间隔过多，请整理数据后重新导入")
            selected = {field: cells[index] for field, index in indexes.items()}
            if all(cell.value in (None, "") for cell in selected.values()):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise HTTPException(400, "每次最多导入 2000 条商品")
            errors = []

            def error(field, message):
                errors.append(ProductImportError(field=field, message=message))

            labels = {"sku": "SKU", "name": "商品名称", "product_type": "商品类型",
                      "safe_stock_quantity": "安全库存"}
            invalid_fields = {field for field, cell in selected.items() if cell.data_type in ("f", "e")}
            for field in invalid_fields:
                error(labels[field], "不接受公式，请填写实际值" if selected[field].data_type == "f"
                      else "Excel 单元格存在错误，请填写实际值")
            raw_sku = selected["sku"].value
            sku = _display(raw_sku, 128)
            if "sku" not in invalid_fields:
                if not isinstance(raw_sku, str) or not raw_sku.strip():
                    error("SKU", "SKU 必须是非空文本；数字 SKU 请设置为文本后重新填写")
                elif len(raw_sku.strip()) > 64:
                    error("SKU", "SKU 不能超过 64 字符")
            raw_name = selected["name"].value
            name = _display(raw_name, 512)
            if "name" not in invalid_fields:
                if not isinstance(raw_name, str) or not raw_name.strip():
                    error("商品名称", "商品名称必须是非空文本")
                elif len(raw_name.strip()) > 255:
                    error("商品名称", "商品名称不能超过 255 字符")
            raw_type = selected["product_type"].value
            type_value = TYPE_VALUES.get(str(raw_type or "").strip().casefold())
            if "product_type" not in invalid_fields and type_value is None:
                error("商品类型", "请选择新品或稳健商品")
            raw_stock = selected["safe_stock_quantity"].value
            stock_value = None
            if "safe_stock_quantity" not in invalid_fields:
                if raw_stock in (None, ""):
                    stock_value = 0
                elif isinstance(raw_stock, bool):
                    pass
                elif isinstance(raw_stock, int):
                    stock_value = raw_stock
                elif isinstance(raw_stock, float) and math.isfinite(raw_stock) and raw_stock.is_integer():
                    stock_value = int(raw_stock)
                elif isinstance(raw_stock, str) and re.fullmatch(r"[0-9]+", raw_stock.strip()):
                    # 先限制文本长度，避免极长整数字符串触发 Python 的转换上限。
                    if len(raw_stock.strip()) <= 10:
                        stock_value = int(raw_stock.strip())
                if stock_value is None or not 0 <= stock_value <= MAX_SAFE_STOCK:
                    error("安全库存", "安全库存必须是 0 至 2147483647 的整数")
                    stock_value = None
                elif type_value == "new" and stock_value != 0:
                    error("安全库存", "新品的安全库存必须为 0")
            rows.append(ProductImportRow(
                row_number=row_number, sku=sku, name=name,
                product_type=type_value or _display(raw_type, 64) or None,
                safe_stock_quantity=stock_value, errors=errors,
                status="error" if errors else "new",
                message="请修正错误" if errors else "待新增",
            ))
        if not rows:
            raise HTTPException(400, "文件中没有可导入的商品，请在商品导入工作表填写数据")
        return ParsedProductImport(hashlib.sha256(content).hexdigest(), rows)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Excel 无法读取，请使用模板重新保存为 .xlsx") from None
    finally:
        if workbook is not None:
            workbook.close()


async def preview_product_import(db: AsyncSession, parsed: ParsedProductImport) -> ProductImportResponse:
    rows = [row.model_copy(deep=True) for row in parsed.rows]
    sku_rows = [row for row in rows if row.sku and not any(error.field == "SKU" for error in row.errors)]
    if sku_rows:
        # 文件内比较与已存在 SKU 的比较均采用实库唯一列的排序规则。
        # 不能只用 Python 大小写转换近似 MySQL 的大小写、重音等比较行为。
        collation = (await db.execute(text(
            "SELECT COLLATION_NAME FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='products' AND COLUMN_NAME='sku'"
        ))).scalar_one_or_none()
        if not collation or not re.fullmatch(r"[a-zA-Z0-9_]+", collation):
            raise HTTPException(500, "无法识别商品 SKU 排序规则，请检查数据库结构")
        incoming = union_all(*[
            select(
                literal(row.row_number).label("row_number"),
                literal(row.sku, type_=String(64)).collate(collation).label("sku"),
            )
            for row in sku_rows
        ]).subquery("incoming_products")
        groups = select(
            incoming.c.sku, func.count().label("duplicate_count")
        ).group_by(incoming.c.sku).subquery("product_sku_groups")
        matches = (await db.execute(
            select(
                incoming.c.row_number, groups.c.duplicate_count,
                Product.id.label("existing_product_id"),
                Product.name.label("existing_product_name"),
            )
            .select_from(incoming.join(groups, incoming.c.sku == groups.c.sku))
            .outerjoin(Product, Product.sku == incoming.c.sku)
        )).mappings().all()
        by_row = {match["row_number"]: match for match in matches}
        for row in sku_rows:
            match = by_row[row.row_number]
            if match["duplicate_count"] > 1:
                row.errors.append(ProductImportError(field="SKU", message="文件内 SKU 重复（按数据库排序规则比较）"))
            row.existing_product_id = match["existing_product_id"]
            row.existing_product_name = match["existing_product_name"]
            if row.errors:
                row.status, row.message = "error", "请修正错误"
            elif row.existing_product_id is not None:
                row.status, row.message = "skipped", "SKU 已存在，保留已有商品资料"
    new_count = sum(row.status == "new" for row in rows)
    skipped_count = sum(row.status == "skipped" for row in rows)
    error_count = sum(row.status == "error" for row in rows)
    return ProductImportResponse(
        total_rows=len(rows), new_count=new_count, skipped_count=skipped_count,
        error_count=error_count, can_import=bool(new_count and not error_count),
        message=(
            "请修正错误后重新预览" if error_count
            else "全部 SKU 已存在，无需重复导入" if not new_count
            else "校验通过，请核对预览后确认导入"
        ),
        rows=rows,
    )


def _preview_digest(report: ProductImportResponse) -> str:
    encoded = json.dumps(
        [row.model_dump() for row in report.rows],
        ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _token_signature(payload: bytes, secret: str) -> bytes:
    return hmac.new(secret.encode("utf-8"), b"product-import-preview-v1:" + payload, hashlib.sha256).digest()


def _b64(value: bytes) -> str:
    return base64.urlsafe_b64encode(value).rstrip(b"=").decode("ascii")


def issue_product_import_token(
    parsed: ParsedProductImport, report: ProductImportResponse, user_id: int, secret: str,
) -> str:
    payload = json.dumps({
        "version": 1, "user_id": user_id, "file_hash": parsed.file_hash,
        "preview_digest": _preview_digest(report),
        "expires": int(time.time()) + PREVIEW_TTL_SECONDS,
    }, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return _b64(payload) + "." + _b64(_token_signature(payload, secret))


def verify_product_import_token(token: str, parsed: ParsedProductImport, user_id: int, secret: str) -> dict:
    try:
        parts = token.split(".")
        if len(parts) != 2 or len(token) > 1024:
            raise ValueError
        payload, signature = [base64.b64decode(
            part + "=" * (-len(part) % 4), altchars=b"-_", validate=True,
        ) for part in parts]
        if not hmac.compare_digest(signature, _token_signature(payload, secret)):
            raise ValueError
        data = json.loads(payload)
        if data["version"] != 1 or data["user_id"] != user_id:
            raise ValueError
        if data["expires"] <= int(time.time()) or data["file_hash"] != parsed.file_hash:
            raise HTTPException(409, "预览已过期或文件已改变，请重新预览")
        return data
    except HTTPException:
        raise
    except (ValueError, KeyError, TypeError, UnicodeError):
        raise HTTPException(400, "导入预览无效，请重新预览") from None


async def confirm_product_import(
    db: AsyncSession, parsed: ParsedProductImport, token: str, user_id: int, secret: str,
) -> ProductImportResponse:
    claims = verify_product_import_token(token, parsed, user_id, secret)
    report = await preview_product_import(db, parsed)
    if claims.get("preview_digest") != _preview_digest(report):
        raise HTTPException(409, "商品数据已变化，请重新预览后确认导入")
    if not report.can_import:
        raise HTTPException(400, report.message)
    candidates = [row for row in report.rows if row.status == "new"]
    products = [Product(
        sku=row.sku, name=row.name, product_type=row.product_type,
        safe_stock_quantity=row.safe_stock_quantity,
    ) for row in candidates]
    try:
        db.add_all(products)
        await db.flush()
        effective_date = date.today()
        db.add_all([ProductImpairmentRule(
            product_id=product.id, product_type=product.product_type,
            safe_stock_quantity=product.safe_stock_quantity, effective_date=effective_date,
        ) for product in products])
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(409, "导入期间出现 SKU 或规则冲突，整批已回滚，请重新预览") from None
    except Exception:
        await db.rollback()
        raise
    for row, product in zip(candidates, products):
        row.status, row.message, row.product_id = "created", "新增成功", product.id
    report.created_count = len(products)
    report.new_count = 0
    report.can_import = False
    report.message = f"导入完成：新增 {report.created_count} 条，跳过 {report.skipped_count} 条"
    return report


def build_product_import_result(rows: list[ProductImportRow]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品导入结果"
    sheet.append(["原文件行号", *IMPORT_HEADERS, "处理结果", "商品 ID", "已有商品名称", "说明"])
    _style_header(sheet)
    for row in rows:
        notes = "；".join(f"{error.field}：{error.message}" for error in row.errors) or row.message
        sheet.append([
            row.row_number, row.sku, row.name, TYPE_NAMES.get(row.product_type, row.product_type),
            row.safe_stock_quantity, STATUS_NAMES[row.status],
            row.product_id or row.existing_product_id, row.existing_product_name, notes,
        ])
        for index in (2, 3, 4, 6, 8, 9):
            _set_text(sheet.cell(sheet.max_row, index), sheet.cell(sheet.max_row, index).value)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in {"A": 16, "B": 24, "C": 36, "D": 18, "E": 18,
                          "F": 20, "G": 14, "H": 36, "I": 65}.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
