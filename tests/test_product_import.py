import unittest
from io import BytesIO
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.sql.elements import TextClause

from auth import get_current_user
from database import get_db
from models import Product
from routers.products import export_products, router
from schemas import ProductImportRow
from services.product_import import (
    IMPORT_HEADERS, MAX_FILE_BYTES, MAX_IMPORT_ROWS,
    build_product_import_result, build_product_import_template,
    confirm_product_import, issue_product_import_token,
    parse_product_import, preview_product_import, verify_product_import_token,
)


def spreadsheet(rows, headers=IMPORT_HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parsed(rows):
    return parse_product_import("products.xlsx", spreadsheet(rows))


class FakeSession:
    """只模拟会话调用；真实排序规则及事务另由本地 MySQL 验证。"""

    def __init__(self, rows, existing=None, duplicates=()):
        self.rows = rows
        self.existing = existing or {}
        self.duplicates = set(duplicates)
        self.added = []
        self.flush_error = None
        self.commit = AsyncMock()
        self.rollback = AsyncMock()
        self.flush = AsyncMock(side_effect=self._flush)
        self.execute = AsyncMock(side_effect=self._execute)

    async def _execute(self, stmt):
        if isinstance(stmt, TextClause):
            return SimpleNamespace(scalar_one_or_none=lambda: "utf8mb4_0900_ai_ci")
        matches = []
        for row in self.rows:
            if any(error.field == "SKU" for error in row.errors):
                continue
            existing = self.existing.get(row.sku)
            matches.append({
                "row_number": row.row_number,
                "duplicate_count": 2 if row.sku in self.duplicates else 1,
                "existing_product_id": existing[0] if existing else None,
                "existing_product_name": existing[1] if existing else None,
            })
        return SimpleNamespace(mappings=lambda: SimpleNamespace(all=lambda: matches))

    def add_all(self, records):
        self.added.extend(records)

    async def _flush(self):
        if self.flush_error:
            raise self.flush_error
        for index, product in enumerate((item for item in self.added if isinstance(item, Product)), 101):
            product.id = index


class ProductImportParsingTests(unittest.TestCase):
    def test_template_has_text_skus_dropdown_and_examples_outside_data_sheet(self):
        content = build_product_import_template().getvalue()
        workbook = load_workbook(BytesIO(content))
        sheet = workbook["商品导入"]
        self.assertEqual([cell.value for cell in sheet[1]], list(IMPORT_HEADERS))
        self.assertEqual(sheet["A2"].number_format, "@")
        self.assertEqual(len(sheet.data_validations.dataValidation), 2)
        self.assertIn("填写说明", workbook.sheetnames)
        with self.assertRaises(HTTPException) as error:
            parse_product_import("template.xlsx", content)
        self.assertIn("没有可导入", error.exception.detail)

    def test_parsing_trims_strings_preserves_zeroes_and_normalizes_types(self):
        data = parsed([[" 0000123 ", " 商品 A ", "新品", None], ["SKU-B", "商品 B", "stable", "12"]])
        self.assertEqual(data.rows[0].sku, "0000123")
        self.assertEqual(data.rows[0].name, "商品 A")
        self.assertEqual(data.rows[0].product_type, "new")
        self.assertEqual(data.rows[0].safe_stock_quantity, 0)
        self.assertEqual(data.rows[1].safe_stock_quantity, 12)
        self.assertTrue(all(not row.errors for row in data.rows))

    def test_invalid_cells_are_reported_by_original_row_and_field(self):
        data = parsed([
            [None, "", "未知类型", -1],
            [1234, "商品", "新品", 1],
            ["A" * 65, "B" * 256, "稳健商品", 1.5],
            ["SKU-D", "商品", "稳健商品", True],
            ["SKU-E", "商品", "稳健商品", 2147483648],
        ])
        self.assertEqual(data.rows[0].row_number, 2)
        self.assertEqual({error.field for error in data.rows[0].errors},
                         {"SKU", "商品名称", "商品类型", "安全库存"})
        self.assertEqual({error.field for error in data.rows[1].errors}, {"SKU", "安全库存"})
        self.assertTrue(all(row.status == "error" for row in data.rows))

    def test_import_columns_reject_formulas_and_ignore_inventory_columns(self):
        data = parse_product_import(
            "products.xlsx",
            spreadsheet([["SKU-A", "=2+2", "新品", 0, 1000]],
                        (*IMPORT_HEADERS, "库存数量")),
        )
        self.assertEqual(data.rows[0].errors[0].field, "商品名称")
        valid = parse_product_import("export.xlsx", spreadsheet(
            [["SKU-A", "商品", "新品", 0, "=2+2"]], (*IMPORT_HEADERS, "库存数量"),
        ))
        self.assertFalse(valid.rows[0].errors)

    def test_excel_error_cells_are_not_imported_as_literal_product_values(self):
        data = parsed([["#N/A", "#VALUE!", "#REF!", "#NUM!"]])
        row = data.rows[0]
        self.assertEqual(row.status, "error")
        self.assertEqual({error.field for error in row.errors}, {"SKU", "商品名称", "商品类型", "安全库存"})
        self.assertIsNone(row.safe_stock_quantity)

    def test_empty_lines_preserve_excel_row_numbers(self):
        data = parsed([[None, None, None, None], ["SKU-A", "商品", "新品", 0]])
        self.assertEqual(data.rows[0].row_number, 3)

    def test_bad_headers_missing_and_duplicate_columns_are_file_errors(self):
        for headers in (("SKU", "商品名称"), ("SKU", "sku", "商品名称", "商品类型", "安全库存")):
            with self.subTest(headers=headers), self.assertRaises(HTTPException):
                parse_product_import("file.xlsx", spreadsheet([], headers))

    def test_extension_empty_corrupt_and_size_limits(self):
        for name, content, status in (
            ("file.xls", b"data", 400), ("file.csv", b"data", 400),
            ("file.xlsx", b"", 400), ("file.xlsx", b"corrupt", 400),
            ("file.xlsx", b"x" * (MAX_FILE_BYTES + 1), 413),
        ):
            with self.subTest(name=name, size=len(content)), self.assertRaises(HTTPException) as error:
                parse_product_import(name, content)
            self.assertEqual(error.exception.status_code, status)

    def test_row_limit_accepts_2000_and_rejects_2001(self):
        rows = [[f"SKU-{index}", "商品", "新品", 0] for index in range(MAX_IMPORT_ROWS)]
        self.assertEqual(len(parsed(rows).rows), MAX_IMPORT_ROWS)
        with self.assertRaises(HTTPException):
            parsed(rows + [["OVER-LIMIT", "商品", "新品", 0]])

    def test_xml_entities_macros_and_archive_expansion_are_rejected(self):
        for filename, content, cap in (
            ("xl/workbook.xml", b'<!DOCTYPE x [<!ENTITY y "bad">]><x/>', 2048),
            ("xl/workbook.xml", '<!DOCTYPE x><x/>'.encode("utf-16"), 2048),
            ("xl/vbaProject.bin", b"macro", 2048),
            ("xl/workbook.xml", b"x" * 2049, 2048),
        ):
            with self.subTest(filename=filename, size=len(content)):
                output = BytesIO()
                with ZipFile(output, "w", ZIP_DEFLATED) as archive:
                    archive.writestr(filename, content)
                with patch("services.product_import.MAX_UNCOMPRESSED_BYTES", cap), self.assertRaises(HTTPException):
                    parse_product_import("file.xlsx", output.getvalue())

    def test_result_export_keeps_text_and_does_not_create_formulas(self):
        workbook = load_workbook(build_product_import_result([
            ProductImportRow(row_number=2, sku="000123", name="=2+2", product_type="new",
                             safe_stock_quantity=0, status="created", product_id=1),
        ]), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet["B2"].value, "000123")
        self.assertEqual(sheet["C2"].value, "=2+2")
        self.assertEqual(sheet["C2"].data_type, "s")


class ProductImportTransactionTests(unittest.IsolatedAsyncioTestCase):
    async def test_product_export_preserves_imported_text_values(self):
        product = Product(id=1, sku="000123", name="=2+2", product_type="new",
                          safe_stock_quantity=0, created_at=datetime(2026, 10, 12))
        db = SimpleNamespace(execute=AsyncMock(side_effect=[
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: [product])),
            SimpleNamespace(all=lambda: []),
        ]))
        response = await export_products(product_ids=None, db=db, user=SimpleNamespace(role="viewer"))
        content = b"".join([chunk async for chunk in response.body_iterator])
        sheet = load_workbook(BytesIO(content), data_only=False).active
        self.assertEqual(sheet["C2"].value, "000123")
        self.assertEqual(sheet["D2"].value, "=2+2")
        self.assertEqual(sheet["D2"].data_type, "s")
        self.assertEqual(sheet["C2"].number_format, "@")

    async def test_preview_does_not_write_and_existing_skus_are_skipped(self):
        data = parsed([["SKU-A", "新名字", "新品", 0], ["SKU-B", "商品 B", "稳健商品", 12]])
        db = FakeSession(data.rows, {"SKU-A": (1, "已有名字")})
        report = await preview_product_import(db, data)
        self.assertEqual((report.new_count, report.skipped_count, report.error_count), (1, 1, 0))
        self.assertEqual(report.rows[0].existing_product_name, "已有名字")
        self.assertEqual(db.added, [])
        db.commit.assert_not_awaited()
        self.assertEqual(data.rows[0].status, "new")

    async def test_file_duplicates_block_the_whole_batch_even_if_sku_already_exists(self):
        data = parsed([["SKU-A", "A", "新品", 0], ["SKU-A", "B", "新品", 0], ["SKU-C", "C", "新品", 0]])
        report = await preview_product_import(FakeSession(data.rows, duplicates={"SKU-A"}), data)
        self.assertFalse(report.can_import)
        self.assertEqual((report.error_count, report.new_count), (2, 1))

    async def test_confirm_creates_products_and_rules_in_one_commit(self):
        data = parsed([["SKU-A", "A", "新品", 0], ["SKU-B", "B", "稳健商品", 12]])
        db = FakeSession(data.rows, {"SKU-A": (1, "保留")})
        preview = await preview_product_import(db, data)
        token = issue_product_import_token(data, preview, 1, "test-signing-key")
        result = await confirm_product_import(db, data, token, 1, "test-signing-key")
        self.assertEqual((result.created_count, result.skipped_count), (1, 1))
        self.assertEqual(len(db.added), 2)
        self.assertEqual(db.added[1].product_id, db.added[0].id)
        self.assertEqual(db.added[1].product_type, "stable")
        self.assertEqual(db.added[1].safe_stock_quantity, 12)
        db.commit.assert_awaited_once()
        self.assertEqual(result.rows[1].status, "created")

    async def test_changed_database_requires_new_preview_without_writes(self):
        data = parsed([["SKU-A", "A", "新品", 0]])
        db = FakeSession(data.rows)
        token = issue_product_import_token(data, await preview_product_import(db, data), 1, "test-key")
        db.existing = {"SKU-A": (2, "其他用户添加")}
        with self.assertRaises(HTTPException) as error:
            await confirm_product_import(db, data, token, 1, "test-key")
        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(db.added, [])
        db.commit.assert_not_awaited()

    async def test_constraint_conflict_rolls_back_entire_transaction(self):
        data = parsed([["SKU-A", "A", "新品", 0]])
        db = FakeSession(data.rows)
        token = issue_product_import_token(data, await preview_product_import(db, data), 1, "test-key")
        db.flush_error = IntegrityError("insert", {}, Exception("duplicate"))
        with self.assertRaises(HTTPException) as error:
            await confirm_product_import(db, data, token, 1, "test-key")
        self.assertEqual(error.exception.status_code, 409)
        db.rollback.assert_awaited_once()
        db.commit.assert_not_awaited()

    async def test_rule_commit_failure_also_rolls_back(self):
        data = parsed([["SKU-A", "A", "新品", 0]])
        db = FakeSession(data.rows)
        token = issue_product_import_token(data, await preview_product_import(db, data), 1, "test-key")
        db.commit.side_effect = RuntimeError("rule write failed")
        with self.assertRaises(RuntimeError):
            await confirm_product_import(db, data, token, 1, "test-key")
        db.rollback.assert_awaited_once()

    async def test_all_existing_or_invalid_rows_have_no_confirmation_token(self):
        data = parsed([["SKU-A", "A", "新品", 0]])
        report = await preview_product_import(FakeSession(data.rows, {"SKU-A": (1, "A")}), data)
        self.assertFalse(report.can_import)
        invalid = parsed([["SKU-B", "", "新品", 0]])
        report = await preview_product_import(FakeSession(invalid.rows), invalid)
        self.assertFalse(report.can_import)
        self.assertEqual(report.error_count, 1)


class ProductImportTokenTests(unittest.TestCase):
    def setUp(self):
        self.data = parsed([["SKU-A", "A", "新品", 0]])
        self.report = SimpleNamespace(rows=self.data.rows)
        with patch("services.product_import.time.time", return_value=100):
            self.token = issue_product_import_token(self.data, self.report, 7, "test-key")

    def test_token_is_bound_to_user_file_signature_and_expiry(self):
        changed = parsed([["SKU-B", "B", "新品", 0]])
        for token, data, user, secret, now in (
            (self.token, self.data, 8, "test-key", 101),
            (self.token, changed, 7, "test-key", 101),
            (self.token, self.data, 7, "wrong-key", 101),
            (self.token, self.data, 7, "test-key", 1000),
            ("invalid.token", self.data, 7, "test-key", 101),
        ):
            with self.subTest(user=user, now=now), patch("services.product_import.time.time", return_value=now):
                with self.assertRaises(HTTPException):
                    verify_product_import_token(token, data, user, secret)
        with patch("services.product_import.time.time", return_value=101):
            self.assertEqual(verify_product_import_token(self.token, self.data, 7, "test-key")["user_id"], 7)


class ProductImportApiTests(unittest.TestCase):
    def setUp(self):
        self.content = spreadsheet([["SKU-A", "商品", "新品", 0]])
        self.db = FakeSession(parse_product_import("products.xlsx", self.content).rows)
        app = FastAPI()
        app.include_router(router)

        async def fake_db():
            yield self.db

        async def fake_user():
            return SimpleNamespace(id=1, role="viewer")

        app.dependency_overrides[get_db] = fake_db
        app.dependency_overrides[get_current_user] = fake_user
        self.client = TestClient(app)

    def test_multipart_preview_confirm_and_result_download(self):
        response = self.client.post("/products/import/preview", files={"file": ("products.xlsx", self.content)})
        self.assertEqual(response.status_code, 200)
        report = response.json()
        self.assertEqual(report["new_count"], 1)
        confirmed = self.client.post("/products/import", files={"file": ("products.xlsx", self.content)},
                                     data={"preview_token": report["preview_token"]})
        self.assertEqual(confirmed.status_code, 200)
        self.assertEqual(confirmed.json()["created_count"], 1)
        downloaded = self.client.post("/products/import/result", json={"rows": confirmed.json()["rows"]})
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(downloaded.content)).active["F2"].value, "已新增")

    def test_template_and_missing_preview_token(self):
        response = self.client.get("/products/import-template")
        self.assertEqual(response.status_code, 200)
        self.assertIn("商品导入", load_workbook(BytesIO(response.content)).sheetnames)
        response = self.client.post("/products/import", files={"file": ("products.xlsx", self.content)})
        self.assertEqual(response.status_code, 422)

    def test_invalid_rows_preview_has_no_token_and_file_errors_use_http_400(self):
        self.content = spreadsheet([["SKU-A", "", "新品", 0]])
        self.db.rows = parse_product_import("products.xlsx", self.content).rows
        response = self.client.post("/products/import/preview", files={"file": ("products.xlsx", self.content)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["error_count"], 1)
        self.assertIsNone(response.json()["preview_token"])
        response = self.client.post("/products/import/preview", files={"file": ("products.xls", self.content)})
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
