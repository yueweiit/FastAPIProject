"""隔离SQLite/真实路由验证导入归属、越店拦截和导出店铺；不启动main。"""
import csv
import io
import unittest
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import patch

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from auth import get_current_user
from database import Base, get_db
from models import (InventoryBatch, PlatformSkuComponent, PlatformSkuMapping, Product, Sale,
                    SalesImportBatch, SettlementEntry, Store, User)
from routers import batches, sales


class Facade:
    def __init__(self, session):
        self.session = session
        self.info = session.info
    async def get(self, model, key): return self.session.get(model, key)
    async def execute(self, statement): return self.session.execute(statement)
    async def scalar(self, statement): return self.session.scalar(statement)
    def add(self, row): self.session.add(row)
    async def flush(self): self.session.flush()
    async def commit(self): self.session.commit()
    async def rollback(self): self.session.rollback()


def order_file():
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(['结算日期','结算单 ID','付款 ID','状态','货币','交易类型','订单ID/调整单ID',
                     'SKU ID','数量','订单创建日期','净商品销售额','结算总金额'])
    for suffix, sku in [('OK','PLAT'),('PENDING','UNKNOWN')]:
        writer.writerow(['2026-09-02','BILL','PAY','已结算','CNY','订单',suffix,sku,1,'2026-09-01',20,18])
    return out.getvalue().encode('utf-8-sig')


class ImportStoreTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://')
        Base.metadata.create_all(self.engine)
        self.session = Session(self.engine, expire_on_commit=False)
        self.session.add_all([Store(id=i,name=f'店铺{i}',platform='test',platform_store_id=str(i),
                                   country_or_region='CN',settlement_currency='CNY',is_active=i!=9) for i in (7,8,9)])
        self.session.flush()
        self.session.add_all([User(id=1,username='admin',password_hash='test',role='admin'),
                              User(id=2,username='operator',password_hash='test',role='operator',store_id=7),
                              Product(id=1,sku='SKU',name='商品')])
        self.session.flush()
        for key, owner in ((1,2),(2,1),(3,None)):
            self.session.add(InventoryBatch(id=key,product_id=1,user_id=owner,batch_no=f'B{key}',
                quantity=10,remaining_quantity=10,purchase_price=10,unit_cost=10,arrived_at=datetime(2026,9,1)))
        mapping=PlatformSkuMapping(platform='tiktok_shop',platform_sku_id='PLAT',is_active=True)
        mapping.components=[PlatformSkuComponent(product_id=1,quantity_per_sale=1)]
        self.session.add(mapping)
        self.session.commit()
        self.actor=SimpleNamespace(id=1,role='admin',store_id=None)
        self.db=Facade(self.session)
        app=FastAPI()
        app.include_router(sales.router)
        app.include_router(batches.router)
        async def user(): return self.actor
        async def db(): yield self.db
        app.dependency_overrides[get_current_user]=user
        app.dependency_overrides[get_db]=db
        self.app=app

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    async def request(self, method, path, **kwargs):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app),base_url='http://test') as client:
            return await client.request(method,path,**kwargs)

    async def upload(self, store=None, detail=False):
        return await self.request('POST','/sales/import',data={} if store is None else {'store_id':str(store)},
            files={'file':('orders.csv',order_file() if detail else b'SKU,2026-09-01,1\n','text/csv')})

    async def test_admin_requires_explicit_store_even_when_bound(self):
        for bound in (None,7):
            self.actor.store_id=bound
            response=await self.upload()
            self.assertEqual(response.status_code,400)
            self.assertEqual(response.json()['detail']['code'],'import_store_required')
        self.assertEqual(self.session.scalars(select(SalesImportBatch)).all(),[])

    async def test_admin_attribution_persists_for_header_sale_pending_and_export(self):
        self.actor.store_id=7
        response=await self.upload(8,detail=True)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual((response.json()['sales_created'],response.json()['pending_confirmation']),(1,1))
        for model in (Sale,SalesImportBatch,SettlementEntry):
            self.assertEqual({r.store_id for r in self.session.scalars(select(model))},{8})
        self.actor=SimpleNamespace(id=3,role='viewer',store_id=8)
        history=await self.request('GET','/sales/imports')
        self.assertEqual(history.json()[0]['store_name'],'店铺8')
        pending=await self.request('GET','/sales/pending-confirmations/export')
        sheet=load_workbook(io.BytesIO(pending.content)).active
        self.assertEqual(sheet['J2'].value,'店铺8')
        self.actor.store_id=7
        self.assertEqual((await self.request('GET','/sales/imports')).json(),[])
        sheet=load_workbook(io.BytesIO((await self.request('GET','/sales/pending-confirmations/export')).content)).active
        self.assertEqual(sheet.max_row,1)

    async def test_unavailable_or_invalid_admin_store_writes_nothing(self):
        for store,status in ((9,400),(999,400),(0,422),('invalid',422)):
            self.assertEqual((await self.upload(store)).status_code,status)
        self.assertEqual(self.session.scalars(select(Sale)).all(),[])
        self.assertEqual(self.session.get(InventoryBatch,1).remaining_quantity,10)

    async def test_operator_legacy_client_without_store_uses_binding(self):
        self.actor=SimpleNamespace(id=2,role='operator',store_id=7)
        response=await self.upload()
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.session.scalar(select(Sale)).store_id,7)
        self.assertEqual(self.session.scalar(select(SalesImportBatch)).store_id,7)

    async def test_operator_cannot_override_binding_or_import_unbound(self):
        self.actor=SimpleNamespace(id=2,role='operator',store_id=7)
        self.assertEqual((await self.upload(8)).status_code,403)
        self.actor.store_id=None
        self.assertEqual((await self.upload(7)).status_code,400)
        self.assertEqual(self.session.scalars(select(SalesImportBatch)).all(),[])

    async def test_viewer_cannot_import(self):
        self.actor=SimpleNamespace(id=2,role='viewer',store_id=7)
        self.assertEqual((await self.upload(7)).status_code,403)

    async def test_batch_export_appends_explicit_current_owner_store(self):
        self.actor=SimpleNamespace(id=2,role='viewer',store_id=7)
        response=await self.request('GET','/batches/export?batch_ids=1&batch_ids=2&batch_ids=3')
        self.assertEqual(response.status_code,200)
        sheet=load_workbook(io.BytesIO(response.content)).active
        self.assertEqual((sheet['E1'].value,sheet['U1'].value),('批次号','创建人ID'))
        self.assertEqual(sheet['V1'].value,'店铺（创建人当前绑定）')
        self.assertEqual([sheet.cell(i,22).value for i in (2,3,4)],['未关联店铺','未关联店铺','店铺7'])
        self.session.get(Store,7).name='=1+1'
        self.session.commit()
        sheet=load_workbook(io.BytesIO((await self.request('GET','/batches/export?batch_ids=1')).content)).active
        self.assertEqual((sheet['V2'].value,sheet['V2'].data_type),('=1+1','s'))

    async def test_batch_export_releases_read_transaction_before_writing_workbook(self):
        original_workbook = batches.Workbook
        saved = []

        def checked_workbook():
            workbook = original_workbook()
            original_save = workbook.save

            def checked_save(output):
                self.assertFalse(self.session.in_transaction())
                saved.append(True)
                return original_save(output)

            workbook.save = checked_save
            return workbook

        with patch.object(batches, 'Workbook', new=checked_workbook):
            response = await self.request('GET', '/batches/export?batch_ids=1')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(saved, [True])

    def test_pending_export_does_not_hide_missing_store_in_a_group(self):
        entry=SimpleNamespace(platform_sku_id='X',mapping_error='缺映射',quantity=1,product_name='',sku_name='',
            source_file='a.csv',transaction_type='订单',currency='CNY',net_product_sales=20,settlement_total=18,
            store=None)
        rows=sales._group_pending_confirmations([entry])
        self.assertEqual(rows[0]['store_names'],'未关联店铺')
        attributed=SimpleNamespace(**{**entry.__dict__,'store':SimpleNamespace(name='店铺7')})
        rows=sales._group_pending_confirmations([entry,attributed])
        self.assertEqual(set(rows[0]['store_names'].split('、')),{'未关联店铺','店铺7'})
