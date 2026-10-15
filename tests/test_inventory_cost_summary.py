import unittest
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook

from auth import RequireAnyRole
from database import get_db
from routers.reports import (
    INVENTORY_IMPAIRMENT_HEADERS, INVENTORY_IMPAIRMENT_SHEET,
    INVENTORY_IMPAIRMENT_SUMMARY_SHEET, _build_inventory_impairment_summary_sheet,
    _build_inventory_impairment_workbook, _load_inventory_impairment_report_rows,
    _summarize_inventory_impairment_rows, _workbook_bytes_with_formula_cache,
    inventory_impairment_report, router,
)

REPORT_DATE = date(2026, 9, 30)


class InventoryCostSummaryTests(unittest.IsolatedAsyncioTestCase):
    async def _rows(self, *, sold_cost="400", store_id=None):
        batches = [
            SimpleNamespace(id=1, product_id=10, quantity=100, unit_cost=Decimal("10"),
                            arrived_at=datetime(2026, 8, 21), batch_no="A"),
            SimpleNamespace(id=2, product_id=10, quantity=200, unit_cost=Decimal("12"),
                            arrived_at=datetime(2026, 9, 20), batch_no="B"),
            SimpleNamespace(id=3, product_id=10, quantity=50, unit_cost=Decimal("8"),
                            arrived_at=datetime(2026, 7, 1), batch_no="C"),
        ]
        batch_rows = [(batch, "SKU-1", "同一物料", "店铺A", "new", 0, 7) for batch in batches]
        rule = SimpleNamespace(product_id=10, product_type="new", safe_stock_quantity=0,
                               effective_date=date(2026, 7, 1))
        deductions = [
            SimpleNamespace(detail_id=1, sale_id=1, batch_id=1, sold_at=datetime(2026, 8, 22),
                            report_quantity=20, saved_unit_cost=Decimal(sold_cost) / 40, sale_store_id=7),
            SimpleNamespace(detail_id=2, sale_id=2, batch_id=1, sold_at=datetime(2026, 8, 23),
                            report_quantity=20, saved_unit_cost=Decimal(sold_cost) / 40, sale_store_id=7),
            SimpleNamespace(detail_id=3, sale_id=3, batch_id=3, sold_at=datetime(2026, 7, 3),
                            report_quantity=50, saved_unit_cost=Decimal("8"), sale_store_id=7),
        ]
        db = SimpleNamespace(execute=AsyncMock(side_effect=[
            SimpleNamespace(all=lambda: batch_rows),
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: [rule])),
            SimpleNamespace(all=lambda: deductions),
        ]), commit=AsyncMock())
        rows = await _load_inventory_impairment_report_rows(db, REPORT_DATE, store_id)
        db.commit.assert_not_awaited()
        return rows, db

    async def test_confirmed_balance_preserves_saved_outgoing_cost(self):
        rows, db = await self._rows(sold_cost="360.1234")
        with patch("routers.reports._load_inventory_impairment_report_rows", new_callable=AsyncMock, return_value=rows):
            response = await inventory_impairment_report(REPORT_DATE, None, None)
        first = response[0]
        self.assertEqual(first["original_quantity"], 100)
        self.assertEqual(first["original_inventory_amount"], Decimal("1000"))
        self.assertEqual(first["sold_quantity"], 40)
        self.assertEqual(first["sold_cost_amount"], Decimal("360.1234"))
        self.assertEqual(first["quantity"], 60)
        self.assertEqual(first["inventory_amount"], Decimal("600"))
        self.assertEqual(first["impairment_amount"], Decimal("120"))
        self.assertEqual(first["book_value"], Decimal("480"))
        self.assertEqual(first["monthly_impairment_amount"], Decimal("90"))
        self.assertEqual(first["book_value_difference"], Decimal("-90"))
        statement = db.execute.call_args_list[2].args[0]
        self.assertIn("sale_cost_details.unit_cost AS saved_unit_cost", str(statement))
        self.assertIn("ORDER BY sales.sold_at, sales.id, sale_cost_details.id", str(statement))
        dates = [v for v in statement.compile().params.values() if isinstance(v, datetime)]
        self.assertEqual(dates, [datetime(2026, 10, 1)])
        self.assertIn(datetime(2026, 10, 1), db.execute.call_args_list[0].args[0].compile().params.values())
        self.assertFalse(any("accounting_period" in str(call.args[0]) for call in db.execute.call_args_list))

    async def test_receipts_and_sold_out_history_are_in_both_views(self):
        rows, _ = await self._rows()
        self.assertEqual([row["batch_no"] for row in rows], ["A", "B", "C"])
        self.assertEqual([row["unit_cost"] for row in rows], [Decimal("10"), Decimal("12"), Decimal("8")])
        self.assertEqual([row["impairment_rate"] for row in rows], [Decimal("0.20"), Decimal("0.05"), Decimal("0")])
        self.assertEqual(rows[2]["quantity"], 0)
        self.assertEqual(rows[2]["impairment_amount"], Decimal("0"))
        summary = _summarize_inventory_impairment_rows(rows)[0]
        self.assertEqual(summary["batch_count"], 3)
        self.assertEqual(summary["original_quantity"], 350)
        self.assertEqual(summary["original_inventory_amount"], Decimal("3800"))
        self.assertEqual(summary["sold_quantity"], 90)
        self.assertEqual(summary["sold_cost_amount"], Decimal("800"))
        self.assertEqual(summary["quantity"], 260)
        self.assertEqual(summary["inventory_amount"], Decimal("3000"))
        self.assertEqual(summary["impairment_amount"], Decimal("240"))
        self.assertEqual(summary["monthly_impairment_amount"], Decimal("210"))
        self.assertEqual(summary["book_value"], Decimal("2760"))
        self.assertEqual(summary["book_value_difference"], Decimal("2190"))
        self.assertEqual(summary["original_inventory_amount"] - summary["sold_cost_amount"], summary["inventory_amount"])
        self.assertEqual(_summarize_inventory_impairment_rows([]), [])

    async def test_defensive_unknown_comparison_is_not_filled_with_zero(self):
        rows, _ = await self._rows()
        rows[2]["previous_book_value"] = None
        summary = _summarize_inventory_impairment_rows(rows)[0]
        self.assertIsNone(summary["book_value_difference"])
        self.assertIn("1个批次缺少比较基准", summary["comparison_note"])
        workbook = Workbook()
        _build_inventory_impairment_summary_sheet(workbook, [summary], REPORT_DATE, use_active_sheet=True)
        sheet = load_workbook(_workbook_bytes_with_formula_cache(workbook, {}), data_only=True).active
        self.assertIsNone(sheet["O2"].value)
        self.assertIsNone(sheet["O3"].value)
        self.assertEqual(sheet["N3"].value, 2760)

    async def test_sku_grouping_does_not_merge_names_or_replace_batch_prices(self):
        rows, _ = await self._rows()
        other = {**rows[1], "product_id": 20, "sku": "SKU-2", "store_name": "店铺B"}
        rows[1]["store_name"] = "店铺B"
        summaries = _summarize_inventory_impairment_rows(rows + [other])
        self.assertEqual([row["sku"] for row in summaries], ["SKU-1", "SKU-2"])
        self.assertEqual(summaries[0]["store_name"], "店铺A / 店铺B")
        self.assertEqual(summaries[0]["book_value"], Decimal("2760"))
        self.assertEqual(summaries[1]["book_value"], Decimal("2280"))

    async def test_excel_remaining_cost_formula_cache_headers_and_literal_text(self):
        rows, _ = await self._rows()
        rows[0]["sku"] = '=HYPERLINK("https://example.invalid")'
        workbook, cache = _build_inventory_impairment_workbook(rows, REPORT_DATE)
        output = _workbook_bytes_with_formula_cache(workbook, cache)
        sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(sheet["G2"].value, 100)
        self.assertEqual(sheet["H2"].value, 1000)
        self.assertEqual(sheet["I2"].value, 40)
        self.assertEqual(sheet["J2"].value, 400)
        self.assertEqual(sheet["O2"].value, 120)
        self.assertEqual(sheet["P2"].value, 480)
        self.assertEqual(sheet["R2"].value, 90)
        self.assertEqual(sheet["S4"].value, "已售完")
        self.assertEqual(sheet["O4"].value, 0)
        self.assertEqual(sheet["M5"].value, 3000)
        self.assertEqual(sheet["P5"].value, 2760)
        self.assertEqual(sheet["N2"].number_format, "0.00%")
        self.assertEqual(sheet["T2"].data_type, "s")
        for cell in ("E5", "F5", "N5"):
            self.assertIsNone(sheet[cell].value)
        output.seek(0)
        formulas = load_workbook(output)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(formulas["B2"].data_type, "s")
        self.assertEqual(formulas["P2"].value, '=IF(OR(M2="",O2=""),"",M2-O2)')
        self.assertEqual(formulas["P5"].value, '=SUM(P2:P4)')
        self.assertEqual([cell.value for cell in formulas[1]], INVENTORY_IMPAIRMENT_HEADERS)
        self.assertEqual(formulas["P1"].value, "账面价值(元)")
        _build_inventory_impairment_summary_sheet(workbook, _summarize_inventory_impairment_rows(rows), REPORT_DATE)
        saved = load_workbook(_workbook_bytes_with_formula_cache(workbook, cache), data_only=True)
        self.assertEqual(saved[INVENTORY_IMPAIRMENT_SHEET]["P5"].value, 2760)
        self.assertEqual(saved[INVENTORY_IMPAIRMENT_SUMMARY_SHEET]["N3"].value, 2760)

    async def test_new_batch_previous_zero_and_historical_comparison_without_snapshot(self):
        rows, _ = await self._rows()
        self.assertEqual(rows[0]["previous_book_value"], Decimal("570"))
        self.assertEqual(rows[1]["previous_book_value"], Decimal("0"))
        self.assertIn("本月新增批次", rows[1]["comparison_note"])
        self.assertEqual(rows[2]["previous_book_value"], Decimal("0"))

    async def test_empty_reports_have_numeric_totals_without_circular_formulas(self):
        workbook, cache = _build_inventory_impairment_workbook([], REPORT_DATE)
        sheet = load_workbook(_workbook_bytes_with_formula_cache(workbook, cache), data_only=True).active
        self.assertEqual(sheet["A2"].value, "合计")
        for cell in ("G2", "H2", "I2", "J2", "K2", "L2", "M2", "O2", "P2", "Q2", "R2"):
            self.assertEqual(sheet[cell].value, 0)
        workbook = Workbook()
        _build_inventory_impairment_summary_sheet(workbook, [], REPORT_DATE, use_active_sheet=True)
        for column in range(4, 16):
            self.assertEqual(workbook.active.cell(2, column).value, 0)

    async def test_local_store_scope_is_preserved(self):
        rows, db = await self._rows(store_id=7)
        stmt = db.execute.call_args_list[0].args[0]
        self.assertNotIn("WHERE users.store_id =", str(stmt))
        self.assertTrue(all(row["store_id"] == 7 for row in rows))

    def _client(self, role="viewer", db=None):
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[RequireAnyRole] = lambda: SimpleNamespace(role=role, store_id=7)
        return AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver")

    async def test_http_both_report_and_export_views(self):
        rows, _ = await self._rows()
        async with self._client() as client:
            with patch("routers.reports._load_inventory_impairment_report_rows", new_callable=AsyncMock, return_value=rows) as load:
                detail = await client.get("/reports/inventory-impairment?report_date=2026-09-30")
                self.assertEqual(detail.status_code, 200)
                self.assertEqual(len(detail.json()), 3)
                self.assertEqual(detail.json()[0]["book_value"], 480)
                response = await client.get("/reports/inventory-impairment/summary?report_date=2026-09-30")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()[0]["book_value"], 2760)
                load.assert_awaited_with(None, REPORT_DATE)
                for view, title, cell in (("batch", INVENTORY_IMPAIRMENT_SHEET, "P5"),
                                          ("sku", INVENTORY_IMPAIRMENT_SUMMARY_SHEET, "N2")):
                    exported = await client.get(f"/reports/inventory-impairment/export?report_date=2026-09-30&view={view}")
                    self.assertEqual(exported.status_code, 200)
                    sheet = load_workbook(BytesIO(exported.content), data_only=True).active
                    self.assertEqual(sheet.title, title)
                    self.assertEqual(sheet[cell].value, 2760)
                self.assertEqual((await client.get("/reports/inventory-impairment/export?report_date=2026-09-30&view=other")).status_code, 422)
                load.reset_mock()
                future = (date.today() + timedelta(days=1)).isoformat()
                for suffix in ("", "/summary", "/export"):
                    self.assertEqual((await client.get(f"/reports/inventory-impairment{suffix}?report_date={future}")).status_code, 400)
                load.assert_not_awaited()

    async def test_financial_attachment_preserves_scope_and_sheet_order(self):
        rows, _ = await self._rows()
        for role, expected_scope in (("admin", None), ("viewer", 7)):
            fake_db = SimpleNamespace(scalar=AsyncMock(return_value=None),
                execute=AsyncMock(return_value=SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: []))))
            store = SimpleNamespace(id=7, name="店铺A", platform="tiktok_shop")
            async with self._client(role, fake_db) as client:
                with patch("routers.reports._resolve_report_store", new_callable=AsyncMock, return_value=store), \
                     patch("routers.reports._store_profit_loss_auto_values", new_callable=AsyncMock, return_value={}), \
                     patch("routers.reports._product_line_profit_data", new_callable=AsyncMock, return_value=[]), \
                     patch("routers.reports._load_inventory_impairment_report_rows", new_callable=AsyncMock, return_value=rows) as load:
                    response = await client.get("/reports/financial-attachments/export?report_date=2026-09-30")
                    self.assertEqual(response.status_code, 200)
                    load.assert_awaited_once_with(fake_db, REPORT_DATE, expected_scope)
                    workbook = load_workbook(BytesIO(response.content), data_only=True)
                    self.assertEqual(len(workbook.sheetnames), 5 if role == "admin" else 4)
                    self.assertEqual(workbook.sheetnames[-1], INVENTORY_IMPAIRMENT_SUMMARY_SHEET)
                    self.assertEqual(workbook[INVENTORY_IMPAIRMENT_SHEET]["G5"].value, 350)
                    self.assertEqual(workbook[INVENTORY_IMPAIRMENT_SHEET]["P5"].value, 2760)
                    self.assertEqual(workbook[INVENTORY_IMPAIRMENT_SUMMARY_SHEET]["N2"].value, 2760)


if __name__ == "__main__":
    unittest.main()
