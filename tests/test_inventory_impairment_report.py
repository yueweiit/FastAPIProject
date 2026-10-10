import asyncio
import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from routers.reports import (
    INVENTORY_IMPAIRMENT_HEADERS,
    INVENTORY_IMPAIRMENT_SHEET,
    _build_inventory_impairment_workbook,
    _load_inventory_impairment_report_rows,
    _previous_month_end,
    _workbook_bytes_with_formula_cache,
    inventory_impairment_report,
)
from services.inventory_impairment import (
    ImpairmentRule,
    InventoryLayer,
    PRODUCT_TYPE_NEW,
    PRODUCT_TYPE_STABLE,
    batch_impairments,
    impairment_rate,
)


class InventoryImpairmentReportTests(unittest.TestCase):
    def _report_rows(self, *, arrived_at, quantity, deducted=0, snapshot=None, rules=()):
        batch = SimpleNamespace(
            id=1, product_id=10, quantity=quantity, unit_cost=Decimal("10"),
            arrived_at=arrived_at, batch_no="BATCH-001",
        )
        results = [
            [snapshot] if snapshot is not None else [],
            [(batch, "SKU-001", "测试商品", "测试店铺", PRODUCT_TYPE_NEW, 0, 1)],
            list(rules),
            [SimpleNamespace(batch_id=1, report_quantity=deducted)],
        ]
        db = SimpleNamespace(execute=AsyncMock(side_effect=[
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: results[0])),
            SimpleNamespace(all=lambda: results[1]),
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: results[2])),
            SimpleNamespace(all=lambda: results[3]),
        ]))
        with patch("routers.reports.ensure_period_snapshot", new_callable=AsyncMock) as ensure:
            ensure.return_value = SimpleNamespace(id=1)
            return asyncio.run(_load_inventory_impairment_report_rows(db, date(2026, 9, 30)))

    def _report_response(self, rows):
        with patch("routers.reports._load_inventory_impairment_report_rows", new_callable=AsyncMock) as load:
            load.return_value = rows
            return asyncio.run(inventory_impairment_report(date(2026, 9, 30), None, None))

    def test_monthly_provision_on_remaining_stock_is_separate_from_sales_reduction(self):
        snapshot = SimpleNamespace(batch_id=1, quantity=100, unit_cost=Decimal("10"))
        rows = self._report_rows(
            arrived_at=datetime(2026, 8, 1), quantity=100, deducted=50, snapshot=snapshot,
        )
        response = self._report_response(rows)[0]
        self.assertEqual(response["impairment_amount"], Decimal("150"))
        self.assertEqual(response["monthly_impairment_amount"], Decimal("75"))
        self.assertEqual(response["book_value"], Decimal("350"))
        self.assertEqual(response["book_value_difference"], Decimal("-500"))

        workbook, formula_cache = _build_inventory_impairment_workbook(rows, date(2026, 9, 30))
        output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
        sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(sheet["K2"].value, 150)
        self.assertEqual(sheet["M2"].value, -500)
        self.assertEqual(sheet["N2"].value, 75)
        self.assertEqual(sheet["N3"].value, 75)

    def test_monthly_provision_preserves_rule_history_and_safe_stock(self):
        layers = [InventoryLayer(
            1, 10, 1, date(2026, 8, 1), 200, PRODUCT_TYPE_STABLE, 100,
            (
                ImpairmentRule(PRODUCT_TYPE_NEW, 0, date(2026, 8, 1)),
                ImpairmentRule(PRODUCT_TYPE_STABLE, 100, date(2026, 9, 11)),
            ),
        )]
        result = batch_impairments(layers, date(2026, 9, 30), period_start=date(2026, 9, 1))[1]
        self.assertEqual(result.impairment_units, Decimal("60"))
        self.assertEqual(result.period_impairment_units, Decimal("30"))

    def test_monthly_provision_obeys_cap_and_backfilled_first_rule(self):
        capped = batch_impairments([
            InventoryLayer(1, 10, 1, date(2026, 7, 1), 100, PRODUCT_TYPE_STABLE, 0),
        ], date(2026, 9, 30), period_start=date(2026, 9, 1))[1]
        self.assertEqual(capped.impairment_units, Decimal("90"))
        self.assertEqual(capped.period_impairment_units, Decimal("29"))

        backfilled = batch_impairments([
            InventoryLayer(
                1, 10, 1, date(2026, 8, 1), 100, PRODUCT_TYPE_NEW, 0,
                (ImpairmentRule(PRODUCT_TYPE_NEW, 0, date(2026, 9, 22)),),
            ),
        ], date(2026, 9, 30), period_start=date(2026, 9, 1))[1]
        self.assertEqual(backfilled.impairment_units, Decimal("30"))
        self.assertEqual(backfilled.period_impairment_units, Decimal("15"))

    def test_monthly_provision_handles_partial_month_across_year_boundary(self):
        result = batch_impairments([
            InventoryLayer(1, 10, 1, date(2025, 12, 1), 100, PRODUCT_TYPE_NEW, 0),
        ], date(2026, 1, 15), period_start=date(2026, 1, 1))[1]
        self.assertEqual(result.impairment_units, Decimal("22.5"))
        self.assertEqual(result.period_impairment_units, Decimal("7.5"))

    def test_missing_snapshot_is_unknown_in_page_and_export(self):
        rows = self._report_rows(arrived_at=datetime(2026, 8, 1), quantity=100)
        response = self._report_response(rows)[0]
        self.assertIsNone(response["book_value_difference"])
        self.assertEqual(response["comparison_note"], "缺少上月快照，无法比较")
        self.assertEqual(response["monthly_impairment_amount"], Decimal("150"))

        workbook, formula_cache = _build_inventory_impairment_workbook(rows, date(2026, 9, 30))
        output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
        sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertIsNone(sheet["M2"].value)
        self.assertIsNone(sheet["M3"].value)
        self.assertEqual(sheet["N2"].value, 150)
        self.assertIn("缺少上月快照", sheet["P2"].value)
        self.assertIn("上月末库存：待核实", sheet["P2"].value)

    def test_missing_snapshot_keeps_difference_total_unknown_with_other_valid_rows(self):
        missing_rows = self._report_rows(arrived_at=datetime(2026, 8, 1), quantity=100)
        snapshot = SimpleNamespace(batch_id=1, quantity=100, unit_cost=Decimal("10"))
        valid_rows = self._report_rows(
            arrived_at=datetime(2026, 8, 1), quantity=100, deducted=50, snapshot=snapshot,
        )
        workbook, formula_cache = _build_inventory_impairment_workbook(
            missing_rows + valid_rows, date(2026, 9, 30),
        )
        sheet = load_workbook(
            _workbook_bytes_with_formula_cache(workbook, formula_cache), data_only=True,
        )[INVENTORY_IMPAIRMENT_SHEET]
        self.assertIsNone(sheet["M4"].value)
        self.assertEqual(sheet["N4"].value, 225)

    def test_new_batch_has_zero_previous_book_value_and_only_current_month_provision(self):
        rows = self._report_rows(arrived_at=datetime(2026, 9, 10), quantity=100)
        response = self._report_response(rows)[0]
        self.assertEqual(rows[0]["previous_book_value"], Decimal("0"))
        self.assertEqual(response["impairment_amount"], Decimal("100"))
        self.assertEqual(response["monthly_impairment_amount"], Decimal("100"))
        self.assertEqual(response["book_value_difference"], Decimal("900"))
        self.assertEqual(response["comparison_note"], "本月新增批次，上月账面价值为0")

    def test_fully_sold_batch_has_no_provision_for_remaining_stock(self):
        snapshot = SimpleNamespace(batch_id=1, quantity=100, unit_cost=Decimal("10"))
        rows = self._report_rows(
            arrived_at=datetime(2026, 8, 1), quantity=100, deducted=100, snapshot=snapshot,
        )
        response = self._report_response(rows)[0]
        self.assertEqual(response["impairment_amount"], Decimal("0"))
        self.assertEqual(response["monthly_impairment_amount"], Decimal("0"))
        self.assertEqual(response["book_value_difference"], Decimal("-850"))
        self.assertEqual(response["status"], "已售完")

    def test_previous_month_end_handles_year_boundary(self):
        self.assertEqual(_previous_month_end(date(2026, 1, 15)), date(2025, 12, 31))
        self.assertEqual(_previous_month_end(date(2026, 9, 30)), date(2026, 8, 31))

    def test_workbook_matches_first_template_sheet_and_formulas(self):
        workbook, formula_cache = _build_inventory_impairment_workbook([
            {
                "store_name": "测试店铺",
                "sku": "SKU-001",
                "product_name": "测试商品",
                "arrived_at": date(2026, 8, 1),
                "unit_cost": Decimal("12.5"),
                "quantity": 2,
                "previous_month_quantity": 0,
                "batch_no": "BATCH-001",
            }
        ], date(2026, 8, 31))

        output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
        saved = load_workbook(output, data_only=False)
        sheet = saved[INVENTORY_IMPAIRMENT_SHEET]

        self.assertEqual(saved.sheetnames, [INVENTORY_IMPAIRMENT_SHEET])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual([cell.value for cell in sheet[1]], INVENTORY_IMPAIRMENT_HEADERS)
        self.assertEqual(sheet["E2"].value, '=IF(D2="","",DATE(2026,8,31)-D2)')
        self.assertEqual(sheet["H2"].value, 2)
        self.assertEqual(sheet["I2"].value, '=IF(OR(F2="",G2=""),"",F2*G2)')
        self.assertEqual(sheet["J2"].value, '=IF(E2="","",MIN(E2*0.01,0.9))')
        self.assertEqual(sheet["K2"].value, '=IF(OR(F2="",H2="",J2=""),"",F2*H2*J2)')
        self.assertEqual(sheet["L2"].value, '=IF(OR(I2="",K2=""),"",I2-K2)')
        self.assertEqual(
            sheet["M2"].value,
            '=IF(OR(D2="",F2="",L2=""),"",L2-F2*(0-0*MIN(MAX(DATE(2026,7,31)-D2,0)*0.01,0.9)))',
        )
        self.assertIn("J2>=0.6", sheet["O2"].value)
        self.assertEqual(sheet["E3"].value, "=SUM(E2:E2)")
        self.assertEqual(sheet["F3"].value, "=SUM(F2:F2)")
        self.assertEqual(sheet["G3"].value, "=SUM(G2:G2)")
        self.assertEqual(sheet["M3"].value, "=SUM(M2:M2)")
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00C00000")

        output.seek(0)
        cached_sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(cached_sheet["E2"].value, 30)
        self.assertEqual(cached_sheet["H2"].value, 2)
        self.assertEqual(cached_sheet["I2"].value, 25)
        self.assertEqual(cached_sheet["J2"].value, 0.3)
        self.assertEqual(cached_sheet["K2"].value, 7.5)
        self.assertEqual(cached_sheet["L2"].value, 17.5)
        self.assertEqual(cached_sheet["M2"].value, 17.5)
        self.assertEqual(cached_sheet["N2"].value, 7.5)
        self.assertEqual(cached_sheet["O2"].value, "正常计提")
        self.assertEqual(cached_sheet["H3"].value, 2)

    def test_previous_month_book_value_uses_actual_previous_inventory_quantity(self):
        workbook, formula_cache = _build_inventory_impairment_workbook([
            {
                "store_name": "测试店铺",
                "sku": "SKU-002",
                "product_name": "历史库存商品",
                "arrived_at": date(2026, 7, 1),
                "unit_cost": Decimal("12.5"),
                "quantity": 2,
                "previous_month_quantity": 3,
                "batch_no": "BATCH-002",
            }
        ], date(2026, 8, 31))

        formula = workbook[INVENTORY_IMPAIRMENT_SHEET]["M2"].value
        self.assertIn("L2-F2*(3-3", formula)
        self.assertIn("DATE(2026,7,31)-D2", formula)

        output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
        cached_sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(cached_sheet["K2"].value, 15.25)
        self.assertEqual(cached_sheet["L2"].value, 9.75)
        self.assertEqual(cached_sheet["M2"].value, -16.5)

    def test_previous_month_book_value_prefers_saved_snapshot(self):
        workbook, formula_cache = _build_inventory_impairment_workbook([
            {
                "store_name": "测试店铺",
                "sku": "SKU-002",
                "product_name": "历史库存商品",
                "arrived_at": date(2026, 7, 1),
                "unit_cost": Decimal("12.5"),
                "quantity": 2,
                "previous_month_quantity": 3,
                "previous_book_value": Decimal("7.5"),
                "impairment_amount": Decimal("0"),
                "batch_no": "BATCH-002",
            }
        ], date(2026, 8, 31))

        sheet = workbook[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(sheet["M2"].value, Decimal("17.5"))

        output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
        cached_sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(cached_sheet["M2"].value, 17.5)

    def test_fully_sold_current_batch_keeps_previous_month_reduction(self):
        workbook, formula_cache = _build_inventory_impairment_workbook([
            {
                "store_name": "测试店铺",
                "sku": "SKU-003",
                "product_name": "本月售完商品",
                "arrived_at": date(2026, 7, 1),
                "unit_cost": Decimal("10"),
                "quantity": 0,
                "previous_month_quantity": 5,
                "batch_no": "BATCH-003",
            }
        ], date(2026, 8, 31))

        output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
        cached_sheet = load_workbook(output, data_only=True)[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(cached_sheet["G2"].value, 0)
        self.assertEqual(cached_sheet["K2"].value, 0)
        self.assertEqual(cached_sheet["M2"].value, -35)
        self.assertEqual(cached_sheet["N2"].value, 0)
        self.assertEqual(cached_sheet["O2"].value, "已售完")

    def test_empty_report_has_numeric_totals_without_circular_formulas(self):
        workbook, _ = _build_inventory_impairment_workbook([], date(2026, 8, 31))
        sheet = workbook[INVENTORY_IMPAIRMENT_SHEET]
        self.assertEqual(sheet["A2"].value, "合计")
        for cell in ("E2", "F2", "G2", "H2", "I2", "J2", "K2", "L2", "M2", "N2"):
            self.assertEqual(sheet[cell].value, 0)

    def test_new_rule_allocates_safe_stock_to_newest_stable_batch(self):
        impairments = batch_impairments([
            InventoryLayer(1, 10, 1, date(2026, 8, 1), 4, PRODUCT_TYPE_STABLE, 7),
            InventoryLayer(2, 10, 2, date(2026, 9, 1), 6, PRODUCT_TYPE_STABLE, 7),
        ], date(2026, 9, 30))

        self.assertEqual(impairments[1].provision_quantity, 3)
        self.assertEqual(impairments[2].provision_quantity, 0)
        self.assertEqual(impairments[1].rate, Decimal("0.6"))

    def test_new_product_uses_half_percent_rate_from_arrival(self):
        self.assertEqual(
            impairment_rate(date(2026, 9, 30), date(2026, 9, 10), PRODUCT_TYPE_NEW),
            Decimal("0.1"),
        )
        self.assertEqual(
            impairment_rate(date(2026, 8, 31), date(2026, 8, 1), PRODUCT_TYPE_NEW),
            Decimal("0.15"),
        )

    def test_new_product_rate_starts_at_arrival_at_new_product_rate(self):
        self.assertEqual(
            impairment_rate(date(2026, 9, 29), date(2026, 8, 1), PRODUCT_TYPE_NEW),
            Decimal("0.295"),
        )

    def test_same_rule_comparison_keeps_decreasing_inventory_book_value_negative(self):
        current_layer = InventoryLayer(
            1, 10, 1, date(2026, 7, 8), 371, PRODUCT_TYPE_NEW, 0
        )
        previous_layer = InventoryLayer(
            1, 10, 1, date(2026, 7, 8), 394, PRODUCT_TYPE_NEW, 0
        )
        current_impairment = batch_impairments([current_layer], date(2026, 9, 29))[1]
        previous_impairment = batch_impairments([previous_layer], date(2026, 8, 31))[1]
        current_book_value = Decimal("5.689") * current_layer.quantity - (
            Decimal("5.689") * current_impairment.impairment_units
        )
        previous_book_value = Decimal("5.689") * previous_layer.quantity - (
            Decimal("5.689") * previous_impairment.impairment_units
        )

        self.assertLess(current_book_value - previous_book_value, 0)

    def test_report_restates_previous_book_value_without_overwriting_snapshot(self):
        batch = SimpleNamespace(
            id=1, product_id=10, quantity=394, unit_cost=Decimal("5.689"),
            arrived_at=datetime(2026, 7, 8), batch_no="BATCH-001",
        )
        snapshot = SimpleNamespace(
            batch_id=1, quantity=394, unit_cost=Decimal("5.689"),
            book_value=Decimal("1031.0744"),
        )
        rule = SimpleNamespace(
            product_id=10, product_type=PRODUCT_TYPE_NEW,
            safe_stock_quantity=0, effective_date=date(2026, 9, 22),
        )
        results = [
            [snapshot],
            [(batch, "CW000175", "测试商品", "测试店铺", PRODUCT_TYPE_NEW, 0, 1)],
            [rule],
            [SimpleNamespace(batch_id=1, report_quantity=23)],
        ]
        db = SimpleNamespace(execute=AsyncMock(side_effect=[
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: results[0])),
            SimpleNamespace(all=lambda: results[1]),
            SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: results[2])),
            SimpleNamespace(all=lambda: results[3]),
        ]))
        with patch("routers.reports.ensure_period_snapshot", new_callable=AsyncMock) as ensure:
            ensure.return_value = SimpleNamespace(id=1)
            rows = asyncio.run(
                _load_inventory_impairment_report_rows(db, date(2026, 9, 29))
            )

        self.assertEqual(rows[0]["previous_month_quantity"], 394)
        self.assertEqual(rows[0]["previous_book_value"], Decimal("1636.270180"))
        self.assertEqual(snapshot.book_value, Decimal("1031.0744"))
        current_book_value = rows[0]["unit_cost"] * rows[0]["quantity"] - rows[0]["impairment_amount"]
        self.assertLess(current_book_value - rows[0]["previous_book_value"], 0)

    def test_first_known_rule_backfills_from_earliest_arrival(self):
        impairments = batch_impairments([
            InventoryLayer(
                1, 10, 1, date(2026, 8, 1), 236, PRODUCT_TYPE_NEW, 0,
                (ImpairmentRule(PRODUCT_TYPE_NEW, 0, date(2026, 9, 22)),),
            ),
        ], date(2026, 9, 29))

        self.assertEqual(impairments[1].provision_quantity, 236)
        self.assertEqual(impairments[1].rate, Decimal("0.295"))

    def test_stable_product_applies_safe_stock_from_arrival(self):
        impairments = batch_impairments([
            InventoryLayer(1, 10, 1, date(2026, 8, 1), 200, PRODUCT_TYPE_STABLE, 100),
        ], date(2026, 9, 29))

        self.assertEqual(impairments[1].provision_quantity, 100)
        self.assertEqual(impairments[1].impairment_units, Decimal("59.000"))
        self.assertEqual(impairments[1].rate, Decimal("0.59"))

    def test_stable_product_newly_arrived_uses_excess_over_safe_stock(self):
        impairments = batch_impairments([
            InventoryLayer(1, 10, 1, date(2026, 9, 10), 200, PRODUCT_TYPE_STABLE, 100),
        ], date(2026, 9, 29))

        self.assertEqual(impairments[1].provision_quantity, 100)
        self.assertEqual(impairments[1].impairment_units, Decimal("19.000"))
        self.assertEqual(impairments[1].rate, Decimal("0.19"))

    def test_new_product_history_is_preserved_when_all_stock_becomes_safe(self):
        impairments = batch_impairments([
            InventoryLayer(
                1, 10, 1, date(2026, 9, 1), 200, PRODUCT_TYPE_STABLE, 200,
                (
                    ImpairmentRule(PRODUCT_TYPE_NEW, 0, date(2026, 9, 1)),
                    ImpairmentRule(PRODUCT_TYPE_STABLE, 200, date(2026, 9, 11)),
                ),
            ),
        ], date(2026, 9, 20))

        self.assertEqual(impairments[1].provision_quantity, 200)
        self.assertEqual(impairments[1].impairment_units, Decimal("9.000"))
        self.assertEqual(impairments[1].rate, Decimal("0.045"))

    def test_new_product_history_continues_for_stock_outside_safe_quantity(self):
        impairments = batch_impairments([
            InventoryLayer(
                1, 10, 1, date(2026, 9, 1), 200, PRODUCT_TYPE_STABLE, 100,
                (
                    ImpairmentRule(PRODUCT_TYPE_NEW, 0, date(2026, 9, 1)),
                    ImpairmentRule(PRODUCT_TYPE_STABLE, 100, date(2026, 9, 11)),
                ),
            ),
        ], date(2026, 9, 20))

        # 200 units keep the 9 days of new-product provision (4.5%), while
        # the 100 units beyond safe stock add 10 days at 1%.
        self.assertEqual(impairments[1].provision_quantity, 200)
        self.assertEqual(impairments[1].impairment_units, Decimal("19.000"))
        self.assertEqual(impairments[1].rate, Decimal("0.095"))

    def test_mixed_history_caps_safe_and_excess_stock_separately(self):
        impairments = batch_impairments([
            InventoryLayer(
                1, 10, 1, date(2026, 9, 1), 200, PRODUCT_TYPE_STABLE, 100,
                (
                    ImpairmentRule(PRODUCT_TYPE_NEW, 0, date(2026, 9, 1)),
                    ImpairmentRule(PRODUCT_TYPE_STABLE, 100, date(2027, 2, 8)),
                ),
            ),
        ], date(2027, 4, 18))

        # The safe 100 units remain at 79.5%; the other 100 units reach 90%.
        self.assertEqual(impairments[1].impairment_units, Decimal("169.500"))


if __name__ == "__main__":
    unittest.main()
