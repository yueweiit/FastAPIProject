"""独立SQLite覆盖追溯成本、费用来源切换、回款核对和版本保留；不证明InnoDB锁。"""
import copy
import os
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from fastapi import FastAPI, HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, func, inspect, select
from sqlalchemy.orm import Session

from auth import get_current_user
from database import Base, get_db
from models import (InventoryBatch, InventoryCostAdjustment, InventoryCostAdjustmentLine,
                    FinancialAuditEvent, FinancialExpense, FinancialExpenseSource, FinancialReportVersion, SettlementEntry,
                    Store, StoreProfitLossReport, User)
from routers import financial_control
from routers.products import _product_summary_stmt
from routers.reports import (_store_profit_loss_auto_values, _store_profit_loss_response,
                             _build_store_profit_loss_values, _build_store_profit_loss_sheet,
                             _build_store_profit_loss_summary_sheet, STORE_PROFIT_LOSS_MANUAL_KEYS,
                             STORE_PROFIT_LOSS_ROWS)
from routers.sales import monthly_sales_summary
from schemas import (CostAdjustmentRequest, CostAllocationRequest, ExpenseRequest, ExpenseSourceRequest,
                     FinanceReversalRequest, FinancialVersionRequest, ReceiptRequest,
                     SettlementClassificationRequest)
from services import financial_controls as controls
from services.finance_config import FINANCE_TABLES
from services.cost_layers import load_cost_layers, unit_adjustment, ensure_batch_cost_editable
from services.financial_facts import manual_period_values
from services.financial_versions import save_version, compare_version, version_workbook
from tests import test_financial_facts as facts


class WritableFacade(facts.AsyncFacade):
    def get_bind(self):
        return self.session.get_bind()

    def add(self, value):
        self.session.add(value)

    async def get(self, model, key):
        return self.session.get(model, key)

    async def flush(self):
        self.session.flush()

    async def commit(self):
        self.session.commit()

    async def rollback(self):
        self.session.rollback()

    async def refresh(self, value, **kwargs):
        self.session.refresh(value, **kwargs)


class FinancialControlsTests(facts.FinancialFactsTests):
    def setUp(self):
        super().setUp()
        self.flag = patch.dict(os.environ, {"FIFO_FINANCIAL_CONTROLS_ENABLED": "1"})
        self.flag.start()
        self.db = WritableFacade(self.session)
        self.user = SimpleNamespace(id=1, role="operator", store_id=7)
        self.admin = SimpleNamespace(id=1, role="admin", store_id=None)

    def tearDown(self):
        self.flag.stop()
        super().tearDown()

    def cost(self, **overrides):
        values = dict(request_key="cost-request-001", source_ref="运费账单1", category="head_shipping",
            amount="200", rate_date=date(2026, 10, 1), document_date=date(2026, 10, 1),
            allocations=[CostAllocationRequest(batch_id=1)], is_final=True)
        values.update(overrides)
        return CostAdjustmentRequest(**values)

    def expense(self, **overrides):
        values = dict(request_key="expense-request-001", store_id=7, expense_date=date(2026, 9, 10),
                      category="advertising", source_ref="广告账单1", amount="100")
        values.update(overrides)
        return ExpenseRequest(**values)

    def require_loaded_record(self, original):
        def checked(row):
            columns = {item.key for item in row.__mapper__.column_attrs}
            self.assertFalse(inspect(row).unloaded & columns,
                             f"{row.__tablename__}: audit/response must not implicitly load columns")
            return original(row)
        return checked

    async def test_financial_inserts_load_defaults_before_audit_without_returning(self):
        # SQLite normally returns defaults eagerly, hiding MySQL's expired-column path.
        self.engine.dialect.insert_returning = False
        self.engine.dialect.update_returning = False
        with patch.object(controls, "record", side_effect=self.require_loaded_record(controls.record)):
            cost = await controls.post_cost_adjustment(self.db, self.user, self.cost())
            expense = await controls.create_expense(self.db, self.user, self.expense())
            source = await controls.set_expense_source(self.db, self.user, ExpenseSourceRequest(
                store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="ledger"))
            receipt = await controls.create_receipt(self.db, self.user, ReceiptRequest(
                request_key="defaults-receipt-001", store_id=7, payout_id="defaults-payout",
                receipt_date=date(2026, 9, 20), source_ref="defaults-receipt", amount="1"))
            classification = await controls.classify_settlement(self.db, self.user, 1,
                SettlementClassificationRequest(breakdown={"tax": Decimal(10)}))
            version = await save_version(self.db, self.admin, FinancialVersionRequest(
                request_key="defaults-version-001", store_id=7, report_month=date(2026, 9, 1)))
            await self.db.commit()
            for row, field in ((cost, "created_at"), (expense, "created_at"), (source, "updated_at"),
                               (receipt, "created_at"), (classification, "confirmed_at"), (version, "created_at")):
                values = controls.record(row)
                self.assertIsInstance(values[field], datetime)
                audit = self.session.scalar(select(FinancialAuditEvent).where(
                    FinancialAuditEvent.entity_type == row.__tablename__, FinancialAuditEvent.entity_id == row.id))
                self.assertEqual(audit.after_values[field], str(values[field]))

    async def test_source_update_audit_uses_persisted_onupdate_timestamp(self):
        self.engine.dialect.insert_returning = False
        self.engine.dialect.update_returning = False
        source = await controls.set_expense_source(self.db, self.user, ExpenseSourceRequest(
            store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="ledger"))
        source.updated_at = datetime(2000, 1, 1)
        await self.db.commit()
        with patch.object(controls, "record", side_effect=self.require_loaded_record(controls.record)):
            updated = await controls.set_expense_source(self.db, self.user, ExpenseSourceRequest(
                store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="manual"))
            await self.db.commit()
            response = controls.record(updated)
        audit = self.session.scalar(select(FinancialAuditEvent).where(
            FinancialAuditEvent.entity_type == source.__tablename__).order_by(FinancialAuditEvent.id.desc()))
        self.assertEqual(audit.before_values["updated_at"], "2000-01-01 00:00:00")
        self.assertEqual(audit.after_values["updated_at"], str(response["updated_at"]))
        self.assertNotEqual(response["updated_at"], datetime(2000, 1, 1))
        self.assertEqual((audit.before_values["mode"], audit.after_values["mode"]), ("ledger", "manual"))

    async def test_audit_failure_does_not_commit_flushed_voucher(self):
        self.engine.dialect.insert_returning = False
        with patch.object(controls, "FinancialAuditEvent", side_effect=RuntimeError("audit failure")):
            with self.assertRaisesRegex(RuntimeError, "audit failure"):
                await controls.create_expense(self.db, self.user, self.expense())
        await self.db.rollback()
        self.assertEqual(self.session.scalar(select(func.count()).select_from(FinancialExpense)), 0)
        self.assertEqual(self.session.scalar(select(func.count()).select_from(FinancialAuditEvent)), 0)

    async def test_late_cost_restates_original_sale_month_without_mutating_fifo(self):
        batch = self.session.get(InventoryBatch, 1)
        from models import Sale, SaleCostDetail
        batch.quantity, batch.remaining_quantity = 100, 60
        sale = self.session.get(Sale, 1)
        sale.quantity, sale.total_cost = 40, Decimal(400)
        detail = self.session.scalar(select(SaleCostDetail).where(SaleCostDetail.sale_id == 1))
        detail.quantity = 40
        self.session.commit()
        before = await monthly_sales_summary(None, None, "2026-09", self.db, self.user)
        _batches, preview, amount = await controls.prepare_cost_adjustment(self.db, self.user, self.cost())
        self.assertEqual(amount, Decimal(200))
        self.assertEqual(preview[0]["sold_amount_cny"], Decimal(80))
        self.assertEqual(preview[0]["remaining_amount_cny"], Decimal(120))
        row = await controls.post_cost_adjustment(self.db, self.user, self.cost())
        await self.db.commit()
        after = await monthly_sales_summary(None, None, "2026-09", self.db, self.user)
        self.assertEqual(after[0].sales_cost_cny - before[0].sales_cost_cny, Decimal(80))
        self.assertEqual(after[0].gross_profit_cny - before[0].gross_profit_cny, Decimal(-80))
        self.assertEqual(batch.unit_cost, Decimal(10))
        self.assertEqual(sale.total_cost, Decimal(400))
        self.assertEqual(detail.unit_cost, Decimal(10))
        self.assertEqual(batch.remaining_quantity, 60)
        product = (await self.db.execute(_product_summary_stmt())).one()
        self.assertEqual(product.stock_value, Decimal(720))
        duplicate = await controls.post_cost_adjustment(self.db, self.user, self.cost())
        self.assertEqual(duplicate.id, row.id)

    async def test_negative_cost_retry_and_reversal_are_idempotent(self):
        row = await controls.post_cost_adjustment(self.db, self.user,
            self.cost(category="purchase", amount="-50", source_ref="采购折让"))
        await self.db.commit()
        again = await controls.post_cost_adjustment(self.db, self.user,
            self.cost(category="purchase", amount="-50", source_ref="采购折让"))
        self.assertEqual(again.id, row.id)
        request = FinanceReversalRequest(request_key="reverse-cost-001", note="折让单撤回")
        reverse = await controls.reverse_cost_adjustment(self.db, self.user, row.id, request)
        await self.db.commit()
        repeat = await controls.reverse_cost_adjustment(self.db, self.user, row.id, request)
        self.assertEqual(reverse.id, repeat.id)
        layers = await load_cost_layers(self.db, {1})
        self.assertEqual(unit_adjustment(layers[1]), Decimal(0))
        with self.assertRaises(HTTPException):
            await ensure_batch_cost_editable(self.db, 1)

    async def test_adjustments_validate_owner_duplicate_batches_and_nonnegative_component(self):
        for payload, user in [
            (self.cost(allocations=[CostAllocationRequest(batch_id=1), CostAllocationRequest(batch_id=1)]), self.user),
            (self.cost(), SimpleNamespace(id=99, role="operator", store_id=7)),
            (self.cost(category="purchase", amount=-101), self.user),
        ]:
            with self.assertRaises(HTTPException):
                await controls.prepare_cost_adjustment(self.db, user, payload)
        self.assertEqual(self.session.scalar(select(func.count()).select_from(InventoryCostAdjustment)), 0)

    async def test_allocation_rounding_preserves_document_total(self):
        for key in (2, 3):
            self.session.add(InventoryBatch(id=key, product_id=10, batch_no=f"B{key}", quantity=10,
                remaining_quantity=10, unit_cost=10, purchase_price=10, user_id=1, arrived_at=datetime(2026, 9, 1)))
        self.session.commit()
        payload = self.cost(amount="0.0001", allocations=[CostAllocationRequest(batch_id=i) for i in (1, 2, 3)])
        _, rows, amount = await controls.prepare_cost_adjustment(self.db, self.user, payload)
        self.assertEqual(sum(row["amount_cny"] for row in rows), amount)
        self.assertEqual(rows[-1]["amount_cny"], Decimal("0.0001"))

    async def test_small_and_weighted_allocations_keep_sign_and_exact_total(self):
        for key in (2, 3, 4):
            self.session.add(InventoryBatch(id=key, product_id=10, batch_no=f"SPLIT-{key}", quantity=10,
                remaining_quantity=10, unit_cost=10, purchase_price=10, user_id=1, arrived_at=datetime(2026, 9, 1)))
        await self.db.commit()
        for amount, weights in (("0.0002", None), ("-0.0002", None), ("0", None),
                                ("200", ["1", "1", "1", "0.00000001"]),
                                ("-200", ["1", "1", "1", "0.00000001"])):
            with self.subTest(amount=amount, weights=weights):
                allocations = [CostAllocationRequest(batch_id=i, weight=weights[i-1] if weights else None)
                               for i in (1, 2, 3, 4)]
                payload = self.cost(amount=amount, allocations=allocations,
                    category="purchase" if Decimal(amount) < 0 else "head_shipping",
                    allocation_basis="weight" if weights else "quantity")
                _, rows, total = await controls.prepare_cost_adjustment(self.db, self.user, payload)
                self.assertEqual(sum(row["amount_cny"] for row in rows), total)
                self.assertTrue(all(row["amount_cny"] >= 0 if total >= 0 else row["amount_cny"] <= 0 for row in rows))
                denominator = sum(row["allocation_weight"] for row in rows)
                for row in rows:
                    ideal = total * row["allocation_weight"] / denominator
                    self.assertLess(abs(row["amount_cny"] - ideal), controls.MONEY)
                payload.allocations.reverse()
                _, again, _ = await controls.prepare_cost_adjustment(self.db, self.user, payload)
                self.assertEqual(rows, again)
        saved = await controls.post_cost_adjustment(self.db, self.user, self.cost(amount="0.0002", allocations=allocations,
            allocation_basis="weight"))
        await self.db.commit()
        self.assertEqual(saved.amount_cny, Decimal("0.0002"))

    async def test_legacy_cost_reversal_validates_exact_saved_allocations(self):
        for key in (2, 3, 4):
            self.session.add(InventoryBatch(id=key, product_id=10, batch_no=f"LEGACY-{key}", quantity=10,
                remaining_quantity=10, unit_cost=10, purchase_price=10, user_id=1,
                shipping_cost=Decimal("0.0001") if key == 4 else 0, arrived_at=datetime(2026, 9, 1)))
        self.session.flush()
        payload = self.cost(amount="0.0002")
        header = InventoryCostAdjustment(**payload.model_dump(exclude={"allocations"}),
            amount_cny=Decimal("0.0002"), created_by=1)
        self.session.add(header)
        self.session.flush()
        amounts = [Decimal("0.0001")] * 3 + [Decimal("-0.0001")]
        for batch_id, amount in enumerate(amounts, 1):
            self.session.add(InventoryCostAdjustmentLine(adjustment_id=header.id, batch_id=batch_id,
                quantity=10, allocation_weight=10, amount_cny=amount))
        await self.db.commit()
        reverse = await controls.reverse_cost_adjustment(self.db, self.user, header.id,
            FinanceReversalRequest(request_key="legacy-cost-reverse", note="更正历史分摊"))
        await self.db.commit()
        lines = (await self.db.execute(select(InventoryCostAdjustmentLine).where(
            InventoryCostAdjustmentLine.adjustment_id == reverse.id).order_by(InventoryCostAdjustmentLine.batch_id))).scalars().all()
        self.assertEqual([line.amount_cny for line in lines], [-amount for amount in amounts])
        layers = await load_cost_layers(self.db, {1, 2, 3, 4})
        self.assertTrue(all(unit_adjustment(group) == 0 for group in layers.values()))

    async def test_rounding_residual_cannot_create_negative_total_cost(self):
        batch = self.session.get(InventoryBatch, 1)
        batch.purchase_price = batch.unit_cost = Decimal("0.3333")
        batch.shipping_cost = Decimal("0.0001")
        await self.db.commit()
        await controls.post_cost_adjustment(self.db, self.user, self.cost(
            category="purchase", amount="-3.3330", source_ref="采购退款", request_key="rounding-purchase"))
        await self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            await controls.post_cost_adjustment(self.db, self.user, self.cost(
                category="head_shipping", amount="-0.0001", source_ref="运费退款", request_key="rounding-shipping"))
        self.assertEqual(raised.exception.detail["code"], "negative_total_cost")

    async def test_negative_minimum_cost_is_rejected_even_for_large_quantity(self):
        batch = self.session.get(InventoryBatch, 1)
        batch.quantity, batch.remaining_quantity = 2_147_483_647, 2_147_483_643
        batch.purchase_price = batch.unit_cost = Decimal(0)
        await self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            await controls.post_cost_adjustment(self.db, self.user, self.cost(
                category="purchase", amount="-0.0001", source_ref="最小金额负成本"))
        self.assertEqual(raised.exception.detail["code"], "negative_component_cost")
        self.assertEqual(self.session.scalars(select(InventoryCostAdjustment)).all(), [])

    async def test_expense_source_is_exclusive_and_drafts_mask_profit_until_posted(self):
        report = StoreProfitLossReport(store_id=7, report_month=date(2026, 9, 1),
            manual_values={"advertising": {"current": "500"}}, report_metadata={"manual_period_mode": "monthly"})
        self.session.add(report)
        await controls.create_expense(self.db, self.user, self.expense())
        await self.db.commit()
        manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
        self.assertEqual(manual["advertising"]["current"], "500")
        source = ExpenseSourceRequest(store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="ledger")
        await controls.set_expense_source(self.db, self.user, source)
        await self.db.commit()
        manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
        values = _build_store_profit_loss_values(await self.auto_values(), manual)
        self.assertIsNone(values["advertising"]["current"])
        self.assertIsNone(values["net_profit"]["current"])
        row = self.session.scalar(select(FinancialExpense))
        with self.assertRaises(HTTPException):
            await controls.post_expense(self.db, self.user, row.id)
        await controls.post_expense(self.db, self.admin, row.id)
        await self.db.commit()
        manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
        self.assertEqual(Decimal(manual["advertising"]["current"]), Decimal(100))
        self.assertEqual(manual["advertising"]["ytd"], Decimal(100))
        self.assertEqual(report.manual_values["advertising"]["current"], "500")
        reverse = await controls.reverse_expense(self.db, self.admin, row.id,
            FinanceReversalRequest(request_key="expense-reverse-001", note="账单退回"))
        await self.db.commit()
        manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
        self.assertEqual(Decimal(manual["advertising"]["current"]), Decimal(0))
        self.assertEqual(reverse.amount_cny, Decimal(-100))

    async def test_platform_linked_voucher_is_not_deducted_twice(self):
        row = await controls.create_expense(self.db, self.user, self.expense(settlement_entry_id=1))
        await controls.post_expense(self.db, self.admin, row.id)
        await controls.set_expense_source(self.db, self.user,
            ExpenseSourceRequest(store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="ledger"))
        await self.db.commit()
        manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
        self.assertEqual(Decimal(manual["advertising"]["current"]), Decimal(0))
        values = _build_store_profit_loss_values(await self.auto_values(), manual)
        self.assertEqual(values["net_profit"]["current"], Decimal(27))

    async def test_cross_store_excel_keeps_unconfirmed_ledger_blank_until_posted(self):
        expense = await controls.create_expense(self.db, self.user, self.expense())
        await controls.set_expense_source(self.db, self.user, ExpenseSourceRequest(
            store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="ledger"))
        await self.db.commit()
        auto = await self.auto_values()
        other = _build_store_profit_loss_values(auto, {"advertising": {"current": "25"}})
        for confirmed in (False, True):
            with self.subTest(confirmed=confirmed):
                if confirmed:
                    await controls.post_expense(self.db, self.admin, expense.id)
                    await self.db.commit()
                manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
                values = _build_store_profit_loss_values(auto, manual)
                book = Workbook()
                _build_store_profit_loss_sheet(book, self.session.get(Store, 7), date(2026, 9, 1), values, {})
                _build_store_profit_loss_summary_sheet(book,
                    [(self.session.get(Store, 7), values), (self.session.get(Store, 8), other)], date(2026, 9, 1))
                data = BytesIO()
                book.save(data)
                data.seek(0)
                exported = load_workbook(data, data_only=True)
                advertising_row = next(i+2 for i, row in enumerate(STORE_PROFIT_LOSS_ROWS) if row[0] == "advertising")
                single, summary = exported.worksheets[-2:]
                self.assertEqual(single.cell(advertising_row, 2).value, 100 if confirmed else None)
                self.assertEqual(summary.cell(advertising_row, 2).value, 100 if confirmed else None)
                self.assertEqual(summary.cell(advertising_row, 3).value, 25)
                self.assertEqual(summary.cell(advertising_row, 4).value, 125 if confirmed else None)
                self.assertEqual(summary.cell(advertising_row, 5).value is None, not confirmed)
                empty_manual_row = next(i+2 for i, row in enumerate(STORE_PROFIT_LOSS_ROWS) if row[0] == "warehousing")
                self.assertEqual(summary.cell(empty_manual_row, 4).value, 0)

    async def test_saving_profit_loss_keeps_manual_fallback_for_ledger_category(self):
        from routers.reports import update_store_profit_loss
        from schemas import StoreProfitLossUpdateRequest
        report = StoreProfitLossReport(store_id=7, report_month=date(2026, 9, 1),
            manual_values={"advertising": {"current": "500"}, "salary": {"current": "20", "ytd": "999"}})
        self.session.add(report)
        await controls.set_expense_source(self.db, self.user, ExpenseSourceRequest(
            store_id=7, report_month=date(2026, 9, 1), category="advertising", mode="ledger"))
        await self.db.commit()
        manual = await manual_period_values(self.db, 7, date(2026, 9, 1), STORE_PROFIT_LOSS_MANUAL_KEYS)
        self.assertEqual(Decimal(manual["salary"]["ytd"]), Decimal(999))
        with patch("routers.reports._store_profit_loss_auto_values", new=AsyncMock(return_value=await self.auto_values())):
            result = await update_store_profit_loss(StoreProfitLossUpdateRequest(
                store_id=7, report_month=date(2026, 9, 1), update_scope="current"), db=self.db, user=self.user)
        self.assertEqual(report.manual_values["advertising"]["current"], "500")
        self.assertEqual(next(row for row in result.rows if row.key == "advertising").input_source, "ledger")

    async def test_locked_cost_validation_uses_current_read_instead_of_cached_layers(self):
        from sqlalchemy.dialects.mysql import dialect
        await load_cost_layers(self.db, {1})
        captured = []
        original_execute = self.db.execute
        async def capture(statement):
            captured.append(str(statement.compile(dialect=dialect())))
            return await original_execute(statement)
        with patch.object(self.db, "execute", side_effect=capture):
            await controls.prepare_cost_adjustment(self.db, self.user, self.cost(), lock=True)
        self.assertTrue(any("inventory_batches" in sql and "FOR UPDATE" in sql for sql in captured))
        self.assertTrue(any("inventory_cost_adjustment_lines" in sql and "FOR UPDATE" in sql for sql in captured))

    async def test_classification_balances_existing_difference_without_changing_profit(self):
        before = await self.auto_values()
        with self.assertRaises(HTTPException):
            await controls.classify_settlement(self.db, self.user, 1,
                SettlementClassificationRequest(breakdown={"commission": Decimal(11)}))
        await controls.classify_settlement(self.db, self.user, 1,
            SettlementClassificationRequest(breakdown={"commission": Decimal(7), "tax": Decimal(3)}))
        await self.db.commit()
        result = await controls.reconciliation(self.db, 7, date(2026, 9, 1))
        row = next(row for row in result["entries"] if row["id"] == 1)
        self.assertEqual(row["unclassified_difference"], Decimal(0))
        after = await self.auto_values()
        self.assertEqual(before["other_direct_cost"], after["other_direct_cost"])

    async def test_receipts_match_all_months_by_payout_without_currency_mixing(self):
        self.session.get(SettlementEntry, 1).payout_id = "PAY-1"
        self.session.commit()
        def request(key, amount, currency="CNY"):
            return ReceiptRequest(request_key=key, store_id=7, payout_id="PAY-1", receipt_date=date(2026, 10, 1),
                                  source_ref=key, amount=amount, currency=currency, exchange_rate=1)
        row = await controls.create_receipt(self.db, self.user, request("receipt-001", 40))
        await self.db.commit()
        result = await controls.reconciliation(self.db, 7, date(2026, 9, 1))
        self.assertEqual(result["payouts"][0]["difference"], Decimal(50))
        await controls.create_receipt(self.db, self.user, request("receipt-002", 50))
        await controls.create_receipt(self.db, self.user, request("receipt-003", 10, "USD"))
        await self.db.commit()
        result = await controls.reconciliation(self.db, 7, date(2026, 9, 1))
        groups = {r["currency"]: r for r in result["payouts"]}
        self.assertEqual(groups["CNY"]["difference"], Decimal(0))
        self.assertEqual(groups["USD"]["difference"], Decimal(-10))
        reversal = await controls.reverse_receipt(self.db, self.user, row.id,
            FinanceReversalRequest(request_key="receipt-reverse-001", note="更正回款"))
        await self.db.commit()
        self.assertEqual(reversal.amount, Decimal(-40))

    async def test_old_version_and_export_remain_unchanged_after_restatement(self):
        patchers = [patch("routers.reports._store_provision_financial_values", new=AsyncMock(return_value=self.zero_provisions())),
                    patch("routers.reports.office_space_totals_by_application_date", new=AsyncMock(return_value={})),
                    patch("routers.reports.shared_admin_totals_by_application_date", new=AsyncMock(return_value={})),
                    patch("routers.reports.china_salary_totals_by_store_and_application_date", new=AsyncMock(return_value={}))]
        for patcher in patchers:
            patcher.start()
            self.addCleanup(patcher.stop)
        original = await save_version(self.db, self.admin, FinancialVersionRequest(
            request_key="version-001", store_id=7, report_month=date(2026, 9, 1)))
        await self.db.commit()
        saved = copy.deepcopy(original.payload)
        await controls.post_cost_adjustment(self.db, self.user, self.cost())
        await self.db.commit()
        comparison = await compare_version(self.db, self.admin, original)
        self.assertTrue(comparison["changed"])
        cost = next(row for row in comparison["changes"] if row["key"] == "head_logistics_cost" and row["period"] == "current")
        self.assertEqual(cost["difference"], Decimal(80))
        second = await save_version(self.db, self.admin, FinancialVersionRequest(
            request_key="version-002", store_id=7, report_month=date(2026, 9, 1)))
        await self.db.commit()
        self.assertEqual((original.version, second.version), (1, 2))
        self.assertEqual(original.payload, saved)
        old_excel = load_workbook(version_workbook(original), data_only=True)
        new_excel = load_workbook(version_workbook(second), data_only=True)
        self.assertEqual(old_excel.worksheets[0]["F2"].value, 10)
        self.assertEqual(new_excel.worksheets[0]["F2"].value, 30)
        duplicate = await save_version(self.db, self.admin, FinancialVersionRequest(
            request_key="version-003", store_id=7, report_month=date(2026, 9, 1)))
        self.assertEqual(duplicate.id, second.id)
        adjustment = self.session.scalar(select(InventoryCostAdjustment))
        await controls.reverse_cost_adjustment(self.db, self.admin, adjustment.id,
            FinanceReversalRequest(request_key="cost-reversal-new-version", note="撤销补差"))
        await self.db.commit()
        third = await save_version(self.db, self.admin, FinancialVersionRequest(
            request_key="version-after-reversal", store_id=7, report_month=date(2026, 9, 1)))
        self.assertEqual(third.version, 3)
        comparison = await compare_version(self.db, self.admin, original)
        self.assertEqual(comparison["changes"], [])
        self.assertTrue(comparison["changed"])
        self.assertEqual(len(third.payload["cost_manifest"]), 2)

    async def test_legacy_settlement_owner_fallback_is_visible_and_warned(self):
        entry = self.session.get(SettlementEntry, 1)
        entry.store_id = None
        await self.db.commit()
        result = await controls.reconciliation(self.db, 7, date(2026, 9, 1))
        self.assertTrue(any(row["id"] == 1 and not row["can_classify"] for row in result["entries"]))
        self.assertTrue(any("推导店铺" in warning for warning in result["warnings"]))

    async def test_financial_references_block_physical_import_rollback_before_inventory_changes(self):
        from routers.sales import rollback_sales_import
        await controls.classify_settlement(self.db, self.user, 1,
            SettlementClassificationRequest(breakdown={"tax": Decimal(10)}))
        await self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            await rollback_sales_import(1, db=self.db, user=self.user)
        self.assertEqual(raised.exception.detail["code"], "import_has_financial_references")
        self.assertEqual(self.session.get(InventoryBatch, 1).remaining_quantity, 6)

    async def test_rolled_back_import_cannot_acquire_new_financial_references(self):
        from models import SalesImportBatch
        self.session.get(SalesImportBatch, 1).status = "rolled_back"
        await self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            await controls.create_expense(self.db, self.user, self.expense(settlement_entry_id=1))
        self.assertEqual(raised.exception.detail["code"], "settlement_source_rolled_back")

    async def test_disabling_finance_after_posting_cannot_silently_revert_amounts(self):
        from contextlib import asynccontextmanager
        from database import init_db
        await controls.post_cost_adjustment(self.db, self.user, self.cost())
        await self.db.commit()
        class Connection:
            async def run_sync(inner, operation):
                with self.engine.connect() as connection:
                    return operation(connection)
        @asynccontextmanager
        async def begin():
            yield Connection()
        with patch("database.engine", SimpleNamespace(begin=begin)), patch.dict(
            os.environ, {"FIFO_FINANCIAL_CONTROLS_ENABLED": "0"}):
            with self.assertRaisesRegex(RuntimeError, "不能关闭财务开关"):
                await init_db()

    async def test_version_capture_rejects_mysql_read_committed(self):
        connection = SimpleNamespace(get_isolation_level=AsyncMock(return_value="READ COMMITTED"))
        with patch.object(self.db, "get_bind", return_value=SimpleNamespace(dialect=SimpleNamespace(name="mysql"))), patch.object(
            self.db, "connection", new=AsyncMock(return_value=connection), create=True):
            with self.assertRaises(HTTPException) as raised:
                await save_version(self.db, self.admin, FinancialVersionRequest(
                    request_key="snapshot-rejected", store_id=7, report_month=date(2026, 9, 1)))
        self.assertEqual(raised.exception.detail["code"], "snapshot_isolation_required")

    async def test_version_routes_read_oa_after_releasing_mysql_transaction(self):
        app = FastAPI()
        app.include_router(financial_control.router)
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.admin
        oa_reads = []

        async def read_oa(*_args):
            self.assertFalse(self.session.in_transaction())
            oa_reads.append(True)
            return {}

        body = FinancialVersionRequest(request_key="version-prefetch-001", store_id=7,
                                       report_month=date(2026, 9, 1)).model_dump(mode="json")
        with patch("routers.reports._store_provision_financial_values",
                   new=AsyncMock(return_value=self.zero_provisions())), \
             patch("routers.reports.office_space_totals_by_application_date", new=read_oa), \
             patch("routers.reports.shared_admin_totals_by_application_date", new=read_oa), \
             patch("routers.reports.china_salary_totals_by_store_and_application_date", new=read_oa), \
             patch("services.financial_versions.today", return_value=date(2026, 10, 27)), \
             patch.object(controls, "today", return_value=date(2026, 10, 27)):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                overview = await client.get("/financial-control/overview?store_id=7&report_month=2026-09-01")
                self.assertEqual(overview.status_code, 200, overview.text)
                captured = await client.post("/financial-control/versions", json=body)
                self.assertEqual(captured.status_code, 200, captured.text)
                version_id = captured.json()["id"]
                compared = await client.get(f"/financial-control/versions/{version_id}/compare")
                self.assertEqual(compared.status_code, 200, compared.text)
                self.assertEqual(len(oa_reads), 9)  # Three OA sources, once per request.

                async def unavailable(*_args):
                    raise AssertionError("Idempotent replay must not read OA")

                with patch("routers.reports.office_space_totals_by_application_date", new=unavailable):
                    replay = await client.post("/financial-control/versions", json=body)
                    conflict = await client.post("/financial-control/versions", json={**body, "note": "changed"})
                self.assertEqual(replay.status_code, 200, replay.text)
                self.assertEqual(replay.json()["id"], version_id)
                self.assertEqual(conflict.status_code, 409, conflict.text)
                self.assertEqual(conflict.json()["detail"]["code"], "idempotency_conflict")

            app.dependency_overrides[get_current_user] = lambda: self.user
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                forbidden = await client.get("/financial-control/overview?store_id=8&report_month=2026-09-01")
            self.assertEqual(forbidden.status_code, 403)
            self.assertEqual(len(oa_reads), 9)

    async def test_api_authorization_feature_gate_and_duplicate_source_errors(self):
        app = FastAPI()
        app.include_router(financial_control.router)
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.user
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get("/financial-control/expenses?store_id=8&report_month=2026-09-01")
            self.assertEqual(response.status_code, 403)
            body = self.expense().model_dump(mode="json")
            response = await client.post("/financial-control/expenses", json=body)
            self.assertEqual(response.status_code, 200)
            self.assertIsInstance(response.json()["amount_cny"], str)
            self.assertEqual(Decimal(response.json()["amount_cny"]), Decimal(100))
            row_id = response.json()["id"]
            response = await client.post(f"/financial-control/expenses/{row_id}/post")
            self.assertEqual(response.status_code, 403)
            body["request_key"] = "expense-new-request"
            response = await client.post("/financial-control/expenses", json=body)
            self.assertEqual(response.status_code, 409)
            with patch.dict(os.environ, {"FIFO_FINANCIAL_CONTROLS_ENABLED": "0"}):
                response = await client.get("/financial-control/expenses?store_id=7&report_month=2026-09-01")
                self.assertEqual(response.status_code, 503)
            with self.assertRaises(ValueError):
                self.expense(currency="USD")

    async def test_disabled_cost_endpoints_do_not_query_unmigrated_finance_tables(self):
        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine, tables=[table for table in Base.metadata.sorted_tables if table.name not in FINANCE_TABLES])
        self.assertTrue(set(inspect(engine).get_table_names()).isdisjoint(FINANCE_TABLES))
        statements = []
        event.listen(engine, "before_cursor_execute", lambda _conn, _cursor, statement, *_args: statements.append(statement))
        session = Session(engine)
        app = FastAPI()
        app.include_router(financial_control.router)
        app.dependency_overrides[get_db] = lambda: WritableFacade(session)
        app.dependency_overrides[get_current_user] = lambda: self.user
        try:
            with patch.dict(os.environ, {"FIFO_FINANCIAL_CONTROLS_ENABLED": "0"}):
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                    for suffix in ("/preview", ""):
                        response = await client.post("/financial-control/cost-adjustments" + suffix, json=self.cost().model_dump(mode="json"))
                        self.assertEqual(response.status_code, 503)
                        self.assertEqual(response.json()["detail"]["code"], "financial_controls_disabled")
            self.assertEqual(statements, [])
        finally:
            session.close()
            engine.dispose()
