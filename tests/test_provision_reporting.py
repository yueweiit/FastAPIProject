"""增提/转销的API、Excel、店铺权限、损益及产品线联动回归。"""

import unittest
from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook, load_workbook

from routers.reports import (
    _build_inventory_impairment_workbook, _build_product_line_profit_sheet,
    _build_store_profit_loss_summary_sheet, _build_store_profit_loss_values,
    _load_inventory_impairment_report_rows, _month_periods, _product_line_profit_data,
    _sale_provision_adjustments, _store_provision_financial_values,
    _store_profit_loss_auto_values,
    _summarize_inventory_impairment_rows, _workbook_bytes_with_formula_cache,
    inventory_impairment_report,
)

REPORT = date(2026, 9, 30)


def result(rows):
    return SimpleNamespace(all=lambda: rows, scalars=lambda: SimpleNamespace(all=lambda: rows))


def fixture(*, owner=7, sale_store=7, invalid=False, second_owner=None):
    batch = SimpleNamespace(id=1, product_id=10, quantity=100, unit_cost=Decimal(10),
                            arrived_at=datetime(2026, 8, 31), batch_no="A")
    batches = [(batch, "SKU-1", "测试物料", "店铺A", "new", 0, owner)]
    rules = [SimpleNamespace(product_id=10, product_type="new", safe_stock_quantity=0,
                             effective_date=date(2026, 8, 31))]
    sales = [SimpleNamespace(detail_id=1, sale_id=1, batch_id=1, report_quantity=10,
                             saved_unit_cost=Decimal(10), sold_at=datetime(2026, 8, 30) if invalid else datetime(2026, 9, 10),
                             sale_store_id=sale_store),
             SimpleNamespace(detail_id=2, sale_id=2, batch_id=1, report_quantity=30,
                             saved_unit_cost=Decimal(10), sold_at=datetime(2026, 9, 20), sale_store_id=sale_store)]
    if second_owner is not None:
        batch.arrived_at = datetime(2026, 8, 1)
        batches[0] = (batch, "SKU-1", "测试物料", "店铺A", "stable", 100, owner)
        batches.append((SimpleNamespace(id=2, product_id=10, quantity=100, unit_cost=Decimal(12),
                                         arrived_at=datetime(2026, 9, 20), batch_no="B"),
                        "SKU-1", "测试物料", "他店名称不能泄漏", "stable", 100, second_owner))
        rules = [SimpleNamespace(product_id=10, product_type="stable", safe_stock_quantity=100,
                                 effective_date=date(2026, 8, 1))]
        sales = []
    return SimpleNamespace(execute=AsyncMock(side_effect=[result(batches), result(rules), result(sales)]),
                           commit=AsyncMock(), info={})


class ProvisionReportingTests(unittest.IsolatedAsyncioTestCase):
    async def test_accepted_example_detail_summary_and_cached_excel(self):
        db = fixture()
        rows = await _load_inventory_impairment_report_rows(db, REPORT)
        with patch("routers.reports._load_inventory_impairment_report_rows", new_callable=AsyncMock, return_value=rows):
            detail = (await inventory_impairment_report(REPORT, db, None))[0]
        for key, value in {"opening_impairment_amount": 0, "monthly_total_impairment_amount": 125,
                           "monthly_writeoff_amount": 35, "impairment_amount": 90, "book_value": 510}.items():
            self.assertEqual(detail[key], Decimal(value))
            self.assertEqual(_summarize_inventory_impairment_rows(rows)[0][key], Decimal(value))
        workbook, cache = _build_inventory_impairment_workbook(rows, REPORT)
        sheet = load_workbook(_workbook_bytes_with_formula_cache(workbook, cache), data_only=True).active
        self.assertEqual(sheet.max_column, 23)
        self.assertEqual([sheet[key].value for key in ("U2", "V2", "W2", "O2", "P2")], [0,125,35,90,510])
        self.assertEqual([sheet[key].value for key in ("V3", "W3", "P3")], [125,35,510])
        db.commit.assert_not_awaited()

    async def test_invalid_history_remains_null_in_api_summary_formulas_and_totals(self):
        rows = await _load_inventory_impairment_report_rows(fixture(invalid=True), REPORT)
        self.assertEqual(rows[0]["original_inventory_amount"], Decimal(1000))
        self.assertEqual(rows[0]["sold_cost_amount"], Decimal(400))
        self.assertIn("销售日期早于", rows[0]["comparison_note"])
        self.assertIsNone(rows[0]["impairment_amount"])
        self.assertIsNone(_summarize_inventory_impairment_rows(rows)[0]["book_value"])
        workbook, cache = _build_inventory_impairment_workbook(rows, REPORT)
        sheet = load_workbook(_workbook_bytes_with_formula_cache(workbook, cache), data_only=True).active
        for cell in ("O2", "P2", "V2", "W2", "O3", "P3", "V3", "W3"):
            self.assertIsNone(sheet[cell].value, cell)
        self.assertEqual(sheet["H2"].value, 1000)
        self.assertEqual(sheet["S2"].value, "历史流水待核对")

    async def test_store_filter_follows_global_safe_stock_allocation_without_name_leak(self):
        rows = await _load_inventory_impairment_report_rows(fixture(second_owner=8), REPORT, 7)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["store_id"], 7)
        self.assertEqual(rows[0]["impairment_amount"], Decimal(110))
        self.assertNotIn("他店名称", str(rows))

    async def test_financial_accrual_and_cost_writeoff_do_not_double_charge_profit(self):
        values = await _store_provision_financial_values(fixture(), 7, _month_periods(date(2026,9,1)))
        self.assertEqual(values["inventory_impairment"]["current"], Decimal(125))
        self.assertEqual(values["inventory_writeoff"]["current"], Decimal(35))
        values.update({"product_sales_revenue": {"current": Decimal(600)},
                       "product_purchase_cost": {"current": Decimal(400)}})
        report = _build_store_profit_loss_values(values, {})
        self.assertEqual(report["operating_cost"]["current"], Decimal(365))
        self.assertEqual(report["gross_profit"]["current"], Decimal(235))
        self.assertEqual(report["net_profit"]["current"], Decimal(110))

    async def test_cross_store_and_missing_owner_do_not_guess_store_profit(self):
        for owner in (8, None):
            auto = await _store_provision_financial_values(fixture(owner=owner), 7, _month_periods(date(2026,9,1)))
            self.assertIsNone(auto["inventory_impairment"]["current"])
            self.assertIsNone(auto["inventory_writeoff"]["current"])
            self.assertTrue(auto["_warnings"])
            values = _build_store_profit_loss_values(auto, {})
            self.assertIsNone(values["net_profit"]["current"])
            self.assertIsNone(values["gross_profit"]["current"])
            workbook = Workbook()
            _build_store_profit_loss_summary_sheet(workbook, [(SimpleNamespace(name="店铺"), values)], date(2026,9,1))
            sheet = workbook.worksheets[-1]
            row = next(row for row in sheet if row[0].value == "三、净利润")
            self.assertIsNone(row[1].value)
            self.assertIsNone(row[2].value)

    async def test_replay_is_cached_only_within_the_same_db_session(self):
        db = fixture()
        await _load_inventory_impairment_report_rows(db, REPORT)
        await _store_provision_financial_values(db, 7, _month_periods(date(2026,9,1)))
        adjustments = await _sale_provision_adjustments(db, REPORT)
        self.assertEqual(adjustments[1], (Decimal(5), ""))
        self.assertEqual(adjustments[2], (Decimal(30), ""))
        self.assertEqual(db.execute.await_count, 3)

    async def test_later_safety_rule_preserves_other_store_previous_month_amounts(self):
        def data():
            db = fixture(owner=8, sale_store=8, invalid=True)
            batches, rules, sales = [item.all() for item in db.execute.side_effect]
            first = batches[0]
            batches[0] = (*first[:4], "stable", 50, first[6])
            batches.append((SimpleNamespace(id=2, product_id=10, quantity=100, unit_cost=Decimal(10),
                                             arrived_at=datetime(2026, 8, 31), batch_no="B"),
                            "SKU-1", "测试物料", "店铺A", "stable", 50, 7))
            rules.append(SimpleNamespace(product_id=10, product_type="stable", safe_stock_quantity=50,
                                         effective_date=date(2026, 10, 1)))
            return SimpleNamespace(execute=AsyncMock(side_effect=[result(batches), result(rules), result(sales)]),
                                   commit=AsyncMock(), info={})

        september_db, october_db = data(), data()
        september = await _load_inventory_impairment_report_rows(september_db, REPORT, 7)
        october = await _load_inventory_impairment_report_rows(october_db, date(2026, 10, 31), 7)
        self.assertEqual(september[0]["impairment_amount"], Decimal(150))
        self.assertEqual(october[0]["opening_impairment_amount"], Decimal(150))
        self.assertEqual(october[0]["previous_book_value"], Decimal(850))
        september_profit = await _store_provision_financial_values(september_db, 7, _month_periods(date(2026, 9, 1)))
        october_profit = await _store_provision_financial_values(october_db, 7, _month_periods(date(2026, 10, 1)))
        self.assertEqual(september_profit["inventory_impairment"]["current"], Decimal(150))
        self.assertEqual(october_profit["inventory_impairment"]["previous"], Decimal(150))
        self.assertIsNone(october_profit["inventory_impairment"]["current"])
        workbook, cache = _build_inventory_impairment_workbook(october, date(2026, 10, 31))
        sheet = load_workbook(_workbook_bytes_with_formula_cache(workbook, cache), data_only=True).active
        self.assertEqual(sheet["U2"].value, 150)

    async def test_january_reads_december_sales_and_oa_for_comparison_writeoffs(self):
        db = SimpleNamespace(execute=AsyncMock(return_value=result([])), scalar=AsyncMock(return_value=1))
        zero = {key: {period: Decimal(0) for period in ("current", "previous", "ytd")}
                for key in ("inventory_impairment", "inventory_writeoff")}
        with patch("routers.reports._store_provision_financial_values", new_callable=AsyncMock, return_value=zero), \
             patch("routers.reports.office_space_totals_by_application_date", new_callable=AsyncMock, return_value={}) as office, \
             patch("routers.reports.shared_admin_totals_by_application_date", new_callable=AsyncMock, return_value={}), \
             patch("routers.reports.china_salary_totals_by_store_and_application_date", new_callable=AsyncMock, return_value={}):
            await _store_profit_loss_auto_values(db, 7, date(2026,1,1), "店铺")
        params = db.execute.await_args.args[0].compile().params.values()
        self.assertIn(datetime(2025,12,1), params)
        self.assertIn(date(2025,12,1), office.await_args.args)

    async def test_product_line_deducts_the_same_writeoff_and_masks_unknown_cost(self):
        sale = SimpleNamespace(id=1, store_id=7, product_id=10, sold_at=datetime(2026,9,10),
                               selling_price=Decimal(20), quantity=10, platform_fee=Decimal(0))
        batch = SimpleNamespace(quantity=100, purchase_price=Decimal(10), shipping_cost=Decimal(0), last_mile_cost=Decimal(0), other_cost=Decimal(1000))
        detail = SimpleNamespace(id=1, quantity=10, unit_cost=Decimal("20.0001"))
        for adjustment in (Decimal(5), None):
            db = SimpleNamespace(execute=AsyncMock(side_effect=[result([]), result([(sale,7)]),
                                                                 result([]), result([(detail,batch,sale)]), result([])]))
            with patch("routers.sales._sale_import_context", new_callable=AsyncMock, return_value=({},{})), \
                 patch("routers.reports._sale_provision_adjustments", new_callable=AsyncMock,
                       return_value={1:(adjustment,"转销归属待核对" if adjustment is None else "")}):
                rows = await _product_line_profit_data(db, {7}, date(2026,9,1), REPORT)
            if adjustment is None:
                self.assertIsNone(rows[0]["current"]["cost"])
                self.assertIn("待核对", rows[0]["note"])
            else:
                self.assertEqual(rows[0]["current"]["cost"], Decimal("195.001"))
            workbook = Workbook()
            _build_product_line_profit_sheet(workbook, rows, date(2026,9,1))
            sheet = workbook.worksheets[-1]
            self.assertEqual(sheet["D2"].value, None if adjustment is None else 195.001)
            if adjustment is None:
                self.assertIsNone(sheet["E2"].value)
                self.assertIsNone(sheet["D3"].value)


if __name__ == "__main__":
    unittest.main()
