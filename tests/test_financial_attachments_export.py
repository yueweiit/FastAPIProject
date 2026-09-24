import unittest
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import Workbook

from routers.reports import (
    PRODUCT_LINE_PROFIT_SHEET,
    STORE_PROFIT_LOSS_SUMMARY_SHEET,
    STORE_PROFIT_LOSS_ROWS,
    _build_product_line_profit_sheet,
    _build_store_profit_loss_sheet,
    _build_store_profit_loss_summary_sheet,
    _build_store_profit_loss_values,
    _allocate_store_profit_loss_to_product_lines,
    _product_line_mapping_for_sale,
)


class FinancialAttachmentsWorkbookTests(unittest.TestCase):
    def test_store_level_revenue_and_costs_are_allocated_by_product_line_sales(self):
        line_data = {
            1: {
                "current": {"revenue": Decimal("80"), "cost": Decimal("30")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("80"), "cost": Decimal("30")},
            },
            2: {
                "current": {"revenue": Decimal("20"), "cost": Decimal("10")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("20"), "cost": Decimal("10")},
            },
        }
        revenue_by_store_line = {
            7: {
                1: {"current": Decimal("80"), "previous": Decimal("0"), "ytd": Decimal("80")},
                2: {"current": Decimal("20"), "previous": Decimal("0"), "ytd": Decimal("20")},
            },
        }
        manual_values_by_store = {
            7: {
                "shipping_revenue": {"current": "10", "ytd": "10"},
                "other_revenue": {"current": "20", "ytd": "20"},
                "customs_import_tax": {"current": "30", "ytd": "30"},
                "local_logistics_cost": {"current": "40", "ytd": "40"},
                "fulfillment_cost": {"current": "50", "ytd": "50"},
            },
        }

        _allocate_store_profit_loss_to_product_lines(
            line_data, revenue_by_store_line, manual_values_by_store
        )

        self.assertEqual(line_data[1]["current"], {
            "revenue": Decimal("104"), "cost": Decimal("126"),
        })
        self.assertEqual(line_data[2]["current"], {
            "revenue": Decimal("26"), "cost": Decimal("34"),
        })

    def test_store_level_allocations_do_not_mix_between_stores(self):
        line_data = {
            1: {
                "current": {"revenue": Decimal("100"), "cost": Decimal("0")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("100"), "cost": Decimal("0")},
            },
            2: {
                "current": {"revenue": Decimal("100"), "cost": Decimal("0")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("100"), "cost": Decimal("0")},
            },
        }
        revenue_by_store_line = {
            1: {1: {"current": Decimal("100"), "previous": Decimal("0"), "ytd": Decimal("100")}},
            2: {2: {"current": Decimal("100"), "previous": Decimal("0"), "ytd": Decimal("100")}},
        }
        manual_values_by_store = {
            1: {"shipping_revenue": {"current": "10", "ytd": "10"}},
            2: {"shipping_revenue": {"current": "30", "ytd": "30"}},
        }

        _allocate_store_profit_loss_to_product_lines(
            line_data, revenue_by_store_line, manual_values_by_store
        )

        self.assertEqual(line_data[1]["current"]["revenue"], Decimal("110"))
        self.assertEqual(line_data[2]["current"]["revenue"], Decimal("130"))

    def _values(self, sales: Decimal, purchase_cost: Decimal) -> dict:
        auto_values = {
            "product_sales_revenue": {"current": sales, "previous": Decimal("80"), "ytd": sales + Decimal("80")},
            "product_purchase_cost": {"current": purchase_cost, "previous": Decimal("30"), "ytd": purchase_cost + Decimal("30")},
        }
        return _build_store_profit_loss_values(auto_values, {})

    def test_store_profit_loss_sheet_uses_all_period_values(self):
        workbook = Workbook()
        values = self._values(Decimal("100"), Decimal("40"))
        store = SimpleNamespace(name="测试店铺", platform="tiktok_shop")

        _build_store_profit_loss_sheet(workbook, store, date(2026, 9, 1), values, {})

        sheet = workbook["09-店铺损益表"]
        self.assertEqual(sheet.max_column, 5)
        self.assertEqual(sheet[1][0].value, "项目")
        self.assertEqual(sheet[3][0].value, "  商品销售收入")
        self.assertEqual(sheet[3][1].value, 100)
        self.assertEqual(sheet[3][2].value, 80)
        self.assertEqual(sheet[3][3].value, 180)
        self.assertNotIn("23", [cell.value for row in sheet.iter_rows() for cell in row])

    def test_admin_summary_reconciles_store_totals(self):
        workbook = Workbook()
        store_a = SimpleNamespace(name="店铺A")
        store_b = SimpleNamespace(name="店铺B")
        reports = [
            (store_a, self._values(Decimal("100"), Decimal("40"))),
            (store_b, self._values(Decimal("50"), Decimal("10"))),
        ]

        _build_store_profit_loss_summary_sheet(workbook, reports, date(2026, 9, 1))

        sheet = workbook[STORE_PROFIT_LOSS_SUMMARY_SHEET]
        self.assertEqual(sheet[1][1].value, "店铺A")
        self.assertEqual(sheet[1][2].value, "店铺B")
        self.assertEqual(sheet[3][1].value, 100)
        self.assertEqual(sheet[3][2].value, 50)
        self.assertEqual(sheet[3][3].value, 150)
        self.assertAlmostEqual(sheet[3][4].value, 1.0)
        self.assertEqual(sheet[3][4].number_format, "0.00%")
        self.assertFalse(any(isinstance(cell.value, str) and cell.value.startswith("=") for row in sheet.iter_rows() for cell in row))

    def test_product_line_sheet_calculates_margins_and_exposes_unassigned_sales(self):
        workbook = Workbook()
        rows = [
            {
                "name": "宠物玩具",
                "sku_count": 2,
                "sku_ids": {1, 2},
                "current": {"revenue": Decimal("100"), "cost": Decimal("40")},
                "previous": {"revenue": Decimal("80"), "cost": Decimal("40")},
                "ytd": {"revenue": Decimal("180"), "cost": Decimal("80")},
                "note": "",
            },
            {
                "name": "未分配产品线",
                "sku_count": 1,
                "sku_ids": {3},
                "current": {"revenue": Decimal("20"), "cost": Decimal("5")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("20"), "cost": Decimal("5")},
                "note": "销售记录未匹配到有效店铺 SKU 产品线映射",
            },
        ]

        _build_product_line_profit_sheet(workbook, rows, date(2026, 9, 1))

        sheet = workbook[PRODUCT_LINE_PROFIT_SHEET]
        self.assertEqual(sheet[2][0].value, "宠物玩具")
        self.assertEqual(sheet[2][4].value, 60)
        self.assertAlmostEqual(sheet[2][5].value, 0.6)
        self.assertEqual(sheet[3][11].value, "销售记录未匹配到有效店铺 SKU 产品线映射")
        self.assertEqual(sheet[4][1].value, 3)
        self.assertAlmostEqual(sheet[4][5].value, 75 / 120)

    def test_product_line_cost_can_be_attributed_when_sale_revenue_rate_is_missing(self):
        """FIFO cost remains usable even when the sales revenue FX rate is unavailable."""
        rows = [
            {
                "name": "缺少汇率产品线",
                "sku_count": 1,
                "sku_ids": {9},
                "current": {"revenue": Decimal("0"), "cost": Decimal("45")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("0"), "cost": Decimal("45")},
                "note": "销售收入因汇率缺失不可用；FIFO 成本仍按销售成本明细计入。",
            },
        ]
        workbook = Workbook()

        _build_product_line_profit_sheet(workbook, rows, date(2026, 9, 1))

        sheet = workbook[PRODUCT_LINE_PROFIT_SHEET]
        self.assertEqual(sheet[2][2].value, 0)
        self.assertEqual(sheet[2][3].value, 45)
        self.assertEqual(sheet[2][4].value, -45)
        self.assertIsNone(sheet[2][5].value)

    def test_product_line_mapping_uses_source_sku_and_rejects_ambiguous_fallback(self):
        store = SimpleNamespace(id=1)
        product = SimpleNamespace(id=10)
        line_a = SimpleNamespace(id=1, name="产品线A")
        line_b = SimpleNamespace(id=2, name="产品线B")
        mapping_a = SimpleNamespace(
            store_id=1, product_id=10, platform_sku_id="SKU-A", store_sku="STORE-A",
            effective_from=date(2026, 1, 1), effective_to=None, is_active=True,
        )
        mapping_b = SimpleNamespace(
            store_id=1, product_id=10, platform_sku_id="SKU-B", store_sku="STORE-B",
            effective_from=date(2026, 1, 1), effective_to=None, is_active=True,
        )
        mappings = {(store.id, product.id): [(mapping_a, line_a, product), (mapping_b, line_b, product)]}

        selected = _product_line_mapping_for_sale(
            mappings, 1, 10, date(2026, 9, 1), {"SKU-B"}
        )
        self.assertIs(selected[1], line_b)
        self.assertIsNone(
            _product_line_mapping_for_sale(mappings, 1, 10, date(2026, 9, 1))
        )
        self.assertIsNone(
            _product_line_mapping_for_sale(
                {(1, 10): [(mapping_a, line_a, product)]},
                1,
                10,
                date(2026, 9, 1),
            )
        )


if __name__ == "__main__":
    unittest.main()
