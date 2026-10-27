"""历史批次作废保留销售与审计；使用隔离 SQLite，不接触业务数据库。"""
import io
import unittest
from datetime import datetime, timezone
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import patch

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session

from auth import get_current_user
from database import Base, get_db
from models import BatchChangeLog, InventoryBatch, Product, Sale, SaleCostDetail, User
from routers import batches
from routers.products import _product_summary_stmt
from routers.sales import _build_sale_response, _monthly_sales_summary
from services.cost_layers import sale_cost_context
from services.fifo import lock_fifo_batches
from tests.test_batch_edit import AsyncFacade


class BatchVoidTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        Base.metadata.create_all(self.engine)
        self.session = Session(self.engine, expire_on_commit=False)
        self.session.add_all([
            User(id=1, username="admin", role="admin", password_hash="test"),
            User(id=2, username="operator", role="operator", password_hash="test"),
            Product(id=10, sku="VOID-TEST", name="作废测试"),
            InventoryBatch(id=20, product_id=10, user_id=2, batch_no="VOID-001",
                           quantity=10, remaining_quantity=7, unit_cost=Decimal("2"),
                           purchase_price=Decimal("2"), total_amount=Decimal("20"),
                           shipping_cost=0, last_mile_cost=0, other_cost=0,
                           arrived_at=datetime(2026, 8, 1)),
            Sale(id=30, product_id=10, user_id=2, order_no="VOID-SALE", quantity=3,
                 selling_price=Decimal("5"), total_cost=Decimal("6"), platform_fee=0,
                 profit=Decimal("9"), sold_at=datetime(2026, 9, 1)),
            SaleCostDetail(id=40, sale_id=30, batch_id=20, quantity=3, unit_cost=Decimal("2")),
        ])
        self.session.commit()
        self.actor = SimpleNamespace(id=1, role="admin", username="admin")
        self.app = FastAPI()
        self.app.include_router(batches.router)

        async def actor():
            return self.actor

        async def db():
            try:
                yield AsyncFacade(self.session)
            finally:
                self.session.rollback()

        self.app.dependency_overrides[get_current_user] = actor
        self.app.dependency_overrides[get_db] = db

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    async def request(self, method, path, **kwargs):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test") as client:
            return await client.request(method, path, **kwargs)

    async def void(self, version=None, reason="历史盘点重复入库"):
        return await self.request("POST", "/batches/20/void", json={
            "edit_version": version or self.session.get(InventoryBatch, 20).edit_version,
            "reason": reason,
        })

    async def test_void_preserves_sale_and_references_but_removes_live_stock(self):
        response = await self.void()
        self.assertEqual(response.status_code, 200, response.text)
        batch = self.session.get(InventoryBatch, 20)
        self.assertEqual((batch.quantity, batch.remaining_quantity), (10, 7))
        self.assertEqual((batch.voided_by, batch.void_reason), (1, "历史盘点重复入库"))
        self.assertIsNotNone(batch.voided_at)
        self.assertIsNotNone(self.session.get(Sale, 30))
        self.assertEqual(self.session.get(SaleCostDetail, 40).batch_id, 20)
        log = self.session.scalar(select(BatchChangeLog))
        self.assertEqual((log.action, log.actor_id), ("void", 1))
        self.assertIsNone(log.before_values["voided_at"])
        self.assertEqual(log.after_values["void_reason"], "历史盘点重复入库")

        active = await self.request("GET", "/batches?page=1")
        archived = await self.request("GET", "/batches?page=1&voided_only=true")
        self.assertEqual(active.json()["total"], 0)
        self.assertEqual(archived.json()["total"], 1)
        self.assertEqual(await lock_fifo_batches(AsyncFacade(self.session), [10]), [])
        product_summary = self.session.execute(_product_summary_stmt()).one()
        self.assertEqual(product_summary.stock_quantity, 0)
        self.assertEqual(product_summary.stock_value, Decimal(0))
        self.assertEqual((await self.request("DELETE", "/batches/20")).status_code, 409)
        self.assertEqual((await self.void()).status_code, 409)

    async def test_void_requires_write_role_reason_and_current_version(self):
        version = self.session.get(InventoryBatch, 20).edit_version
        self.actor.role = "viewer"
        self.assertEqual((await self.void(version)).status_code, 403)
        self.assertEqual((await self.request("GET", "/batches?voided_only=true")).status_code, 403)
        self.actor.role = "operator"
        self.assertEqual((await self.void(version, "   ")).status_code, 422)
        self.session.get(InventoryBatch, 20).shipping_cost_no = "new-bill"
        self.session.commit()
        conflict = await self.void(version)
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(conflict.json()["detail"]["code"], "batch_edit_conflict")
        self.assertIsNone(self.session.get(InventoryBatch, 20).voided_at)
        self.assertEqual(self.session.scalars(select(BatchChangeLog)).all(), [])

    async def test_operator_can_void_other_creators_batch_and_read_archive(self):
        self.actor.role = "operator"
        self.assertEqual((await self.request("DELETE", "/batches/20")).status_code, 403)
        response = await self.void()
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["voided_by"], self.actor.id)
        log = self.session.scalar(select(BatchChangeLog))
        self.assertEqual((log.action, log.actor_role, log.actor_id), ("void", "operator", 1))
        archived = await self.request("GET", "/batches?voided_only=true")
        self.assertEqual(len(archived.json()), 1)
        self.assertEqual((await self.request("GET", "/batches/export?batch_ids=20")).status_code, 200)

    async def test_void_audit_and_response_use_persisted_timestamp(self):
        def round_to_database_precision(session, _flush_context, _instances):
            for row in session.dirty:
                if isinstance(row, InventoryBatch) and row.voided_at is not None:
                    row.voided_at = row.voided_at.replace(microsecond=0)

        event.listen(self.session, "before_flush", round_to_database_precision)
        try:
            with patch.object(batches, "datetime") as clock:
                clock.now.return_value = datetime(2026, 10, 27, 12, 34, 56, 123456, timezone.utc)
                response = await self.void()
        finally:
            event.remove(self.session, "before_flush", round_to_database_precision)
        self.assertEqual(response.status_code, 200, response.text)
        persisted = self.session.get(InventoryBatch, 20).voided_at
        logged = self.session.scalar(select(BatchChangeLog)).after_values["voided_at"]
        self.assertEqual(persisted, datetime(2026, 10, 27, 12, 34, 56))
        self.assertEqual(logged, "2026-10-27T12:34:56+00:00")
        self.assertEqual(response.json()["voided_at"], logged)
        self.assertEqual(response.json()["edit_version"],
                         self.session.get(InventoryBatch, 20).edit_version)

    async def test_voided_sale_is_visible_with_cost_pending_review(self):
        await self.void()
        sale = self.session.get(Sale, 30)
        context = (await sale_cost_context(AsyncFacade(self.session), [sale]))[30]
        self.assertEqual(context["cost_status"], "needs_review")
        self.assertIn("VOID-001", context["cost_review_reason"])
        response = _build_sale_response(sale, [{"batch_id": 20, "batch_no": "VOID-001",
                                                "quantity": 3, "unit_cost": Decimal("2")}], cost_context=context)
        self.assertEqual(response.total_cost, Decimal("6"))
        self.assertIsNone(response.profit_cny)
        monthly = _monthly_sales_summary([sale], {}, cost_contexts={30: context})
        self.assertEqual(monthly[0].cost_review_count, 1)
        self.assertEqual(monthly[0].sales_cost_cny, Decimal("6"))
        self.assertIsNone(monthly[0].gross_profit_cny)

    async def test_archived_export_restricts_viewer_and_reason_is_plain_text(self):
        reason = '=HYPERLINK("https://example.invalid", "bad")'
        self.assertEqual((await self.void(reason=reason)).status_code, 200)
        self.actor.role = "viewer"
        self.assertEqual((await self.request("GET", "/batches/export?batch_ids=20")).status_code, 404)
        self.actor.role = "operator"
        response = await self.request("GET", "/batches/export?voided_only=true")
        self.assertEqual(response.status_code, 200, response.text)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(1, sheet.max_column).value, "作废原因")
        self.assertEqual(sheet.cell(2, sheet.max_column).value, reason)
        self.assertEqual(sheet.cell(2, sheet.max_column).data_type, "s")

    async def test_audit_failure_rolls_back_void(self):
        def fail_audit(*_args):
            raise RuntimeError("audit failure")

        event.listen(BatchChangeLog, "before_insert", fail_audit)
        try:
            with self.assertRaisesRegex(RuntimeError, "audit failure"):
                await self.void()
        finally:
            event.remove(BatchChangeLog, "before_insert", fail_audit)
            self.session.rollback()

        batch = self.session.get(InventoryBatch, 20)
        self.assertIsNone(batch.voided_at)
        self.assertIsNone(batch.void_reason)
        self.assertEqual(self.session.scalars(select(BatchChangeLog)).all(), [])


if __name__ == "__main__":
    unittest.main()
