"""商品导入真实 MySQL 验证；只允许本机、专用数据库，不连接项目业务库。

运行前显式设置 PRODUCT_IMPORT_TEST_HOST/PORT/DATABASE/USER/PASSWORD。
数据库名必须以 inventory_fifo_product_import_validation_ 开头；需事先创建。
本脚本创建商品/规则两张表、追加带随机前缀的测试数据，不删除任何数据。
"""

import asyncio
import json
import os
import re
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import httpx
from fastapi import FastAPI, HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.engine import URL
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from auth import get_current_user
from database import Base, get_db
from models import Product, ProductImpairmentRule
from routers.products import router
from services.product_import import (
    IMPORT_HEADERS, MAX_IMPORT_ROWS, confirm_product_import,
    issue_product_import_token, parse_product_import, preview_product_import,
)


def test_url() -> URL:
    host = os.environ.get("PRODUCT_IMPORT_TEST_HOST")
    database = os.environ.get("PRODUCT_IMPORT_TEST_DATABASE", "")
    if host != "127.0.0.1" or not re.fullmatch(
        r"inventory_fifo_product_import_validation_[a-z0-9_]{1,20}", database,
    ):
        raise RuntimeError("仅允许 127.0.0.1 上显式指定的商品导入专用测试库")
    return URL.create(
        "mysql+aiomysql", host=host,
        port=int(os.environ["PRODUCT_IMPORT_TEST_PORT"]),
        database=database, username=os.environ["PRODUCT_IMPORT_TEST_USER"],
        password=os.environ["PRODUCT_IMPORT_TEST_PASSWORD"], query={"charset": "utf8mb4"},
    )


def spreadsheet(rows):
    workbook = Workbook()
    workbook.active.append(IMPORT_HEADERS)
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def main() -> None:
    engine = create_async_engine(test_url(), hide_parameters=True)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    prefix = "IMPORT-" + uuid4().hex[:12] + "-"
    secret = uuid4().hex  # 仅本次验证使用，不打印、不保存。
    checks = []

    async def counts():
        async with sessions() as db:
            return tuple([await db.scalar(select(func.count()).select_from(model))
                          for model in (Product, ProductImpairmentRule)])

    async def prepare(rows):
        parsed = parse_product_import("test.xlsx", spreadsheet(rows))
        async with sessions() as db:
            report = await preview_product_import(db, parsed)
        token = issue_product_import_token(parsed, report, 1, secret)
        return parsed, report, token

    try:
        async with engine.begin() as connection:
            version = (await connection.execute(text("SELECT VERSION()"))).scalar_one()
            await connection.run_sync(lambda sync: Base.metadata.create_all(
                sync, tables=[Product.__table__, ProductImpairmentRule.__table__],
            ))

        # 真实排序规则识别文件内大小写/重音差异造成的唯一键冲突。
        before = await counts()
        _, report, _ = await prepare([
            [prefix + "CAFÉ", "甲", "稳健商品", 1],
            [prefix.lower() + "cafe", "乙", "新品", 0],
        ])
        assert report.error_count == 2 and not report.can_import
        assert await counts() == before
        checks.append("实库排序规则查重及预览零写入")

        content = spreadsheet([
            [prefix + "0001", "前导零商品", "新品", 0],
            [prefix + "CAFÉ", "稳健商品", "稳健商品", 12],
        ])
        parsed = parse_product_import("test.xlsx", content)
        async with sessions() as db:
            report = await preview_product_import(db, parsed)
            token = issue_product_import_token(parsed, report, 1, secret)
            result = await confirm_product_import(db, parsed, token, 1, secret)
        assert result.created_count == 2
        assert await counts() == (before[0] + 2, before[1] + 2)
        async with sessions() as db:
            rules = (await db.execute(select(ProductImpairmentRule).join(Product).where(
                Product.sku.in_([prefix + "0001", prefix + "CAFÉ"]),
            ))).scalars().all()
            assert len(rules) == 2 and all(rule.effective_date == date.today() for rule in rules)
        checks.append("商品与当日规则统一提交")

        _, report, _ = await prepare([[prefix.lower() + "cafe", "不覆盖旧名称", "新品", 0]])
        assert report.skipped_count == 1 and report.rows[0].existing_product_name == "稳健商品"
        _, report, _ = await prepare([
            [prefix + "0001", "前导零商品", "新品", 0],
            [prefix + "CAFÉ", "稳健商品", "稳健商品", 12],
        ])
        assert report.skipped_count == 2 and not report.can_import
        async with sessions() as db:
            try:
                await confirm_product_import(db, parsed, token, 1, secret)
            except HTTPException as error:
                assert error.status_code == 409
            else:
                raise AssertionError("旧预览被重复执行")
        checks.append("重复导入跳过且旧预览拒绝再次提交")

        parsed, _, token = await prepare([[prefix + "STALE", "旧预览", "新品", 0]])
        async with sessions() as db:
            db.add(Product(sku=prefix + "STALE", name="并发手工新增", product_type="new"))
            await db.commit()
        async with sessions() as db:
            try:
                await confirm_product_import(db, parsed, token, 1, secret)
            except HTTPException as error:
                assert error.status_code == 409
            else:
                raise AssertionError("数据库变化未使预览失效")
        checks.append("确认阶段重新检查数据库变化")

        # 在商品已 flush、规则尚未 commit 时制造实库唯一键错误，检查两表回滚。
        class FailingRuleSession(AsyncSession):
            async def commit(self):
                rule = next(item for item in self.new if isinstance(item, ProductImpairmentRule))
                self.add(ProductImpairmentRule(
                    product_id=rule.product_id, product_type=rule.product_type,
                    safe_stock_quantity=rule.safe_stock_quantity, effective_date=rule.effective_date,
                ))
                await super().commit()

        failing_sessions = async_sessionmaker(engine, class_=FailingRuleSession, expire_on_commit=False)
        parsed, _, token = await prepare([
            [prefix + "ROLLBACK-A", "回滚甲", "新品", 0],
            [prefix + "ROLLBACK-B", "回滚乙", "稳健商品", 1],
        ])
        before = await counts()
        async with failing_sessions() as db:
            try:
                await confirm_product_import(db, parsed, token, 1, secret)
            except HTTPException as error:
                assert error.status_code == 409
            else:
                raise AssertionError("规则唯一键冲突没有阻止提交")
        assert await counts() == before
        checks.append("真实规则写入失败时商品与规则整批回滚")

        # 两个事务均完成确认校验后，才开始 INSERT，强制覆盖实际并发冲突窗口。
        barrier = asyncio.Barrier(2)

        class ConcurrentSession(AsyncSession):
            async def flush(self, objects=None):
                if any(isinstance(item, Product) for item in self.new):
                    await asyncio.wait_for(barrier.wait(), timeout=10)
                await super().flush(objects)

        concurrent_sessions = async_sessionmaker(engine, class_=ConcurrentSession, expire_on_commit=False)
        candidates = []
        for side in ("A", "B"):
            candidates.append(await prepare([
                [prefix + "CONCURRENT", "并发共享 SKU", "新品", 0],
                [prefix + "CONCURRENT-" + side, "并发独有 SKU", "新品", 0],
            ]))

        async def contender(candidate):
            parsed, _, token = candidate
            async with concurrent_sessions() as db:
                try:
                    result = await confirm_product_import(db, parsed, token, 1, secret)
                    return result.created_count
                except HTTPException as error:
                    return error.status_code

        before = await counts()
        outcomes = await asyncio.wait_for(asyncio.gather(*[contender(item) for item in candidates]), timeout=20)
        assert sorted(outcomes) == [2, 409], outcomes
        assert await counts() == (before[0] + 2, before[1] + 2)
        checks.append("真实并发冲突仅一批成功且另一批完整回滚")

        rows = [[prefix + f"LIMIT-{index:04d}", "上限商品", "新品", 0]
                for index in range(MAX_IMPORT_ROWS)]
        parsed, report, token = await prepare(rows)
        assert report.new_count == MAX_IMPORT_ROWS and report.can_import
        before = await counts()
        async with sessions() as db:
            result = await confirm_product_import(db, parsed, token, 1, secret)
        assert result.created_count == MAX_IMPORT_ROWS
        assert await counts() == (before[0] + MAX_IMPORT_ROWS, before[1] + MAX_IMPORT_ROWS)
        checks.append("2000 条上限真实预览和批量写入")

        # 不导入 main，不运行生产启动迁移；只覆盖测试应用中的数据库和用户依赖。
        app = FastAPI()
        app.include_router(router)

        async def test_db():
            async with sessions() as db:
                yield db

        async def test_user():
            return SimpleNamespace(id=1, role="viewer")

        app.dependency_overrides[get_db] = test_db
        app.dependency_overrides[get_current_user] = test_user
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://local-test") as client:
            response = await client.get("/products/import-template")
            assert response.status_code == 200
            assert load_workbook(BytesIO(response.content)).sheetnames == ["商品导入", "填写说明"]
            content = spreadsheet([[prefix + "HTTP", "HTTP 商品", "新品", 0]])
            response = await client.post("/products/import/preview", files={"file": ("products.xlsx", content)})
            assert response.status_code == 200 and response.json()["can_import"]
            response = await client.post("/products/import", files={"file": ("products.xlsx", content)},
                                         data={"preview_token": response.json()["preview_token"]})
            assert response.status_code == 200 and response.json()["created_count"] == 1
            rows = response.json()["rows"]
            response = await client.post("/products/import/result", json={"rows": rows})
            assert response.status_code == 200
            assert load_workbook(BytesIO(response.content)).active["F2"].value == "已新增"
        checks.append("查看者通过真实数据库完成模板、预览、确认与结果下载")
        print(json.dumps({"version": version, "database": test_url().database, "checks": checks}, ensure_ascii=False))
    finally:
        await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
