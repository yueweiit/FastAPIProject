from datetime import datetime
from decimal import Decimal
from io import BytesIO
import re
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from database import get_db
from models import InventoryBatch, Product, User, SaleCostDetail
from schemas import BatchCreate, BatchResponse
from auth import get_current_user, RequireOperator, RequireAnyRole

router = APIRouter(prefix="/batches", tags=["入库管理"])
BASE_DIR = Path(__file__).parent.parent


def _calc_unit_cost(price, qty, ship, last_mile, other):
    price = Decimal(str(price))
    ship = Decimal(str(ship))
    last_mile = Decimal(str(last_mile))
    other = Decimal(str(other))
    return price + (ship + last_mile + other) / Decimal(qty)


def _resolve_purchase_values(data: BatchCreate, existing_batch: InventoryBatch | None = None):
    quantity = Decimal(data.quantity)
    if "total_amount" in data.model_fields_set and data.total_amount is not None:
        total_amount = Decimal(str(data.total_amount))
        return total_amount / quantity, total_amount

    if "purchase_price" in data.model_fields_set:
        purchase_price = Decimal(str(data.purchase_price))
        return purchase_price, purchase_price * quantity

    if existing_batch is not None and existing_batch.total_amount is not None:
        total_amount = Decimal(str(existing_batch.total_amount))
        return total_amount / quantity, total_amount

    purchase_price = Decimal(str(existing_batch.purchase_price if existing_batch else data.purchase_price))
    return purchase_price, purchase_price * quantity


async def _next_batch_no(db: AsyncSession, prefix: str) -> str:
    result = await db.execute(
        select(InventoryBatch.batch_no).where(InventoryBatch.batch_no.like(f"{prefix}-%"))
    )
    pattern = re.compile(rf"^{re.escape(prefix)}-(\d+)$")
    max_sequence = 0
    for batch_no in result.scalars():
        match = pattern.fullmatch(batch_no)
        if match:
            max_sequence = max(max_sequence, int(match.group(1)))
    return f"{prefix}-{max_sequence + 1:03d}"


@router.post("", response_model=BatchResponse)
async def create_batch(
    data: BatchCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    """创建入库批次 - admin和operator可操作"""
    product = await db.get(Product, data.product_id)
    if not product:
        raise HTTPException(status_code=404, detail=f"商品不存在: {data.product_id}")

    purchase_price, total_amount = _resolve_purchase_values(data)
    unit_cost = _calc_unit_cost(purchase_price, data.quantity, data.shipping_cost, data.last_mile_cost, data.other_cost)

    arrived = data.arrived_at or datetime.now()
    prefix = f"LatinGo{arrived.strftime('%Y%m%d')}"

    for attempt in range(3):
        batch_no = await _next_batch_no(db, prefix)
        batch = InventoryBatch(
            product_id=data.product_id,
            user_id=user.id,
            batch_no=batch_no,
            quantity=data.quantity,
            remaining_quantity=data.quantity,
            unit_cost=unit_cost,
            purchase_price=purchase_price,
            dingtalk_order_no=data.dingtalk_order_no or None,
            total_amount=total_amount,
            shipping_cost=data.shipping_cost,
            shipping_cost_no=data.shipping_cost_no or None,
            last_mile_cost=data.last_mile_cost,
            last_mile_cost_no=data.last_mile_cost_no or None,
            other_cost=data.other_cost,
            other_cost_no=data.other_cost_no or None,
            shipping_date=data.shipping_date,
            arrived_at=arrived,
        )
        db.add(batch)
        try:
            await db.commit()
        except IntegrityError:
            await db.rollback()
            if attempt == 2:
                raise HTTPException(status_code=409, detail="批次号生成冲突，请稍后重试")
            continue
        await db.refresh(batch)
        return batch

    raise HTTPException(status_code=409, detail="批次号生成冲突，请稍后重试")


@router.put("/{batch_id}", response_model=BatchResponse)
async def update_batch(
    batch_id: int,
    data: BatchCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    """编辑批次 - 只能改未被销售消耗的批次"""
    batch = await db.get(InventoryBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")

    # operator只能改自己的
    if user.role == "operator" and batch.user_id != user.id:
        raise HTTPException(status_code=403, detail="无权修改此批次")

    # 检查是否已被消耗
    if batch.remaining_quantity != batch.quantity:
        raise HTTPException(status_code=400, detail="批次已有出库记录，无法修改")

    purchase_price, total_amount = _resolve_purchase_values(data, batch)

    batch.quantity = data.quantity
    batch.remaining_quantity = data.quantity
    batch.purchase_price = purchase_price
    if "dingtalk_order_no" in data.model_fields_set:
        batch.dingtalk_order_no = data.dingtalk_order_no or None
    batch.total_amount = total_amount
    batch.shipping_cost = data.shipping_cost
    if "shipping_cost_no" in data.model_fields_set:
        batch.shipping_cost_no = data.shipping_cost_no or None
    batch.last_mile_cost = data.last_mile_cost
    if "last_mile_cost_no" in data.model_fields_set:
        batch.last_mile_cost_no = data.last_mile_cost_no or None
    batch.other_cost = data.other_cost
    if "other_cost_no" in data.model_fields_set:
        batch.other_cost_no = data.other_cost_no or None
    batch.shipping_date = data.shipping_date
    batch.arrived_at = data.arrived_at or batch.arrived_at
    batch.unit_cost = _calc_unit_cost(purchase_price, data.quantity, data.shipping_cost, data.last_mile_cost, data.other_cost)

    await db.commit()
    await db.refresh(batch)
    return batch


@router.delete("/{batch_id}")
async def delete_batch(
    batch_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    """删除批次 - 只能删未被消耗的批次"""
    batch = await db.get(InventoryBatch, batch_id)
    if not batch:
        raise HTTPException(status_code=404, detail="批次不存在")

    if user.role == "operator" and batch.user_id != user.id:
        raise HTTPException(status_code=403, detail="无权删除此批次")

    if batch.remaining_quantity != batch.quantity:
        raise HTTPException(status_code=400, detail="批次已有出库记录，无法删除")

    await db.delete(batch)
    await db.commit()
    return {"ok": True}


@router.get("", response_model=list[BatchResponse])
async def list_batches(
    product_id: int | None = None,
    keyword: str | None = Query(default=None, max_length=255),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """列出批次，可按商品名称、SKU 或批次号搜索。"""
    stmt = select(InventoryBatch)
    if product_id:
        stmt = stmt.where(InventoryBatch.product_id == product_id)
    if keyword and keyword.strip():
        pattern = f"%{keyword.strip()}%"
        stmt = stmt.join(Product).where(or_(
            Product.name.ilike(pattern),
            Product.sku.ilike(pattern),
            InventoryBatch.batch_no.ilike(pattern),
        ))
    if user.role == "operator":
        stmt = stmt.where(InventoryBatch.user_id == user.id)
    stmt = stmt.order_by(InventoryBatch.arrived_at.desc())
    result = await db.execute(stmt)
    return result.scalars().all()


@router.get("/export")
async def export_batches(
    batch_ids: list[int] = Query(..., min_length=1, description="要导出的批次 ID，可重复传入"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """导出一个或多个批次为 Excel，并嵌入商品图片。"""
    stmt = (
        select(InventoryBatch)
        .where(InventoryBatch.id.in_(batch_ids))
        .order_by(InventoryBatch.arrived_at.desc(), InventoryBatch.id.desc())
    )
    if user.role == "operator":
        stmt = stmt.where(InventoryBatch.user_id == user.id)
    result = await db.execute(stmt)
    batches = list(result.scalars().all())
    if not batches:
        raise HTTPException(status_code=404, detail="没有可导出的批次")

    product_ids = {batch.product_id for batch in batches}
    products_result = await db.execute(select(Product).where(Product.id.in_(product_ids)))
    products = {product.id: product for product in products_result.scalars().all()}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "批次数据"
    headers = [
        "图片", "批次ID", "商品名称", "SKU", "批次号", "钉钉单号", "入库数", "剩余数",
        "采购价", "总金额", "头程运费", "头程运费钉钉单号", "尾程派送", "尾程派送费钉钉单号",
        "其他成本", "其他成本钉钉单号", "单件成本", "头程发货时间", "入库时间", "创建时间", "创建人ID",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"

    for row_index, batch in enumerate(batches, start=2):
        product = products.get(batch.product_id)
        sheet.append([
            "", batch.id, product.name if product else "", product.sku if product else "", batch.batch_no,
            batch.dingtalk_order_no or "", batch.quantity, batch.remaining_quantity, batch.purchase_price,
            batch.total_amount, batch.shipping_cost, batch.shipping_cost_no or "", batch.last_mile_cost,
            batch.last_mile_cost_no or "", batch.other_cost, batch.other_cost_no or "", batch.unit_cost,
            batch.shipping_date, batch.arrived_at, batch.created_at, batch.user_id,
        ])
        image_path = None
        if product and product.image:
            candidate = (BASE_DIR / product.image.lstrip("/")).resolve()
            uploads_root = (BASE_DIR / "uploads").resolve()
            if candidate.is_file() and uploads_root in candidate.parents:
                image_path = candidate
        if image_path:
            try:
                image = ExcelImage(str(image_path))
                max_size = 72
                scale = min(max_size / image.width, max_size / image.height, 1)
                image.width = max(1, int(image.width * scale))
                image.height = max(1, int(image.height * scale))
                sheet.add_image(image, f"A{row_index}")
                sheet.row_dimensions[row_index].height = 60
            except (OSError, ValueError, TypeError):
                pass

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in ("I", "J", "K", "M", "O", "Q"):
        for cell in sheet[column][1:]:
            cell.number_format = '¥#,##0.0000'
    for column in ("R", "S", "T"):
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
    for column, width in {
        "A": 12, "B": 10, "C": 22, "D": 14, "E": 24, "F": 25, "G": 10, "H": 10,
        "I": 13, "J": 13, "K": 13, "L": 20, "M": 13, "N": 22, "O": 13, "P": 20,
        "Q": 13, "R": 20, "S": 20, "T": 20, "U": 12,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"batch_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
