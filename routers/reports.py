from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import and_, case, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from database import get_db
from models import (
    InventoryBatch,
    InventoryPeriodSnapshot,
    Product,
    ProductImpairmentRule,
    ProductLine,
    Sale,
    SaleCostDetail,
    Store,
    StoreProduct,
    StoreProfitLossReport,
    SettlementEntry,
    SettlementEntryAllocation,
    User,
)
from schemas import (
    MonthlyReportItem,
    StoreProfitLossCell,
    StoreProfitLossResponse,
    StoreProfitLossRowResponse,
    StoreProfitLossUpdateRequest,
)
from auth import RequireAnyRole, RequireOperator
from services.accounting_periods import ensure_period_snapshot
from services.oa_office_expenses import (
    china_salary_totals_by_store_and_application_date,
    office_space_totals_by_application_date,
)
from services.inventory_impairment import (
    InventoryLayer,
    ImpairmentRule,
    PRODUCT_TYPE_NEW,
    batch_impairments,
    daily_impairment_rate,
    impairment_rate,
)

router = APIRouter(prefix="/reports", tags=["报表"])

INVENTORY_IMPAIRMENT_SHEET = "07-存货跌价准备明细表"
INVENTORY_IMPAIRMENT_HEADERS = [
    "店铺",
    "SKU编码",
    "商品名称",
    "入库日期",
    "已入库天数",
    "单件到仓成本(元)",
    "库存数量",
    "计提数量",
    "库存总金额(元)",
    "累计计提比例",
    "累计跌价准备(元)",
    "账面价值(元)",
    "与上月末账面价值差额(元)",
    "状态",
    "备注",
]

STORE_PROFIT_LOSS_ROWS = (
    ("operating_revenue", "一、营业收入", "formula", None),
    ("product_sales_revenue", "  商品销售收入", "auto", "销售记录中的净商品收入"),
    ("shipping_revenue", "  运费收入", "manual", None),
    ("other_revenue", "  其他收入", "manual", None),
    ("operating_cost", "减：营业成本（对应）", "formula", None),
    ("product_purchase_cost", "  应商品采购成本", "auto", "FIFO 实际扣除数量 × 入库采购单价"),
    ("head_logistics_cost", "  头程尾程物流成本", "auto", "FIFO 实际扣除数量分摊入库头程和尾程物流费用"),
    ("customs_import_tax", "  关税及进口税费（正报）", "manual", None),
    ("local_logistics_cost", "  本地物流配送成本", "manual", None),
    ("fulfillment_cost", "  代发成本", "manual", None),
    ("other_direct_cost", "  其他直接成本", "auto", "销售记录中的其他费用"),
    ("gross_profit", "毛利", "formula", None),
    ("gross_margin", "毛利率", "formula", None),
    ("tax_surcharge", "减：税金及附加", "manual", None),
    ("selling_expenses", "减：销售费用（包括样本）", "formula", None),
    ("salary", "  人员薪资", "manual", "OA 工资中国明细按部门名称匹配店铺；可人工修改"),
    ("advertising", "  广告推广费", "manual", None),
    ("platform_subscription", "  平台月费/年费", "manual", None),
    ("warehousing", "  仓储费用", "manual", None),
    ("delivery", "  配送费", "manual", None),
    ("platform_fines", "  平台罚款/赔偿", "manual", None),
    ("admin_expenses", "减：管理费用", "formula", None),
    ("rent_utilities", "  房租+水电+网费", "auto", "OA 办公场地费用中 LatínGo 的租金和电费，均分至有效店铺"),
    ("shared_admin", "  平摊人事+财务管理费用", "manual", None),
    ("research_development", "减：研发费用", "manual", None),
    ("finance_expenses", "减：财务费用", "manual", None),
    ("asset_impairment_loss", "减：资产减值损失", "formula", None),
    ("inventory_impairment", "    存货跌价准备", "auto", "本月末计提额减上月末计提额"),
    ("credit_impairment_loss", "减：信用减值损失", "manual", None),
    ("operating_profit", "二、营业利润", "formula", None),
    ("non_operating_income", "加：营业外收入", "manual", None),
    ("non_operating_expense", "减：营业外支出", "manual", None),
    ("income_tax_expense", "减：所得税费用", "manual", None),
    ("net_profit", "三、净利润", "formula", None),
)
STORE_PROFIT_LOSS_AUTO_KEYS = {
    key for key, _label, kind, _source in STORE_PROFIT_LOSS_ROWS if kind == "auto"
}
STORE_PROFIT_LOSS_FORMULA_KEYS = {
    key for key, _label, kind, _source in STORE_PROFIT_LOSS_ROWS if kind == "formula"
}
STORE_PROFIT_LOSS_MANUAL_KEYS = {
    key for key, _label, kind, _source in STORE_PROFIT_LOSS_ROWS if kind == "manual"
}
STORE_PROFIT_LOSS_ROW_KEYS = {
    key for key, _label, _kind, _source in STORE_PROFIT_LOSS_ROWS
}
STORE_PROFIT_LOSS_PERIODS = ("current", "previous", "ytd")


def _previous_month_end(value: date) -> date:
    return value.replace(day=1) - timedelta(days=1)


def _excel_date(value: date) -> str:
    return f"DATE({value.year},{value.month},{value.day})"


def _next_month_start(value: date) -> date:
    if value.month == 12:
        return date(value.year + 1, 1, 1)
    return date(value.year, value.month + 1, 1)


def _month_periods(report_month: date) -> dict[str, tuple[date, date]]:
    month_start = report_month.replace(day=1)
    previous_start = _previous_month_end(month_start).replace(day=1)
    return {
        "current": (month_start, _next_month_start(month_start)),
        "previous": (previous_start, month_start),
        "ytd": (date(month_start.year, 1, 1), _next_month_start(month_start)),
    }


def _decimal_or_zero(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    return Decimal(str(value))


def _blank_profit_loss_cells() -> dict[str, dict[str, Decimal | None]]:
    return {
        key: {period: None for period in STORE_PROFIT_LOSS_PERIODS}
        for key, _label, _kind, _source in STORE_PROFIT_LOSS_ROWS
    }


def _build_store_profit_loss_values(
    auto_values: dict[str, dict[str, Decimal]],
    manual_values: dict,
) -> dict[str, dict[str, Decimal | None]]:
    """Combine generated/manual cells and apply the template's formulas."""
    values = _blank_profit_loss_cells()
    for key in STORE_PROFIT_LOSS_AUTO_KEYS:
        for period in STORE_PROFIT_LOSS_PERIODS:
            values[key][period] = auto_values.get(key, {}).get(period, Decimal("0"))
    for key in STORE_PROFIT_LOSS_MANUAL_KEYS:
        stored = manual_values.get(key, {}) if isinstance(manual_values, dict) else {}
        for period in STORE_PROFIT_LOSS_PERIODS:
            value = stored.get(period) if isinstance(stored, dict) else None
            if value in (None, ""):
                values[key][period] = auto_values.get(key, {}).get(period)
            else:
                values[key][period] = Decimal(str(value))

    for period in STORE_PROFIT_LOSS_PERIODS:
        get = lambda key: _decimal_or_zero(values[key][period])
        values["operating_revenue"][period] = sum(
            (get(key) for key in ("product_sales_revenue", "shipping_revenue", "other_revenue")),
            Decimal("0"),
        )
        values["operating_cost"][period] = sum(
            (get(key) for key in (
                "product_purchase_cost", "head_logistics_cost", "customs_import_tax",
                "local_logistics_cost", "fulfillment_cost", "other_direct_cost",
            )),
            Decimal("0"),
        )
        values["gross_profit"][period] = (
            get("operating_revenue") - get("operating_cost")
        )
        revenue = get("operating_revenue")
        values["gross_margin"][period] = (
            get("gross_profit") / revenue if revenue else None
        )
        values["selling_expenses"][period] = sum(
            (get(key) for key in (
                "salary", "advertising", "platform_subscription", "warehousing",
                "delivery", "platform_fines",
            )),
            Decimal("0"),
        )
        values["admin_expenses"][period] = get("rent_utilities") + get("shared_admin")
        values["asset_impairment_loss"][period] = get("inventory_impairment")
        values["operating_profit"][period] = (
            get("gross_profit")
            - get("tax_surcharge")
            - get("selling_expenses")
            - get("admin_expenses")
            - get("research_development")
            - get("finance_expenses")
            - get("asset_impairment_loss")
            - get("credit_impairment_loss")
        )
        values["net_profit"][period] = (
            get("operating_profit")
            + get("non_operating_income")
            - get("non_operating_expense")
            - get("income_tax_expense")
        )
    return values


def _store_sale_scope(store_id: int):
    return or_(
        Sale.store_id == store_id,
        and_(Sale.store_id.is_(None), User.store_id == store_id),
    )


def _add_period_value(
    target: dict[str, Decimal], sold_at: date, amount: Decimal,
    periods: dict[str, tuple[date, date]],
) -> None:
    for period, (start, end) in periods.items():
        if start <= sold_at < end:
            target[period] += amount


async def _store_profit_loss_auto_values(
    db: AsyncSession, store_id: int, report_month: date, store_name: str | None = None
) -> dict[str, dict[str, Decimal]]:
    periods = _month_periods(report_month)
    ytd_start, ytd_end = periods["ytd"]
    sales = list((await db.execute(
        select(Sale)
        .outerjoin(User, User.id == Sale.user_id)
        .where(
            _store_sale_scope(store_id),
            Sale.sold_at >= datetime.combine(ytd_start, time.min),
            Sale.sold_at < datetime.combine(ytd_end, time.min),
        )
    )).scalars().all())

    auto_values = {
        key: {period: Decimal("0") for period in STORE_PROFIT_LOSS_PERIODS}
        for key in STORE_PROFIT_LOSS_AUTO_KEYS
    }
    if sales:
        # Imported rows retain their source amount and historical FX rate in the
        # sales module; manual sales fall back to their stored CNY amount.
        from routers.sales import _sale_import_context

        import_context_by_sale, _ = await _sale_import_context(db, sales)
        for sale in sales:
            context = import_context_by_sale.get(sale.id, {})
            rate = context.get("exchange_rate_to_cny", Decimal("1"))
            if rate is None:
                continue
            revenue = context.get(
                "source_net_product_sales", sale.selling_price * sale.quantity
            ) * rate
            _add_period_value(
                auto_values["product_sales_revenue"],
                sale.sold_at.date(),
                revenue,
                periods,
            )

    if sales:
        sale_ids = [sale.id for sale in sales]
        cost_rows = (await db.execute(
            select(SaleCostDetail, InventoryBatch, Sale.sold_at)
            .join(Sale, Sale.id == SaleCostDetail.sale_id)
            .join(InventoryBatch, InventoryBatch.id == SaleCostDetail.batch_id)
            .where(SaleCostDetail.sale_id.in_(sale_ids))
        )).all()
        for detail, batch, sold_at in cost_rows:
            quantity = Decimal(detail.quantity)
            _add_period_value(
                auto_values["product_purchase_cost"], sold_at.date(),
                batch.purchase_price * quantity, periods,
            )
            divisor = Decimal(batch.quantity or 1)
            _add_period_value(
                auto_values["head_logistics_cost"], sold_at.date(),
                (batch.shipping_cost + batch.last_mile_cost) * quantity / divisor, periods,
            )
        for sale in sales:
            context = import_context_by_sale.get(sale.id, {})
            rate = context.get("exchange_rate_to_cny", Decimal("1"))
            if rate is None:
                continue
            other_expense = context.get("source_other_expense", sale.platform_fee) * rate
            _add_period_value(
                auto_values["other_direct_cost"],
                sale.sold_at.date(),
                other_expense,
                periods,
            )

    previous_month_end = _previous_month_end(report_month.replace(day=1))
    month_before_end = _previous_month_end(previous_month_end.replace(day=1))
    current_impairment = await _store_inventory_impairment_total(
        db, store_id, datetime.combine(_next_month_start(report_month.replace(day=1)), time.min)
    )
    previous_impairment = await _store_inventory_impairment_total(
        db, store_id, month_start_cutoff(previous_month_end)
    )
    month_before_impairment = await _store_inventory_impairment_total(
        db, store_id, month_start_cutoff(month_before_end)
    )
    year_end_previous = date(report_month.year - 1, 12, 31)
    year_end_impairment = await _store_inventory_impairment_total(
        db, store_id, month_start_cutoff(year_end_previous)
    )
    auto_values["inventory_impairment"] = {
        "current": current_impairment - previous_impairment,
        "previous": previous_impairment - month_before_impairment,
        "ytd": current_impairment - year_end_impairment,
    }
    active_store_count = await db.scalar(
        select(func.count()).select_from(Store).where(Store.is_active.is_(True))
    )
    if active_store_count:
        office_totals = await office_space_totals_by_application_date(ytd_start, ytd_end)
        for application_date, total in office_totals.items():
            _add_period_value(
                auto_values["rent_utilities"],
                application_date,
                total / Decimal(active_store_count),
                periods,
            )
    if store_name:
        salary_totals = await china_salary_totals_by_store_and_application_date(
            ytd_start, ytd_end
        )
        auto_values["salary"] = {
            period: Decimal("0") for period in STORE_PROFIT_LOSS_PERIODS
        }
        for application_date, total in salary_totals.get(store_name.strip(), {}).items():
            _add_period_value(
                auto_values["salary"], application_date, total, periods
            )
    return auto_values


def month_start_cutoff(value: date) -> datetime:
    return datetime.combine(value + timedelta(days=1), time.min)


async def _store_inventory_impairment_total(
    db: AsyncSession, store_id: int, cutoff: datetime
) -> Decimal:
    batch_rows = (await db.execute(
        select(
            InventoryBatch,
            User.store_id,
            Product.product_type,
            Product.safe_stock_quantity,
        )
        .join(Product, Product.id == InventoryBatch.product_id)
        .outerjoin(User, User.id == InventoryBatch.user_id)
        .where(InventoryBatch.arrived_at < cutoff)
    )).all()
    if not batch_rows:
        return Decimal("0")
    batch_ids = [batch.id for batch, *_ in batch_rows]
    product_ids = {batch.product_id for batch, *_ in batch_rows}
    rule_rows = (await db.execute(
        select(ProductImpairmentRule).where(
            ProductImpairmentRule.product_id.in_(product_ids)
        )
    )).scalars().all()
    rules_by_product: defaultdict[int, list[ImpairmentRule]] = defaultdict(list)
    for rule in rule_rows:
        rules_by_product[rule.product_id].append(ImpairmentRule(
            product_type=rule.product_type,
            safe_stock_quantity=rule.safe_stock_quantity,
            effective_date=rule.effective_date,
        ))
    deductions = (await db.execute(
        select(
            SaleCostDetail.batch_id,
            func.coalesce(func.sum(SaleCostDetail.quantity), 0),
        )
        .join(Sale, Sale.id == SaleCostDetail.sale_id)
        .where(
            SaleCostDetail.batch_id.in_(batch_ids),
            Sale.sold_at < cutoff,
        )
        .group_by(SaleCostDetail.batch_id)
    )).all()
    deducted_by_batch = {row[0]: int(row[1] or 0) for row in deductions}
    as_of = (cutoff - timedelta(days=1)).date()
    layers = [
        InventoryLayer(
            batch_id=batch.id,
            product_id=batch.product_id,
            store_id=batch_store_id,
            arrived_at=batch.arrived_at.date(),
            quantity=max(0, batch.quantity - deducted_by_batch.get(batch.id, 0)),
            product_type=product_type,
            safe_stock_quantity=safe_stock_quantity,
            rules=tuple(rules_by_product[batch.product_id]),
        )
        for batch, batch_store_id, product_type, safe_stock_quantity in batch_rows
    ]
    impairments = batch_impairments(layers, as_of)
    total = Decimal("0")
    for batch, batch_store_id, _product_type, _safe_stock_quantity in batch_rows:
        if batch_store_id != store_id:
            continue
        impairment = impairments.get(batch.id)
        if impairment:
            total += batch.unit_cost * impairment.impairment_units
    return total


async def _resolve_report_store(
    db: AsyncSession, user: User, store_id: int | None
) -> Store:
    if user.role == "operator":
        if user.store_id is None:
            raise HTTPException(status_code=400, detail="当前账号未绑定店铺，请联系管理员")
        if store_id is not None and store_id != user.store_id:
            raise HTTPException(status_code=403, detail="运营账号只能查看绑定店铺")
        store_id = user.store_id
    if store_id is None:
        store_id = await db.scalar(
            select(Store.id).where(Store.is_active.is_(True)).order_by(Store.id).limit(1)
        )
    if store_id is None:
        raise HTTPException(status_code=404, detail="暂无可用店铺")
    store = await db.get(Store, store_id)
    if not store:
        raise HTTPException(status_code=404, detail="店铺不存在")
    return store


async def _store_profit_loss_response(
    db: AsyncSession, user: User, store: Store, report_month: date,
    report: StoreProfitLossReport | None,
) -> StoreProfitLossResponse:
    auto_values = await _store_profit_loss_auto_values(
        db, store.id, report_month, store.name
    )
    manual_values = report.manual_values if report else {}
    values = _build_store_profit_loss_values(auto_values, manual_values)
    remarks = report.remarks if report else {}
    rows = [
        StoreProfitLossRowResponse(
            key=key,
            label=label,
            kind=kind,
            is_auto=kind == "auto",
            is_formula=kind == "formula",
            source=source,
            value=StoreProfitLossCell(**values[key]),
            remark=str(remarks.get(key, "")),
        )
        for key, label, kind, source in STORE_PROFIT_LOSS_ROWS
    ]
    return StoreProfitLossResponse(
        store_id=store.id,
        store_name=store.name,
        store_platform=store.platform,
        report_month=report_month.strftime("%Y-%m"),
        period_label=f"{report_month.year}年{report_month.month}月",
        can_edit=user.role in ("admin", "operator"),
        updated_at=report.updated_at if report else None,
        rows=rows,
    )


@router.get("/store-profit-loss", response_model=StoreProfitLossResponse)
async def get_store_profit_loss(
    report_month: date = Query(..., description="核算月份，传入该月第一天"),
    store_id: int | None = Query(default=None, ge=1),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """获取单店单月损益表，自动项由系统数据生成，其他项保留人工输入。"""
    if report_month.day != 1:
        raise HTTPException(status_code=400, detail="核算月份必须传入该月第一天")
    store = await _resolve_report_store(db, user, store_id)
    report = await db.scalar(
        select(StoreProfitLossReport).where(
            StoreProfitLossReport.store_id == store.id,
            StoreProfitLossReport.report_month == report_month,
        )
    )
    return await _store_profit_loss_response(db, user, store, report_month, report)


@router.put("/store-profit-loss", response_model=StoreProfitLossResponse)
async def update_store_profit_loss(
    data: StoreProfitLossUpdateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    """保存单店单月损益表中的人工输入。"""
    if data.report_month.day != 1:
        raise HTTPException(status_code=400, detail="核算月份必须传入该月第一天")
    store = await _resolve_report_store(db, user, data.store_id)
    manual_values = {}
    for key, cell in data.manual_values.items():
        if key not in STORE_PROFIT_LOSS_MANUAL_KEYS:
            continue
        manual_values[key] = {
            period: (
                None
                if getattr(cell, period) is None
                else str(getattr(cell, period))
            )
            for period in STORE_PROFIT_LOSS_PERIODS
        }
    remarks = {
        key: str(value).strip()[:500]
        for key, value in data.remarks.items()
        if key in STORE_PROFIT_LOSS_ROW_KEYS and str(value).strip()
    }
    report = await db.scalar(
        select(StoreProfitLossReport).where(
            StoreProfitLossReport.store_id == store.id,
            StoreProfitLossReport.report_month == data.report_month,
        )
    )
    if report is None:
        report = StoreProfitLossReport(
            store_id=store.id,
            report_month=data.report_month,
            manual_values=manual_values,
            remarks=remarks,
            updated_by_user_id=user.id,
        )
        db.add(report)
    else:
        report.manual_values = manual_values
        report.remarks = remarks
        report.updated_by_user_id = user.id
    await db.commit()
    await db.refresh(report)
    return await _store_profit_loss_response(db, user, store, data.report_month, report)


def _build_inventory_impairment_workbook(
    rows: list[dict], report_date: date
) -> tuple[Workbook, dict[str, int | Decimal | str]]:
    previous_month_end = _previous_month_end(report_date)
    report_date_formula = _excel_date(report_date)
    previous_month_end_formula = _excel_date(previous_month_end)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INVENTORY_IMPAIRMENT_SHEET
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.append(INVENTORY_IMPAIRMENT_HEADERS)

    header_fill = PatternFill("solid", fgColor="C00000")
    thin_gray = Side(style="thin", color="BFBFBF")
    table_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="等线", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = table_border
    sheet.row_dimensions[1].height = 32
    formula_cache: dict[str, int | Decimal | str] = {}

    for row_number, item in enumerate(rows, start=2):
        age_days = max((report_date - item["arrived_at"]).days, 0)
        product_type = item.get("product_type", "stable")
        provision_quantity = int(item.get("provision_quantity", item["quantity"]))
        current_rate = item.get(
            "impairment_rate",
            impairment_rate(report_date, item["arrived_at"], product_type),
        )
        inventory_amount = Decimal(item["unit_cost"]) * item["quantity"]
        impairment_amount = Decimal(item.get(
            "impairment_amount",
            Decimal(item["unit_cost"]) * provision_quantity * current_rate,
        ))
        book_value = inventory_amount - impairment_amount
        previous_provision_quantity = int(item.get(
            "previous_provision_quantity", item["previous_month_quantity"]
        ))
        previous_rate = item.get(
            "previous_impairment_rate",
            impairment_rate(previous_month_end, item["arrived_at"], product_type),
        )
        previous_impairment_amount = item.get("previous_impairment_amount")
        calculated_previous_book_value = Decimal(item["unit_cost"]) * Decimal(
            item["previous_month_quantity"]
        ) - (
            Decimal(previous_impairment_amount)
            if previous_impairment_amount is not None
            else Decimal(item["unit_cost"]) * Decimal(previous_provision_quantity) * previous_rate
        )
        snapshot_previous_book_value = item.get("previous_book_value")
        previous_book_value = (
            Decimal(snapshot_previous_book_value)
            if snapshot_previous_book_value is not None
            else calculated_previous_book_value
        )
        book_value_difference = book_value - previous_book_value
        status = "安全库存内" if provision_quantity == 0 and item["quantity"] else (
            "达到上限" if current_rate >= Decimal("0.9") else (
                "接近上限" if current_rate >= Decimal("0.6") else "正常计提"
            )
        )
        daily_rate = daily_impairment_rate(report_date, product_type)
        previous_daily_rate = daily_impairment_rate(previous_month_end, product_type)
        uses_historical_amount = "impairment_amount" in item
        uses_snapshot_previous_book_value = snapshot_previous_book_value is not None
        formula_cache.update({
            f"E{row_number}": age_days,
            f"H{row_number}": provision_quantity,
            f"I{row_number}": inventory_amount,
            f"J{row_number}": current_rate,
            f"K{row_number}": impairment_amount,
            f"L{row_number}": book_value,
            f"M{row_number}": book_value_difference,
            f"N{row_number}": status,
        })
        sheet.append([
            item["store_name"],
            item["sku"],
            item["product_name"],
            item["arrived_at"],
            f'=IF(D{row_number}="","",{report_date_formula}-D{row_number})',
            item["unit_cost"],
            item["quantity"],
            provision_quantity,
            f'=IF(OR(F{row_number}="",G{row_number}=""),"",F{row_number}*G{row_number})',
            (
                current_rate if uses_historical_amount
                else f'=IF(E{row_number}="","",MIN(E{row_number}*{daily_rate},0.9))'
            ),
            (
                impairment_amount if uses_historical_amount
                else f'=IF(OR(F{row_number}="",H{row_number}="",J{row_number}=""),"",F{row_number}*H{row_number}*J{row_number})'
            ),
            f'=IF(OR(I{row_number}="",K{row_number}=""),"",I{row_number}-K{row_number})',
            (
                book_value_difference if uses_historical_amount
                or uses_snapshot_previous_book_value
                else (
                    f'=IF(OR(D{row_number}="",F{row_number}="",L{row_number}=""),"",'
                    f'L{row_number}-F{row_number}*({item["previous_month_quantity"]}-'
                    f'{previous_provision_quantity}*MIN(MAX({previous_month_end_formula}-D{row_number},0)*{previous_daily_rate},0.9)))'
                )
            ),
            (
                f'=IF(H{row_number}=0,"安全库存内",IF(J{row_number}>=0.9,"达到上限",'
                f'IF(J{row_number}>=0.6,"接近上限","正常计提")))'
            ),
            (
                f"批次号：{item['batch_no']}；商品类型："
                f"{'新品' if product_type == PRODUCT_TYPE_NEW else '稳健商品'}；"
                f"上月末库存：{item['previous_month_quantity']}；"
                f"上月末账面价值：{previous_book_value:.4f}；"
                "单件成本含采购、头程、尾程及其他成本，未单列关税"
            ),
        ])
        for cell in sheet[row_number]:
            cell.font = Font(name="等线", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = table_border
        sheet.cell(row_number, 4).number_format = "yyyy-mm-dd"
        sheet.cell(row_number, 5).number_format = "0"
        for column in (6, 9, 11, 12, 13):
            sheet.cell(row_number, column).number_format = '¥#,##0.00'
        for column in (7, 8):
            sheet.cell(row_number, column).number_format = "#,##0"
        sheet.cell(row_number, 10).number_format = "0%"

    first_data_row = 2
    last_data_row = len(rows) + 1
    total_row = last_data_row + 1
    sheet.cell(total_row, 1, "合计")
    if rows:
        for column in range(5, 14):
            letter = sheet.cell(1, column).column_letter
            sheet.cell(total_row, column, f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
            formula_cache[f"{letter}{total_row}"] = sum(
                formula_cache[f"{letter}{row_number}"]
                if letter not in ("F", "G")
                else rows[row_number - first_data_row]["unit_cost" if letter == "F" else "quantity"]
                for row_number in range(first_data_row, last_data_row + 1)
            )
    else:
        for column in range(5, 14):
            sheet.cell(total_row, column, 0)
    for cell in sheet[total_row]:
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.font = Font(name="等线", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = table_border
    for column in (6, 9, 11, 12, 13):
        sheet.cell(total_row, column).number_format = '¥#,##0.00'
    sheet.cell(total_row, 5).number_format = "0"
    for column in (7, 8):
        sheet.cell(total_row, column).number_format = "#,##0"
    sheet.cell(total_row, 10).number_format = "0%"

    notes_row = total_row + 2
    notes = [
        f"报告截止日：{report_date.isoformat()}；上月末：{previous_month_end.isoformat()}。",
        "计提规则：2026-09起，稳健商品仅超出安全库存的数量按每日1%计提；新品全部库存按每日0.5%计提；类型变更前已计提金额保留，累计上限90%。",
        "安全库存口径：按商品全仓汇总，优先豁免最新入库批次；此前月份仍按每日1%计提全部库存。",
        "库存口径：当前库存按截止日前已确认销售计算；上月末使用系统自动生成的月末快照。",
        "成本口径：单件到仓成本使用系统批次单件成本，包含采购、头程、尾程及其他成本；当前未单列关税。",
        "店铺口径：按入库批次创建人的当前绑定店铺；无法确认时显示“未关联店铺”。",
    ]
    for offset, note in enumerate(notes):
        row_number = notes_row + offset
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=15)
        cell = sheet.cell(row_number, 1, note)
        cell.font = Font(name="等线", size=9, color="666666")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = {
        "A": 18, "B": 18, "C": 28, "D": 13, "E": 12, "F": 18, "G": 12,
        "H": 12, "I": 18, "J": 14, "K": 20, "L": 17, "M": 25, "N": 13, "O": 62,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:O{max(1, last_data_row)}"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:O{notes_row + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook, formula_cache


def _workbook_bytes_with_formula_cache(
    workbook: Workbook, formula_cache: dict[str, int | Decimal | str]
) -> BytesIO:
    """Save formulas and their last-calculated values for Excel Protected View."""
    raw_output = BytesIO()
    workbook.save(raw_output)
    raw_output.seek(0)

    spreadsheet_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ElementTree.register_namespace("", spreadsheet_namespace)
    output = BytesIO()
    with ZipFile(raw_output, "r") as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                for cell in root.findall(f".//{{{spreadsheet_namespace}}}c"):
                    coordinate = cell.attrib.get("r")
                    if coordinate not in formula_cache:
                        continue
                    formula = cell.find(f"{{{spreadsheet_namespace}}}f")
                    if formula is None:
                        continue
                    cached_value = formula_cache[coordinate]
                    value = cell.find(f"{{{spreadsheet_namespace}}}v")
                    if value is None:
                        value = ElementTree.SubElement(cell, f"{{{spreadsheet_namespace}}}v")
                    if isinstance(cached_value, str):
                        cell.set("t", "str")
                        value.text = cached_value
                    else:
                        cell.attrib.pop("t", None)
                        value.text = (
                            format(cached_value, "f")
                            if isinstance(cached_value, Decimal)
                            else str(cached_value)
                        )
                content = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(member, content)
    output.seek(0)
    return output


FINANCIAL_ATTACHMENT_SHEET = "09-店铺损益表"
STORE_PROFIT_LOSS_SUMMARY_SHEET = "10-各店铺损益汇总表"
PRODUCT_LINE_PROFIT_SHEET = "12-各产品线毛利分析表"


def _excel_number(value: Decimal | int | float | None):
    if value is None:
        return None
    return float(value)


def _style_financial_sheet(sheet, headers: list[str], fill_color: str = "1F4E78") -> None:
    header_fill = PatternFill("solid", fgColor=fill_color)
    thin_gray = Side(style="thin", color="D9E1F2")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="等线", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            cell.border = border


def _format_profit_loss_sheet(sheet, first_data_row: int, last_data_row: int) -> None:
    for row_number in range(first_data_row, last_data_row + 1):
        key = STORE_PROFIT_LOSS_ROWS[row_number - first_data_row][0]
        for column in range(2, 5):
            sheet.cell(row_number, column).number_format = "0.00%" if key == "gross_margin" else '¥#,##0.00'
        kind = STORE_PROFIT_LOSS_ROWS[row_number - first_data_row][2]
        if kind == "formula":
            for cell in sheet[row_number]:
                cell.fill = PatternFill("solid", fgColor="EEF3F8")
                cell.font = Font(name="等线", size=10, bold=True)
        elif kind == "auto":
            for cell in sheet[row_number][1:4]:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")


def _build_store_profit_loss_sheet(
    workbook: Workbook,
    store: Store,
    report_month: date,
    values: dict[str, dict[str, Decimal | None]],
    remarks: dict,
) -> None:
    sheet = workbook.create_sheet(FINANCIAL_ATTACHMENT_SHEET)
    headers = ["项目", "本月金额", "上月金额", "本年累计", "备注"]
    sheet.append(headers)
    for key, label, _kind, _source in STORE_PROFIT_LOSS_ROWS:
        sheet.append([
            label,
            _excel_number(values[key].get("current")),
            _excel_number(values[key].get("previous")),
            _excel_number(values[key].get("ytd")),
            str(remarks.get(key, "")) if isinstance(remarks, dict) else "",
        ])
    first_data_row = 2
    last_data_row = 1 + len(STORE_PROFIT_LOSS_ROWS)
    _style_financial_sheet(sheet, headers)
    _format_profit_loss_sheet(sheet, first_data_row, last_data_row)
    metadata = [
        ["店铺", store.name, None, None, None],
        ["所属平台", store.platform, None, None, None],
        ["核算期间", f"{report_month.year}年{report_month.month}月", None, None, None],
        ["数据来源", "与系统中的店铺损益表查询结果一致；人工项目保留已保存值。", None, None, None],
    ]
    for row in metadata:
        sheet.append(row)
    for row_number in range(last_data_row + 1, sheet.max_row + 1):
        sheet.cell(row_number, 1).font = Font(name="等线", size=10, bold=row_number == last_data_row + 1)
        sheet.cell(row_number, 2).font = Font(name="等线", size=10, color="666666")
        for column in range(1, 6):
            sheet.cell(row_number, column).alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"A": 28, "B": 16, "C": 16, "D": 16, "E": 42}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:E{last_data_row}"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:E{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _build_store_profit_loss_summary_sheet(
    workbook: Workbook,
    store_reports: list[tuple[Store, dict[str, dict[str, Decimal | None]]]],
    report_month: date,
) -> None:
    sheet = workbook.create_sheet(STORE_PROFIT_LOSS_SUMMARY_SHEET)
    headers = ["项目"] + [store.name for store, _values in store_reports] + ["合计", "占比"]
    sheet.append(headers)
    revenue_total = sum(
        (_decimal_or_zero(values["operating_revenue"].get("current")) for _store, values in store_reports),
        Decimal("0"),
    )
    for key, label, _kind, _source in STORE_PROFIT_LOSS_ROWS:
        store_values = [
            _decimal_or_zero(values[key].get("current"))
            for _store, values in store_reports
        ]
        if key == "gross_margin":
            total = (
                sum((_decimal_or_zero(values["gross_profit"].get("current")) for _store, values in store_reports), Decimal("0"))
                / revenue_total
                if revenue_total else None
            )
            ratio = None
        else:
            total = sum(store_values, Decimal("0"))
            ratio = total / revenue_total if revenue_total else None
        sheet.append([
            label,
            *[_excel_number(value) for value in store_values],
            _excel_number(total),
            _excel_number(ratio),
        ])
    first_data_row = 2
    last_data_row = 1 + len(STORE_PROFIT_LOSS_ROWS)
    _style_financial_sheet(sheet, headers)
    for row_number in range(first_data_row, last_data_row + 1):
        key = STORE_PROFIT_LOSS_ROWS[row_number - first_data_row][0]
        for column in range(2, len(headers)):
            sheet.cell(row_number, column).number_format = "0.00%" if key == "gross_margin" or column == len(headers) else '¥#,##0.00'
        if key == "gross_margin":
            sheet.cell(row_number, len(headers)).number_format = "General"
        for cell in sheet[row_number]:
            if STORE_PROFIT_LOSS_ROWS[row_number - first_data_row][2] == "formula":
                cell.fill = PatternFill("solid", fgColor="EEF3F8")
                cell.font = Font(name="等线", size=10, bold=True)
    sheet.append(["店铺数量", len(store_reports)] + [None] * (len(headers) - 2))
    sheet.append(["核算期间", f"{report_month.year}年{report_month.month}月"] + [None] * (len(headers) - 2))
    sheet.append(["占比口径", "各项目合计 ÷ 合计营业收入；毛利率按合计毛利 ÷ 合计营业收入计算。", *([None] * (len(headers) - 2))])
    for row_number in range(last_data_row + 1, sheet.max_row + 1):
        for column in range(1, len(headers) + 1):
            sheet.cell(row_number, column).alignment = Alignment(vertical="center", wrap_text=True)
    sheet.column_dimensions["A"].width = 28
    for column in range(2, len(headers) + 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 16
    sheet.column_dimensions[sheet.cell(1, len(headers)).column_letter].width = 14
    sheet.auto_filter.ref = f"A1:{sheet.cell(last_data_row, len(headers)).coordinate}"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:{sheet.cell(sheet.max_row, len(headers)).coordinate}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _mapping_effective_on(mapping: StoreProduct, as_of: date) -> bool:
    if not mapping.is_active or mapping.effective_from > as_of:
        return False
    return mapping.effective_to is None or as_of <= mapping.effective_to


def _lookup_key(value) -> str:
    return " ".join(str(value or "").casefold().split())


def _product_line_mapping_for_sale(
    mappings_by_store_product: dict[tuple[int, int], list[tuple[StoreProduct, ProductLine, Product]]],
    store_id: int | None,
    product_id: int,
    as_of: date,
    source_skus: set[str] | None = None,
) -> tuple[StoreProduct, ProductLine, Product] | None:
    if store_id is None:
        return None
    candidates = [
        item for item in mappings_by_store_product.get((store_id, product_id), [])
        if _mapping_effective_on(item[0], as_of)
    ]
    if not source_skus:
        return None
    normalized_source_skus = {_lookup_key(value) for value in source_skus if value}
    if not normalized_source_skus:
        return None
    matched = [
        item for item in candidates
        if _lookup_key(item[0].platform_sku_id) in normalized_source_skus
        or _lookup_key(item[0].store_sku) in normalized_source_skus
    ]
    if not matched:
        return None
    candidates = matched
    product_line_ids = {item[1].id for item in candidates}
    if len(product_line_ids) != 1:
        return None
    return max(candidates, key=lambda item: item[0].effective_from) if candidates else None


async def _product_line_profit_data(
    db: AsyncSession,
    store_ids: set[int],
    report_month: date,
    report_date: date,
) -> list[dict]:
    periods = _month_periods(report_month)
    ytd_start, ytd_end = periods["ytd"]
    mapping_rows = (await db.execute(
        select(StoreProduct, ProductLine, Product)
        .join(ProductLine, ProductLine.id == StoreProduct.product_line_id)
        .join(Product, Product.id == StoreProduct.product_id)
        .where(StoreProduct.store_id.in_(store_ids))
    )).all() if store_ids else []
    mappings_by_store_product: dict[tuple[int, int], list[tuple[StoreProduct, ProductLine, Product]]] = defaultdict(list)
    line_data: dict[int | None, dict] = {}
    for mapping, product_line, product in mapping_rows:
        mappings_by_store_product[(mapping.store_id, mapping.product_id)].append((mapping, product_line, product))
        if _mapping_effective_on(mapping, report_date):
            item = line_data.setdefault(product_line.id, {
                "name": product_line.name,
                "sku_ids": set(),
                "current": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "ytd": {"revenue": Decimal("0"), "cost": Decimal("0")},
                "note": "",
            })
            item["sku_ids"].add(product.id)

    sales = list((await db.execute(
        select(Sale, User.store_id)
        .outerjoin(User, User.id == Sale.user_id)
        .where(
            Sale.sold_at >= datetime.combine(ytd_start, time.min),
            Sale.sold_at < datetime.combine(ytd_end, time.min),
            or_(
                Sale.store_id.in_(store_ids),
                and_(Sale.store_id.is_(None), User.store_id.in_(store_ids)),
            ),
        )
    )).all()) if store_ids else []
    import_context_by_sale: dict[int, dict] = {}
    if sales:
        from routers.sales import _sale_import_context

        import_context_by_sale, _ = await _sale_import_context(db, [sale for sale, _user_store_id in sales])

    sale_ids = [sale.id for sale, _user_store_id in sales]
    source_skus_by_sale: defaultdict[int, set[str]] = defaultdict(set)
    if sale_ids:
        source_sku_rows = await db.execute(
            select(SettlementEntryAllocation.sale_id, SettlementEntry.platform_sku_id)
            .join(
                SettlementEntry,
                SettlementEntry.id == SettlementEntryAllocation.settlement_entry_id,
            )
            .where(
                SettlementEntryAllocation.sale_id.in_(sale_ids),
                SettlementEntry.platform_sku_id.is_not(None),
            )
        )
        for sale_id, platform_sku_id in source_sku_rows.all():
            if platform_sku_id:
                source_skus_by_sale[sale_id].add(platform_sku_id)

    def line_for(sale: Sale, user_store_id: int | None):
        store_id = sale.store_id or user_store_id
        mapping = _product_line_mapping_for_sale(
            mappings_by_store_product,
            store_id,
            sale.product_id,
            sale.sold_at.date(),
            source_skus_by_sale.get(sale.id),
        )
        if mapping is None:
            if source_skus_by_sale.get(sale.id):
                return None, "未分配产品线（来源 SKU 未匹配或产品线映射冲突）"
            return None, "未分配产品线（缺少来源 SKU 或产品线映射冲突）"
        return mapping[1].id, mapping[1].name

    sale_line_by_id: dict[int, int | None] = {}
    for sale, user_store_id in sales:
        line_id, line_name = line_for(sale, user_store_id)
        item = line_data.setdefault(line_id, {
            "name": line_name,
            "sku_ids": set(),
            "current": {"revenue": Decimal("0"), "cost": Decimal("0")},
            "previous": {"revenue": Decimal("0"), "cost": Decimal("0")},
            "ytd": {"revenue": Decimal("0"), "cost": Decimal("0")},
            "note": "销售记录未匹配到有效店铺 SKU 产品线映射" if line_id is None else "",
        })
        item["sku_ids"].add(sale.product_id)
        sale_line_by_id[sale.id] = line_id
        context = import_context_by_sale.get(sale.id, {})
        rate = context.get("exchange_rate_to_cny", Decimal("1"))
        if rate is None:
            warning = "销售收入因汇率缺失未计入；FIFO成本仍按销售成本明细计入。"
            item["note"] = f"{item['note']}；{warning}".strip("；")
            continue
        revenue = context.get("source_net_product_sales", sale.selling_price * sale.quantity) * rate
        for period in STORE_PROFIT_LOSS_PERIODS:
            start, end = periods[period]
            if start <= sale.sold_at.date() < end:
                item[period]["revenue"] += revenue
        other_expense = context.get("source_other_expense", sale.platform_fee) * rate
        for period in STORE_PROFIT_LOSS_PERIODS:
            start, end = periods[period]
            if start <= sale.sold_at.date() < end:
                item[period]["cost"] += other_expense

    if sales:
        cost_rows = (await db.execute(
            select(SaleCostDetail, InventoryBatch, Sale)
            .join(Sale, Sale.id == SaleCostDetail.sale_id)
            .join(InventoryBatch, InventoryBatch.id == SaleCostDetail.batch_id)
            .where(SaleCostDetail.sale_id.in_([sale.id for sale, _user_store_id in sales]))
        )).all()
        for detail, batch, sale in cost_rows:
            line_id = sale_line_by_id.get(sale.id)
            if line_id not in line_data:
                continue
            quantity = Decimal(detail.quantity)
            batch_cost = batch.purchase_price * quantity
            divisor = Decimal(batch.quantity or 1)
            batch_cost += (batch.shipping_cost + batch.last_mile_cost) * quantity / divisor
            for period in STORE_PROFIT_LOSS_PERIODS:
                start, end = periods[period]
                if start <= sale.sold_at.date() < end:
                    line_data[line_id][period]["cost"] += batch_cost

    rows = []
    for item in sorted(line_data.values(), key=lambda value: value["name"]):
        rows.append({
            "name": item["name"],
            "sku_count": len(item["sku_ids"]),
            "sku_ids": item["sku_ids"],
            "current": item["current"],
            "previous": item["previous"],
            "ytd": item["ytd"],
            "note": item["note"],
        })
    return rows


def _build_product_line_profit_sheet(
    workbook: Workbook,
    rows: list[dict],
    report_month: date,
) -> None:
    sheet = workbook.create_sheet(PRODUCT_LINE_PROFIT_SHEET)
    headers = ["产品线/品类", "SKU数量", "本月营业收入", "本月营业成本", "本月毛利", "本月毛利率", "上月毛利率", "环比变化", "本年累计收入", "本年累计毛利", "累计毛利率", "备注"]
    sheet.append(headers)
    totals = {
        period: {"revenue": sum((row[period]["revenue"] for row in rows), Decimal("0")), "cost": sum((row[period]["cost"] for row in rows), Decimal("0"))}
        for period in ("current", "previous", "ytd")
    }
    for row in rows:
        current_gross = row["current"]["revenue"] - row["current"]["cost"]
        previous_gross = row["previous"]["revenue"] - row["previous"]["cost"]
        ytd_gross = row["ytd"]["revenue"] - row["ytd"]["cost"]
        current_margin = current_gross / row["current"]["revenue"] if row["current"]["revenue"] else None
        previous_margin = previous_gross / row["previous"]["revenue"] if row["previous"]["revenue"] else None
        ytd_margin = ytd_gross / row["ytd"]["revenue"] if row["ytd"]["revenue"] else None
        sheet.append([
            row["name"], row["sku_count"], _excel_number(row["current"]["revenue"]), _excel_number(row["current"]["cost"]),
            _excel_number(current_gross), _excel_number(current_margin), _excel_number(previous_margin),
            _excel_number(current_margin - previous_margin if current_margin is not None and previous_margin is not None else None),
            _excel_number(row["ytd"]["revenue"]), _excel_number(ytd_gross), _excel_number(ytd_margin), row["note"],
        ])
    total_current_gross = totals["current"]["revenue"] - totals["current"]["cost"]
    total_previous_gross = totals["previous"]["revenue"] - totals["previous"]["cost"]
    total_ytd_gross = totals["ytd"]["revenue"] - totals["ytd"]["cost"]
    total_current_margin = total_current_gross / totals["current"]["revenue"] if totals["current"]["revenue"] else None
    total_previous_margin = total_previous_gross / totals["previous"]["revenue"] if totals["previous"]["revenue"] else None
    total_ytd_margin = total_ytd_gross / totals["ytd"]["revenue"] if totals["ytd"]["revenue"] else None
    sheet.append([
        "合计", len(set().union(*(row["sku_ids"] for row in rows))) if rows else 0,
        _excel_number(totals["current"]["revenue"]), _excel_number(totals["current"]["cost"]), _excel_number(total_current_gross),
        _excel_number(total_current_margin), _excel_number(total_previous_margin),
        _excel_number(total_current_margin - total_previous_margin if total_current_margin is not None and total_previous_margin is not None else None),
        _excel_number(totals["ytd"]["revenue"]), _excel_number(total_ytd_gross), _excel_number(total_ytd_margin), "",
    ])
    _style_financial_sheet(sheet, headers)
    for row_number in range(2, sheet.max_row + 1):
        for column in (3, 4, 5, 9, 10):
            sheet.cell(row_number, column).number_format = '¥#,##0.00'
        for column in (6, 7, 8, 11):
            sheet.cell(row_number, column).number_format = "0.00%"
        if row_number == sheet.max_row:
            for cell in sheet[row_number]:
                cell.font = Font(name="等线", size=10, bold=True)
                cell.fill = PatternFill("solid", fgColor="EEF3F8")
    note_row = sheet.max_row + 2
    sheet.cell(note_row, 1, "说明")
    sheet.cell(note_row, 2, f"核算期间：{report_month.year}年{report_month.month}月；收入沿用销售记录口径，成本包含可按销售/FIFO明细归属的采购成本、头程尾程物流及销售关联费用。")
    sheet.cell(note_row + 1, 2, "店铺损益表中人工录入的税费、本地物流、代发等未按产品线分摊；未匹配产品线的销售单独列示。")
    for row_number in (note_row, note_row + 1):
        sheet.cell(row_number, 1).font = Font(name="等线", size=10, bold=row_number == note_row)
        sheet.cell(row_number, 2).font = Font(name="等线", size=9, color="666666")
        sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="center")
    widths = {"A": 24, "B": 12, "C": 16, "D": 16, "E": 16, "F": 14, "G": 14, "H": 14, "I": 16, "J": 16, "K": 14, "L": 48}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = f"A1:L{sheet.max_row - 2}"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:L{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


async def _load_inventory_impairment_report_rows(
    db: AsyncSession, report_date: date
) -> list[dict]:
    report_cutoff = datetime.combine(report_date + timedelta(days=1), time.min)
    previous_month_end = _previous_month_end(report_date)
    snapshot_period = await ensure_period_snapshot(db, previous_month_end)
    snapshots_by_batch = {
        snapshot.batch_id: snapshot
        for snapshot in (
            await db.execute(
                select(InventoryPeriodSnapshot).where(
                    InventoryPeriodSnapshot.accounting_period_id == snapshot_period.id
                )
            )
        ).scalars().all()
    }

    batch_result = await db.execute(
        select(
            InventoryBatch,
            Product.sku,
            Product.name.label("product_name"),
            Store.name.label("store_name"),
            Product.product_type,
            Product.safe_stock_quantity,
            User.store_id,
        )
        .join(Product, Product.id == InventoryBatch.product_id)
        .outerjoin(User, User.id == InventoryBatch.user_id)
        .outerjoin(Store, Store.id == User.store_id)
        .where(InventoryBatch.arrived_at < report_cutoff)
        .order_by(Store.name, Product.sku, InventoryBatch.arrived_at, InventoryBatch.id)
    )
    batch_rows = batch_result.all()
    batch_ids = [row[0].id for row in batch_rows]
    product_ids = {row[0].product_id for row in batch_rows}
    rule_rows = (await db.execute(
        select(ProductImpairmentRule).where(
            ProductImpairmentRule.product_id.in_(product_ids)
        )
    )).scalars().all() if product_ids else []
    rules_by_product: defaultdict[int, list[ImpairmentRule]] = defaultdict(list)
    for rule in rule_rows:
        rules_by_product[rule.product_id].append(ImpairmentRule(
            product_type=rule.product_type,
            safe_stock_quantity=rule.safe_stock_quantity,
            effective_date=rule.effective_date,
        ))

    deducted_by_batch: dict[int, int] = {}
    if batch_ids:
        deduction_result = await db.execute(
            select(
                SaleCostDetail.batch_id,
                func.coalesce(func.sum(case(
                    (Sale.sold_at < report_cutoff, SaleCostDetail.quantity),
                    else_=0,
                )), 0).label("report_quantity"),
            )
            .join(Sale, Sale.id == SaleCostDetail.sale_id)
            .where(SaleCostDetail.batch_id.in_(batch_ids))
            .group_by(SaleCostDetail.batch_id)
        )
        deducted_by_batch = {
            row.batch_id: int(row.report_quantity)
            for row in deduction_result.all()
        }

    report_items = []
    for (
        batch,
        sku,
        product_name,
        store_name,
        product_type,
        safe_stock_quantity,
        store_id,
    ) in batch_rows:
        report_deducted = deducted_by_batch.get(batch.id, 0)
        quantity = max(0, batch.quantity - report_deducted)
        previous_snapshot = snapshots_by_batch.get(batch.id)
        previous_month_quantity = int(previous_snapshot.quantity) if previous_snapshot else 0
        previous_book_value = (
            Decimal(previous_snapshot.book_value) if previous_snapshot else None
        )
        if quantity == 0 and previous_month_quantity == 0:
            continue
        report_items.append({
            "batch": batch,
            "store_name": store_name or "未关联店铺",
            "sku": sku,
            "product_name": product_name,
            "store_id": store_id,
            "product_type": product_type,
            "safe_stock_quantity": safe_stock_quantity,
            "arrived_at": batch.arrived_at.date(),
            "unit_cost": batch.unit_cost,
            "quantity": quantity,
            "previous_month_quantity": previous_month_quantity,
            "previous_book_value": previous_book_value,
            "batch_no": batch.batch_no,
        })

    current_impairments = batch_impairments([
        InventoryLayer(
            batch_id=item["batch"].id,
            product_id=item["batch"].product_id,
            store_id=item["store_id"],
            arrived_at=item["arrived_at"],
            quantity=item["quantity"],
            product_type=item["product_type"],
            safe_stock_quantity=item["safe_stock_quantity"],
            rules=tuple(rules_by_product[item["batch"].product_id]),
        )
        for item in report_items
    ], report_date)
    return [
        {
            key: value for key, value in item.items() if key != "batch"
        } | {
            "provision_quantity": (
                current_impairments[item["batch"].id].provision_quantity
                if item["batch"].id in current_impairments else 0
            ),
            "impairment_rate": (
                current_impairments[item["batch"].id].rate
                if item["batch"].id in current_impairments else Decimal("0")
            ),
            "impairment_amount": (
                item["unit_cost"] * current_impairments[item["batch"].id].impairment_units
                if item["batch"].id in current_impairments else Decimal("0")
            ),
        }
        for item in report_items
    ]


@router.get("/inventory-impairment/export")
async def export_inventory_impairment_report(
    report_date: date = Query(..., description="报告截止日"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """按指定截止日导出存货跌价准备明细表。"""
    if report_date > date.today():
        raise HTTPException(status_code=400, detail="报告截止日不能晚于当天")
    report_rows = await _load_inventory_impairment_report_rows(db, report_date)
    workbook, formula_cache = _build_inventory_impairment_workbook(report_rows, report_date)
    output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
    filename = f"inventory_impairment_{report_date.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/financial-attachments/export")
async def export_financial_attachments(
    report_date: date = Query(..., description="存货报表截止日"),
    report_month: date | None = Query(default=None, description="损益报表月份第一天"),
    store_id: int | None = Query(default=None, ge=1),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """导出财务附件。第三张表只对管理员开放，并汇总所有启用店铺。"""
    if report_date > date.today():
        raise HTTPException(status_code=400, detail="报告截止日不能晚于当天")
    report_month = report_month or report_date.replace(day=1)
    if report_month.day != 1:
        raise HTTPException(status_code=400, detail="核算月份必须传入该月第一天")

    selected_store = await _resolve_report_store(db, user, store_id)
    report = await db.scalar(
        select(StoreProfitLossReport).where(
            StoreProfitLossReport.store_id == selected_store.id,
            StoreProfitLossReport.report_month == report_month,
        )
    )
    selected_auto_values = await _store_profit_loss_auto_values(
        db, selected_store.id, report_month, selected_store.name
    )
    selected_values = _build_store_profit_loss_values(
        selected_auto_values, report.manual_values if report else {}
    )
    selected_remarks = report.remarks if report else {}

    inventory_rows = await _load_inventory_impairment_report_rows(db, report_date)
    workbook, formula_cache = _build_inventory_impairment_workbook(inventory_rows, report_date)
    _build_store_profit_loss_sheet(
        workbook, selected_store, report_month, selected_values, selected_remarks
    )

    if user.role == "admin":
        stores = list((await db.execute(
            select(Store).where(Store.is_active.is_(True)).order_by(Store.name, Store.id)
        )).scalars().all())
        reports_by_store = {
            item.store_id: item
            for item in (await db.execute(
                select(StoreProfitLossReport).where(
                    StoreProfitLossReport.report_month == report_month,
                    StoreProfitLossReport.store_id.in_([store.id for store in stores]),
                )
            )).scalars().all()
        } if stores else {}
        store_reports = []
        for store in stores:
            store_report = reports_by_store.get(store.id)
            auto_values = await _store_profit_loss_auto_values(
                db, store.id, report_month, store.name
            )
            store_reports.append((
                store,
                _build_store_profit_loss_values(
                    auto_values, store_report.manual_values if store_report else {}
                ),
            ))
        _build_store_profit_loss_summary_sheet(workbook, store_reports, report_month)
        product_line_store_ids = {store.id for store in stores}
    else:
        product_line_store_ids = {selected_store.id}

    product_line_rows = await _product_line_profit_data(
        db, product_line_store_ids, report_month, report_date
    )
    _build_product_line_profit_sheet(workbook, product_line_rows, report_month)

    output = _workbook_bytes_with_formula_cache(workbook, formula_cache)
    filename = f"financial_attachments_{report_date.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/monthly", response_model=list[MonthlyReportItem])
async def monthly_report(
    start_date: date = Query(..., description="开始日期"),
    end_date: date = Query(..., description="结束日期"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """月度报表 - 所有角色可查看，operator只看自己的数据"""
    month_expr = func.DATE_FORMAT(Sale.sold_at, "%Y-%m").label("month")
    start_at = datetime.combine(start_date, time.min)
    end_at = datetime.combine(end_date + timedelta(days=1), time.min)
    stmt = (
        select(
            month_expr,
            Sale.product_id,
            func.sum(Sale.quantity).label("sold_quantity"),
        )
        .where(
            Sale.sold_at >= start_at,
            Sale.sold_at < end_at,
        )
        .group_by(month_expr, Sale.product_id)
        .order_by(month_expr, Sale.product_id)
    )
    if user.role == "operator":
        stmt = stmt.where(Sale.user_id == user.id)

    result = await db.execute(stmt)
    sold_rows = result.all()

    # 库存信息
    inventory_stmt = select(
        InventoryBatch.product_id,
        func.sum(InventoryBatch.remaining_quantity).label("inventory_quantity"),
        func.sum(InventoryBatch.remaining_quantity * InventoryBatch.unit_cost).label("inventory_value"),
    ).group_by(InventoryBatch.product_id)
    inv_result = await db.execute(inventory_stmt)
    inventory_map: dict[int, tuple[int, Decimal]] = {}
    for row in inv_result.all():
        inventory_map[row.product_id] = (
            row.inventory_quantity or 0,
            row.inventory_value or Decimal("0"),
        )

    products_result = await db.execute(select(Product))
    product_map = {p.id: p for p in products_result.scalars().all()}

    reports = []
    for row in sold_rows:
        product = product_map.get(row.product_id)
        if not product:
            continue
        inv_qty, inv_val = inventory_map.get(row.product_id, (0, Decimal("0")))
        reports.append(MonthlyReportItem(
            month=row.month,
            product_id=row.product_id,
            sku=product.sku,
            product_name=product.name,
            sold_quantity=row.sold_quantity or 0,
            inventory_quantity=inv_qty,
            inventory_value=inv_val,
        ))

    return reports
