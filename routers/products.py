import os
import uuid
from datetime import datetime
from io import BytesIO
from decimal import Decimal
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select, func, or_
from sqlalchemy.ext.asyncio import AsyncSession

from database import get_db
from models import InventoryBatch, PlatformSkuComponent, Product, User
from services.inventory_impairment import (
    PRODUCT_TYPE_NEW,
    PRODUCT_TYPE_STABLE,
    VALID_PRODUCT_TYPES,
)
from schemas import ProductOptionResponse, ProductPageResponse, ProductResponse
from auth import RequireAnyRole, get_current_user

router = APIRouter(prefix="/products", tags=["商品管理"])

UPLOAD_DIR = Path(__file__).parent.parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
BASE_DIR = Path(__file__).parent.parent


def _impairment_settings(product_type: str, safe_stock_quantity: int) -> tuple[str, int]:
    if product_type not in VALID_PRODUCT_TYPES:
        raise HTTPException(status_code=422, detail="商品类型必须是稳健商品或新品")
    if safe_stock_quantity < 0:
        raise HTTPException(status_code=422, detail="安全库存不能小于 0")
    return (
        product_type,
        0 if product_type == PRODUCT_TYPE_NEW else safe_stock_quantity,
    )


@router.post("", response_model=ProductResponse)
async def create_product(
    sku: str = Form(...),
    name: str = Form(...),
    product_type: str = Form(default=PRODUCT_TYPE_STABLE),
    safe_stock_quantity: int = Form(default=0),
    image: UploadFile | None = File(None),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """添加商品 - 所有角色可操作。"""
    existing = await db.execute(select(Product).where(Product.sku == sku))
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=400, detail=f"SKU已存在: {sku}")

    image_path = None
    if image and image.filename:
        ext = os.path.splitext(image.filename)[1]
        filename = f"{uuid.uuid4().hex}{ext}"
        filepath = UPLOAD_DIR / filename
        content = await image.read()
        filepath.write_bytes(content)
        image_path = f"/uploads/{filename}"

    product_type, safe_stock_quantity = _impairment_settings(
        product_type, safe_stock_quantity
    )
    product = Product(
        sku=sku,
        name=name,
        image=image_path,
        product_type=product_type,
        safe_stock_quantity=safe_stock_quantity,
    )
    db.add(product)
    await db.commit()
    await db.refresh(product)
    return _enrich_product(product, 0, Decimal("0"), None)


@router.get("/options", response_model=list[ProductOptionResponse])
async def list_product_options(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """Return lightweight product data for selectors and batch lookups."""
    result = await db.execute(select(Product).order_by(Product.id.desc()))
    return result.scalars().all()


def _product_summary_stmt():
    return (
        select(
            Product,
            func.coalesce(func.sum(InventoryBatch.remaining_quantity), 0).label("stock_quantity"),
            func.coalesce(
                func.sum(InventoryBatch.remaining_quantity * InventoryBatch.unit_cost), 0
            ).label("stock_value"),
            func.max(InventoryBatch.arrived_at).label("last_batch_at"),
        )
        .outerjoin(InventoryBatch, InventoryBatch.product_id == Product.id)
        .group_by(
            Product.id,
            Product.sku,
            Product.name,
            Product.image,
            Product.product_type,
            Product.safe_stock_quantity,
            Product.created_at,
        )
        .order_by(Product.id.desc())
    )


@router.get("", response_model=list[ProductResponse] | ProductPageResponse)
async def list_products(
    keyword: str | None = Query(default=None, max_length=255),
    page: int | None = Query(default=None, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """列出商品；支持按商品名称或 SKU 搜索。"""
    stmt = _product_summary_stmt()
    count_stmt = select(func.count()).select_from(Product)
    if keyword and keyword.strip():
        pattern = f"%{keyword.strip()}%"
        product_filter = or_(Product.name.ilike(pattern), Product.sku.ilike(pattern))
        stmt = stmt.where(product_filter)
        count_stmt = count_stmt.where(product_filter)
    if page is None:
        result = await db.execute(stmt)
        return [_product_summary_from_row(row) for row in result.all()]

    total = (await db.execute(count_stmt)).scalar_one()
    page = min(page, max(1, (total + page_size - 1) // page_size))
    result = await db.execute(stmt.offset((page - 1) * page_size).limit(page_size))
    return ProductPageResponse(
        items=[_product_summary_from_row(row) for row in result.all()],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/export")
async def export_products(
    product_ids: list[int] | None = Query(default=None, description="要导出的商品 ID，可重复传入"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """导出一个或多个商品为 Excel，并嵌入商品图片。"""
    stmt = select(Product).order_by(Product.id.desc())
    if product_ids:
        stmt = stmt.where(Product.id.in_(product_ids))
    result = await db.execute(stmt)
    products = list(result.scalars().all())
    if not products:
        raise HTTPException(status_code=404, detail="没有可导出的商品")

    product_ids = [product.id for product in products]
    summaries = await db.execute(
        select(
            InventoryBatch.product_id,
            func.coalesce(func.sum(InventoryBatch.remaining_quantity), 0).label("stock_quantity"),
            func.coalesce(
                func.sum(InventoryBatch.remaining_quantity * InventoryBatch.unit_cost), 0
            ).label("stock_value"),
            func.max(InventoryBatch.arrived_at).label("last_batch_at"),
        )
        .where(InventoryBatch.product_id.in_(product_ids))
        .group_by(InventoryBatch.product_id)
    )
    summary_map = {
        row.product_id: (row.stock_quantity, row.stock_value, row.last_batch_at)
        for row in summaries.all()
    }
    rows = [
        (product, *summary_map.get(product.id, (0, Decimal("0"), None)))
        for product in products
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品数据"
    sheet.append([
        "图片", "商品ID", "SKU", "商品名称", "商品类型", "安全库存",
        "库存数量", "库存价值", "最近入库", "创建时间",
    ])
    header_fill = PatternFill("solid", fgColor="5B9BD5")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"

    for row_index, (product, stock_qty, stock_value, last_batch) in enumerate(rows, start=2):
        sheet.append([
            "", product.id, product.sku, product.name,
            "新品" if product.product_type == PRODUCT_TYPE_NEW else "稳健商品",
            product.safe_stock_quantity, stock_qty, stock_value, last_batch, product.created_at,
        ])
        image_path = None
        if product.image:
            candidate = (BASE_DIR / product.image.lstrip("/")).resolve()
            uploads_root = (BASE_DIR / "uploads").resolve()
            if candidate.is_file() and uploads_root in candidate.parents:
                image_path = candidate
        if image_path:
            try:
                image = ExcelImage(str(image_path))
                max_size = 72
                scale = min(max_size / image.width, max_size / image.height, 1)
                image.width = max(1, int(image.width * scale))
                image.height = max(1, int(image.height * scale))
                sheet.add_image(image, f"A{row_index}")
                sheet.row_dimensions[row_index].height = 60
            except (OSError, ValueError, TypeError):
                pass

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet["H"][1:]:
        cell.number_format = '¥#,##0.00'
    for column in ("I", "J"):
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
    for column, width in {
        "A": 12, "B": 10, "C": 18, "D": 24, "E": 12, "F": 12,
        "G": 12, "H": 16, "I": 20, "J": 20,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"product_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _enrich_product(product, stock_qty, stock_val, last_batch):
    return ProductResponse(
        id=product.id,
        sku=product.sku,
        name=product.name,
        image=product.image,
        product_type=product.product_type,
        safe_stock_quantity=product.safe_stock_quantity,
        created_at=product.created_at,
        stock_quantity=stock_qty,
        stock_value=stock_val,
        last_batch_at=str(last_batch)[:16] if last_batch else None,
    )


def _product_summary_from_row(row):
    product, stock_qty, stock_val, last_batch = row
    return _enrich_product(product, stock_qty, stock_val, last_batch)


@router.put("/{product_id}", response_model=ProductResponse)
async def update_product(
    product_id: int,
    sku: str = Form(...),
    name: str = Form(...),
    product_type: str | None = Form(default=None),
    safe_stock_quantity: int | None = Form(default=None),
    image: UploadFile | None = File(None),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """编辑商品 - 所有角色可操作。"""
    product = await db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")

    # 检查SKU重复（排除自己）
    existing = await db.execute(select(Product).where(Product.sku == sku, Product.id != product_id))
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=400, detail=f"SKU已存在: {sku}")

    product.sku = sku
    product.name = name
    if product_type is not None or safe_stock_quantity is not None:
        product.product_type, product.safe_stock_quantity = _impairment_settings(
            product_type if product_type is not None else product.product_type,
            safe_stock_quantity
            if safe_stock_quantity is not None
            else product.safe_stock_quantity,
        )

    if image and image.filename:
        # 删除旧图片
        if product.image:
            old_path = Path(__file__).parent.parent / product.image.lstrip("/")
            if old_path.exists():
                old_path.unlink()
        ext = os.path.splitext(image.filename)[1]
        filename = f"{uuid.uuid4().hex}{ext}"
        filepath = UPLOAD_DIR / filename
        content = await image.read()
        filepath.write_bytes(content)
        product.image = f"/uploads/{filename}"

    await db.commit()
    await db.refresh(product)
    return _enrich_product(product, 0, Decimal("0"), None)


@router.delete("/{product_id}")
async def delete_product(
    product_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """删除商品 - 所有角色可操作。"""
    product = await db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="商品不存在")

    platform_sku_mapping_count = await db.scalar(
        select(func.count()).select_from(PlatformSkuComponent).where(
            PlatformSkuComponent.product_id == product_id
        )
    )
    if platform_sku_mapping_count:
        raise HTTPException(status_code=400, detail="商品已被平台 SKU 映射引用，请先删除或修改映射")

    if product.image:
        filepath = Path(__file__).parent.parent / product.image.lstrip("/")
        if filepath.exists():
            filepath.unlink()

    await db.delete(product)
    await db.commit()
    return {"ok": True}
