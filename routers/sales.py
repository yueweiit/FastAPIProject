import csv
import hashlib
import io
import re
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from database import get_db
from models import (
    InventoryBatch,
    PlatformSkuComponent,
    PlatformSkuMapping,
    Product,
    Sale,
    SaleCostDetail,
    SalesImportBatch,
    SettlementEntry,
    SettlementEntryAllocation,
    Store,
    StoreProduct,
    User,
)
from schemas import (
    CostDetailResponse,
    MonthlySalesSummaryResponse,
    SaleCreate,
    SalePageResponse,
    SaleResponse,
    SalesImportError,
    SalesImportBatchResponse,
    SalesImportRollbackResponse,
    SalesImportResponse,
    PendingConfirmationPageResponse,
    PendingConfirmationResponse,
)
from services.fifo import fifo_sell, InsufficientStockError
from services.exchange_rates import ExchangeRateError, get_cny_rates
from auth import RequireOperator, RequireAnyRole

router = APIRouter(prefix="/sales", tags=["销售管理"])

ORDER_DETAIL_SHEET = "订单详情"
ORDER_DETAIL_HEADERS = {
    "结算日期",
    "结算单 ID",
    "付款 ID",
    "状态",
    "货币",
    "交易类型",
    "订单ID/调整单ID",
    "SKU ID",
    "数量",
}
ORDER_DETAIL_HEADER_ALIASES = {
    "结算日期": "结算日期",
    "结算单ID": "结算单 ID",
    "付款ID": "付款 ID",
    "状态": "状态",
    "货币": "货币",
    "交易类型": "交易类型",
    "订单ID/调整单ID": "订单ID/调整单ID",
    "SKUID": "SKU ID",
    "数量": "数量",
}
def _clean_header(value) -> str:
    return str(value or "").strip().lstrip("\ufeff")


def _order_detail_header(value) -> str:
    """Normalize the required TikTok headers while preserving other columns."""
    header = _clean_header(value)
    compact = re.sub(r"\s+", "", header)
    return ORDER_DETAIL_HEADER_ALIASES.get(compact, header)


def _text_value(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal_value(value, default: Decimal = Decimal("0")) -> Decimal:
    if value is None:
        return default
    text = str(value).strip().replace(",", "")
    if not text or text in {"/", "-", "—"}:
        return default
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError(f"金额不是有效数字: {value}")


def _quantity_value(value, allow_negative: bool = False) -> int:
    quantity = _decimal_value(value)
    if (not allow_negative and quantity < 0) or quantity != quantity.to_integral_value():
        raise ValueError(f"数量不是整数: {value}")
    return int(quantity)


def _datetime_value(value) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    text = str(value).strip()
    if not text or text in {"/", "-", "—"}:
        return None
    for candidate in (text, text.replace("/", "-")):
        try:
            return datetime.fromisoformat(candidate.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            pass
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                return datetime.strptime(candidate, fmt)
            except ValueError:
                pass
    raise ValueError(f"日期不是有效格式: {value}")


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if value is None:
        return None
    return value


def _source_value(row: dict, *names: str):
    for name in names:
        if name in row:
            return row[name]
    return None


PRODUCT_MULTIPLIER_RE = re.compile(r"^(?P<name>.+?)\s*(?:\*|×|[xX])\s*(?P<count>\d+)$")


def _product_expression_parts(value: str) -> list[tuple[str, int]]:
    """Split bundle names such as A*3+B into component names and counts."""
    parts = re.split(r"\s*[+＋]\s*", _text_value(value))
    if not parts or any(not part.strip() for part in parts):
        raise ValueError(f"商品名称组合格式无效: {value}")

    result: list[tuple[str, int]] = []
    for part in parts:
        component = part.strip()
        multiplier = 1
        match = PRODUCT_MULTIPLIER_RE.match(component)
        if match:
            component = match.group("name").strip()
            multiplier = int(match.group("count"))
            if multiplier <= 0:
                raise ValueError(f"商品数量必须大于 0: {part}")
        if not component:
            raise ValueError(f"商品名称组合格式无效: {value}")
        result.append((component, multiplier))
    return result


def _read_import_file(filename: str, content: bytes) -> tuple[str, str | None, list[dict]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            # Some TikTok exports keep the worksheet dimension as A1 even when
            # the XML contains all rows. Normal mode reads the actual cells.
            workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Excel 文件无法读取: {exc}")
        if ORDER_DETAIL_SHEET not in workbook.sheetnames:
            raise HTTPException(status_code=400, detail=f"Excel 中未找到“{ORDER_DETAIL_SHEET}”工作表")
        sheet = workbook[ORDER_DETAIL_SHEET]
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [_order_detail_header(value) for value in next(values)]
        except StopIteration:
            raise HTTPException(status_code=400, detail="订单详情工作表为空")
        missing = sorted(ORDER_DETAIL_HEADERS - set(headers))
        if missing:
            raise HTTPException(status_code=400, detail=f"订单详情缺少必要列: {', '.join(missing)}")
        rows = []
        for row_number, values_row in enumerate(values, start=2):
            values_list = list(values_row)
            if not any(value not in (None, "") for value in values_list):
                continue
            rows.append({
                "_row_number": row_number,
                "_headers": headers,
                **{header: _json_value(values_list[index] if index < len(values_list) else None)
                   for index, header in enumerate(headers) if header},
            })
        workbook.close()
        return "order_detail", ORDER_DETAIL_SHEET, rows

    if suffix != ".csv":
        raise HTTPException(status_code=400, detail="只支持 CSV、XLSX 或 XLSM 文件")
    decoded = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise HTTPException(status_code=400, detail="CSV 文件编码无法识别")

    reader = csv.reader(io.StringIO(decoded))
    try:
        first_row = next(reader)
    except StopIteration:
        raise HTTPException(status_code=400, detail="CSV 文件为空")
    first_headers = [_order_detail_header(value) for value in first_row]
    is_order_detail = bool(ORDER_DETAIL_HEADERS & set(first_headers))
    rows = []
    if is_order_detail:
        missing = sorted(ORDER_DETAIL_HEADERS - set(first_headers))
        if missing:
            raise HTTPException(status_code=400, detail=f"订单详情 CSV 缺少必要列: {', '.join(missing)}")
        for row_number, values_row in enumerate(reader, start=2):
            if not any(value.strip() for value in values_row):
                continue
            rows.append({
                "_row_number": row_number,
                "_headers": first_headers,
                **{header: values_row[index].strip() if index < len(values_row) else ""
                   for index, header in enumerate(first_headers) if header},
            })
        return "order_detail", None, rows

    # 保留旧版无表头 CSV：SKU, 日期, 数量。
    legacy_rows = [first_row, *reader]
    for row_number, values_row in enumerate(legacy_rows, start=1):
        values_row = [value.strip() for value in values_row]
        if not any(values_row):
            continue
        if len(values_row) < 3:
            raise HTTPException(status_code=400, detail=f"第 {row_number} 行不是三列 CSV 格式")
        rows.append({
            "_row_number": row_number,
            "_headers": ["SKU", "日期", "数量"],
            "SKU": values_row[0],
            "日期": values_row[1],
            "数量": values_row[2],
        })
    return "legacy_csv", None, rows


def _lookup_key(value) -> str:
    return " ".join(_text_value(value).casefold().split())


def _mapping_is_effective(mapping: StoreProduct, as_of: date | None) -> bool:
    if not mapping.is_active:
        return False
    if as_of is None:
        return True
    if mapping.effective_from > as_of:
        return False
    return mapping.effective_to is None or as_of <= mapping.effective_to


def _resolve_product(
    row: dict,
    products: list[Product],
    store_products: list[StoreProduct],
    as_of: date | None = None,
) -> tuple[Product | None, str | None]:
    products_by_id = {product.id: product for product in products}
    by_sku: dict[str, list[Product]] = {}
    by_name: dict[str, list[Product]] = {}
    for product in products:
        by_sku.setdefault(_lookup_key(product.sku), []).append(product)
        by_name.setdefault(_lookup_key(product.name), []).append(product)

    sku_id = _text_value(_source_value(row, "SKU ID", "SKU"))
    platform_product_id = _text_value(
        _source_value(row, "商品 ID", "平台商品 ID", "平台商品ID")
    )
    product_name_values = [
        ("产品名称", _text_value(row.get("产品名"))),
        ("商品名称", _text_value(row.get("商品名称"))),
        ("SKU 名称", _text_value(row.get("SKU 名称"))),
    ]

    mapping_matches = []
    for value, field_name in (
        (sku_id, "platform_sku_id"),
        (platform_product_id, "platform_product_id"),
    ):
        lookup = _lookup_key(value)
        if not lookup:
            continue
        mapping_matches.extend(
            mapping
            for mapping in store_products
            if _lookup_key(getattr(mapping, field_name)) == lookup
        )
    if sku_id:
        lookup = _lookup_key(sku_id)
        mapping_matches.extend(
            mapping
            for mapping in store_products
            if _lookup_key(mapping.store_sku) == lookup
        )
    effective_matches = [
        mapping for mapping in mapping_matches if _mapping_is_effective(mapping, as_of)
    ]
    if mapping_matches and not effective_matches:
        return None, f"平台 SKU 映射在该订单日期无有效版本: {sku_id or platform_product_id}"
    mapped_product_ids = {mapping.product_id for mapping in effective_matches}
    if len(mapped_product_ids) > 1:
        return None, f"平台 SKU 映射到多个本地商品: {sku_id or platform_product_id}"
    if len(mapped_product_ids) == 1:
        return products_by_id[next(iter(mapped_product_ids))], None

    for value, label, lookup_mapping in (
        (sku_id, "SKU ID", by_sku),
        *[(value, label, by_name) for label, value in product_name_values],
    ):
        lookup = _lookup_key(value)
        if not lookup:
            continue
        matches = lookup_mapping.get(lookup, [])
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, f"{label}匹配到多个本地商品: {value}"
    product_name = next((value for _, value in product_name_values if value), "")
    return None, f"未找到本地商品映射: SKU ID={sku_id or '-'}, 产品名={product_name or '-'}"


def _resolve_product_name(name: str, products: list[Product]) -> tuple[Product | None, str | None]:
    lookup = _lookup_key(name)
    matches = [product for product in products if _lookup_key(product.name) == lookup]
    if len(matches) == 1:
        return matches[0], None
    if len(matches) > 1:
        return None, f"产品名称匹配到多个本地商品: {name}"
    return None, f"未找到本地商品映射: 产品名={name}"


def _resolve_platform_sku_components(
    row: dict,
    mappings_by_sku: dict[str, PlatformSkuMapping],
) -> tuple[list[dict] | None, str | None]:
    """Resolve an order-detail row strictly through its configured platform SKU ID."""
    sku_id = _text_value(row.get("SKU ID"))
    mapping = mappings_by_sku.get(sku_id)
    if not mapping:
        return None, f"未配置平台 SKU ID 映射: {sku_id or '-'}"
    if not mapping.components:
        return None, f"平台 SKU ID 未配置商品组件: {sku_id}"

    return [
        {
            "product": component.product,
            "component_name": component.product.name,
            "multiplier": component.quantity_per_sale,
        }
        for component in mapping.components
    ], None


def _resolve_product_components(
    row: dict,
    products: list[Product],
    store_products: list[StoreProduct],
    as_of: date | None,
) -> tuple[list[dict] | None, str | None]:
    """Resolve a platform product expression into local products and multipliers."""
    product_name = _text_value(_source_value(row, "产品名", "商品名称"))
    if not product_name:
        product, error = _resolve_product(row, products, store_products, as_of)
        return ([{"product": product, "component_name": product.name, "multiplier": 1}] if product else None), error

    try:
        parts = _product_expression_parts(product_name)
    except ValueError as exc:
        return None, str(exc)

    components: list[dict] = []
    for component_name, multiplier in parts:
        product, error = _resolve_product_name(component_name, products)
        if not product:
            # A non-bundle row may still be resolved through platform SKU mapping.
            if len(parts) == 1:
                product, error = _resolve_product(row, products, store_products, as_of)
            if not product:
                return None, error
        existing = next((item for item in components if item["product"].id == product.id), None)
        if existing:
            existing["multiplier"] += multiplier
        else:
            components.append({
                "product": product,
                "component_name": component_name,
                "multiplier": multiplier,
            })
    return components, None


def _standalone_price_from_row(row: dict, quantity: int) -> Decimal | None:
    """Return a positive per-unit standalone price from an order-detail row."""
    if quantity <= 0:
        return None
    for field_name in ("商品原价小计", "净商品销售额"):
        try:
            amount = _decimal_value(row.get(field_name))
        except ValueError:
            continue
        if amount > 0:
            return amount / Decimal(quantity)
    return None


def _select_standalone_price(
    observations: list[dict],
    sold_at: datetime | None,
) -> Decimal | None:
    """Pick the most representative observation nearest to the bundle sale date."""
    if not observations:
        return None
    candidates = observations
    if sold_at is not None:
        target_date = sold_at.date()
        dated = [item for item in observations if item["sold_at"] is not None]
        if dated:
            best_date_rank = min(
                (
                    abs((item["sold_at"].date() - target_date).days),
                    0 if item["sold_at"].date() <= target_date else 1,
                )
                for item in dated
            )
            candidates = [
                item
                for item in dated
                if (
                    abs((item["sold_at"].date() - target_date).days),
                    0 if item["sold_at"].date() <= target_date else 1,
                ) == best_date_rank
            ]

    price_counts = Counter(item["price"] for item in candidates)
    selected = min(
        candidates,
        key=lambda item: (-price_counts[item["price"]], item["row_number"]),
    )
    return selected["price"]


def _component_allocation_shares(components: list[dict], item_quantity: int) -> list[Decimal]:
    total_price_weight = sum(
        (component.get("price_allocation_weight", Decimal("0")) for component in components),
        Decimal("0"),
    )
    if total_price_weight > 0:
        return [
            component["price_allocation_weight"] / total_price_weight
            for component in components
        ]

    component_quantities = [item_quantity * component["multiplier"] for component in components]
    total_component_quantity = sum(component_quantities)
    if total_component_quantity <= 0:
        return [Decimal("0") for _ in components]
    return [
        Decimal(component_quantity) / Decimal(total_component_quantity)
        for component_quantity in component_quantities
    ]


def _allocate_available_stock(candidates: list[dict], stock_map: dict[int, int]) -> None:
    """Allocate stock in file order without partially fulfilling a row."""
    remaining_stock = dict(stock_map)
    for item in candidates:
        if not item["is_inventory_sale"]:
            continue

        requirements = []
        for component in item["components"]:
            required = item["quantity"] * component["multiplier"]
            product = component["product"]
            available = remaining_stock.get(product.id, 0)
            requirements.append((component, required, available))

        shortage_messages = [
            f"商品 {component['product'].sku} 本行需要 {required} 件，当前可用 {available} 件"
            for component, required, available in requirements
            if available < required
        ]
        if shortage_messages:
            item["mapping_error"] = "库存不足：" + "；".join(shortage_messages)
            item["is_inventory_sale"] = False
            continue

        for component, required, available in requirements:
            component["sale_quantity"] = required
            remaining_stock[component["product"].id] = available - required


def _order_no(order_id: str, sku_id: str) -> str:
    value = f"{order_id}-{sku_id}" if sku_id else order_id
    if len(value) <= 64:
        return value
    return f"tt:{hashlib.sha256(value.encode('utf-8')).hexdigest()[:60]}"


def _source_key(row: dict, source_kind: str) -> str:
    if source_kind == "legacy_csv":
        value = f"legacy|{row.get('SKU','')}|{row.get('日期','')}|{row.get('数量','')}"
    else:
        value = "|".join(
            _text_value(row.get(name))
            for name in ("结算单 ID", "付款 ID", "订单ID/调整单ID", "SKU ID")
        )
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _other_expense_from_row(net_product_sales: Decimal, settlement_total: Decimal) -> Decimal:
    """Return the settlement difference shown as other expense."""
    return net_product_sales - settlement_total


def _settlement_entry(
    row: dict,
    source_key: str,
    filename: str,
    sheet_name: str | None,
    product_id: int | None,
    sale_id: int | None,
    store_id: int | None,
    import_batch_id: int | None,
    exchange_rate: dict | None,
    mapping_error: str | None = None,
) -> SettlementEntry:
    def money(name: str) -> Decimal:
        return _decimal_value(row.get(name))

    return SettlementEntry(
        source_key=source_key,
        source_file=filename[:255] if filename else None,
        source_sheet=sheet_name,
        source_row_number=int(row["_row_number"]),
        settlement_id=_text_value(row.get("结算单 ID")) or None,
        payout_id=_text_value(row.get("付款 ID")) or None,
        order_id=_text_value(_source_value(row, "订单ID/调整单ID", "订单号", "SKU")),
        related_order_id=_text_value(row.get("相关订单 ID")) or None,
        platform_sku_id=_text_value(_source_value(row, "SKU ID", "SKU")),
        product_id=product_id,
        sale_id=sale_id,
        store_id=store_id,
        import_batch_id=import_batch_id,
        mapping_status="pending_confirmation" if mapping_error else "confirmed",
        mapping_error=mapping_error[:500] if mapping_error else None,
        status=_text_value(row.get("状态")) or None,
        currency=_text_value(row.get("货币")) or None,
        transaction_type=_text_value(row.get("交易类型")) or "订单",
        exchange_rate_to_cny=exchange_rate["rate"] if exchange_rate else None,
        exchange_rate_date=exchange_rate["rate_date"] if exchange_rate else None,
        exchange_rate_source=exchange_rate["source"] if exchange_rate else None,
        settlement_date=_datetime_value(row.get("结算日期")),
        order_created_at=_datetime_value(row.get("订单创建日期", row.get("日期"))),
        delivered_at=_datetime_value(row.get("订单送达日期")),
        quantity=_quantity_value(row.get("数量", 0), allow_negative=True) if _text_value(row.get("数量")) else 0,
        product_name=_text_value(_source_value(row, "产品名", "商品名称")) or None,
        sku_name=_text_value(row.get("SKU 名称")) or None,
        settlement_total=money("结算总金额"),
        net_product_sales=money("净商品销售额"),
        net_shipping=money("净运费"),
        taxes=money("税费"),
        platform_commission=money("平台佣金费"),
        service_fee=money("服务费"),
        sfp_service_fee=money("SFP 服务费"),
        per_item_fee=money("每件成交商品的费用"),
        tax_withholding=money("税款"),
        affiliate_commission=money("联盟佣金"),
        creator_commission=money("支付给达人的佣金"),
        creator_shop_ad_commission=money("支付给达人的店铺广告佣金"),
        affiliate_shop_ad_commission=money("支付给联盟服务商的店铺广告佣金"),
        gmv_max_ad_fee=money("GMV Max 广告费"),
        adjustment_amount=money("调整金额"),
        customer_payment=money("客户付款"),
        customer_refund=money("客户退款"),
        raw_data={str(key): _json_value(value) for key, value in row.items() if not key.startswith("_")},
    )


def _import_error_response(
    source_kind: str,
    sheet_name: str | None,
    total_rows: int,
    settlement_rows: int,
    errors: list[SalesImportError],
    warnings: list[str] | None = None,
):
    from fastapi.responses import JSONResponse

    response = SalesImportResponse(
        success=False,
        file_type=source_kind,
        sheet_name=sheet_name,
        total_rows=total_rows,
        settlement_rows=settlement_rows,
        sales_created=0,
        skipped_duplicates=0,
        failed_rows=len(errors),
        pending_confirmation=0,
        errors=errors,
        warnings=warnings or [],
    )
    return JSONResponse(status_code=400, content=response.model_dump(mode="json"))


@router.post("/import", response_model=SalesImportResponse)
async def import_sales_file(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    """导入旧版三列 CSV 或 TikTok Shop“订单详情”工作表。"""
    import_store = None
    if user.role != "admin":
        if user.store_id is None:
            raise HTTPException(status_code=400, detail="当前账号未绑定店铺，请联系管理员")
        import_store = await db.get(Store, user.store_id)
        if not import_store or not import_store.is_active:
            raise HTTPException(status_code=400, detail="当前账号绑定的店铺不存在或已停用")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="导入文件为空")
    source_kind, sheet_name, rows = _read_import_file(file.filename or "", content)
    if not rows:
        raise HTTPException(status_code=400, detail="文件中没有可导入的数据行")

    products = list((await db.execute(select(Product).order_by(Product.id))).scalars().all())
    store_products: list[StoreProduct] = []
    platform_sku_mappings_by_sku: dict[str, PlatformSkuMapping] = {}
    if source_kind == "order_detail":
        platform_sku_mappings = (
            await db.execute(
                select(PlatformSkuMapping)
                .options(
                    selectinload(PlatformSkuMapping.components).selectinload(
                        PlatformSkuComponent.product
                    )
                )
                .where(
                    PlatformSkuMapping.platform == "tiktok_shop",
                    PlatformSkuMapping.is_active.is_(True),
                )
            )
        ).scalars().all()
        platform_sku_mappings_by_sku = {
            mapping.platform_sku_id: mapping for mapping in platform_sku_mappings
        }
    else:
        store_products = list(
            (
                await db.execute(
                    select(StoreProduct).where(StoreProduct.is_active.is_(True))
                )
            ).scalars().all()
        )
    source_keys = [_source_key(row, source_kind) for row in rows]
    existing_entries = {
        key
        for key in (
            await db.execute(
                select(SettlementEntry.source_key).where(SettlementEntry.source_key.in_(source_keys))
            )
        ).scalars()
    }
    errors: list[SalesImportError] = []
    warnings: list[str] = []
    candidates = []
    skipped_duplicates = 0
    source_key_counts = Counter(source_keys)

    for row, source_key in zip(rows, source_keys):
        row_number = int(row["_row_number"])
        identifier = _text_value(_source_value(row, "订单ID/调整单ID", "SKU")) or None
        if source_key_counts[source_key] > 1:
            errors.append(SalesImportError(row=row_number, message="文件内存在重复的订单详情行", identifier=identifier))
            continue
        if source_key in existing_entries:
            skipped_duplicates += 1
            continue
        try:
            quantity = _quantity_value(row.get("数量"), allow_negative=source_kind == "order_detail")
        except ValueError as exc:
            errors.append(SalesImportError(row=row_number, message=str(exc), identifier=identifier))
            continue

        transaction_type = _text_value(row.get("交易类型")) or "订单"
        is_inventory_sale = source_kind == "legacy_csv" or (
            transaction_type == "订单" and quantity > 0
        )
        if source_kind == "legacy_csv":
            try:
                sold_at = _datetime_value(row.get("日期"))
            except ValueError as exc:
                errors.append(SalesImportError(row=row_number, message=str(exc), identifier=identifier))
                continue
            if sold_at is None:
                errors.append(SalesImportError(row=row_number, message="销售日期不能为空", identifier=identifier))
                continue
            order_created_at = None
            settlement_date = None
            net_product_sales = Decimal("0")
            settlement_total = Decimal("0")
            other_expense = Decimal("0")
        else:
            try:
                order_created_at = _datetime_value(row.get("订单创建日期"))
                settlement_date = _datetime_value(row.get("结算日期"))
                net_product_sales = _decimal_value(row.get("净商品销售额"))
                settlement_total = _decimal_value(row.get("结算总金额"))
            except ValueError as exc:
                errors.append(SalesImportError(row=row_number, message=str(exc), identifier=identifier))
                continue
            if is_inventory_sale and not (order_created_at or settlement_date):
                errors.append(SalesImportError(row=row_number, message="订单创建日期和结算日期不能同时为空", identifier=identifier))
                continue
            sold_at = order_created_at or settlement_date

        lookup_date = sold_at.date() if sold_at else None
        if is_inventory_sale and source_kind == "order_detail":
            if not _text_value(row.get("订单ID/调整单ID")):
                errors.append(SalesImportError(row=row_number, message="订单 ID 不能为空", identifier=identifier))
                continue
            if not _text_value(row.get("SKU ID")):
                errors.append(SalesImportError(row=row_number, message="SKU ID 不能为空", identifier=identifier))
                continue
        if source_kind == "legacy_csv":
            # Keep the original order number format so old imports remain deduplicated.
            order_no = f"{row['日期']}-{row['SKU']}"[:64]
        else:
            if is_inventory_sale and net_product_sales < 0:
                warnings.append(f"第 {row_number} 行商品销售额为负，仅保存结算行，不扣库存")
                is_inventory_sale = False
            order_id = _text_value(row.get("订单ID/调整单ID"))
            sku_id = _text_value(row.get("SKU ID"))
            order_no = _order_no(order_id, sku_id)
            other_expense = _other_expense_from_row(net_product_sales, settlement_total)

        if is_inventory_sale and quantity <= 0:
            errors.append(SalesImportError(row=row_number, message="销售数量必须大于 0", identifier=identifier))
            continue

        if source_kind == "order_detail":
            components, mapping_error = _resolve_platform_sku_components(
                row, platform_sku_mappings_by_sku
            )
        else:
            components, mapping_error = _resolve_product_components(
                row, products, store_products, lookup_date
            )
        pending_mapping_error = None
        if not components:
            pending_mapping_error = mapping_error or "商品映射失败"
            is_inventory_sale = False

        candidates.append({
            "row": row,
            "source_key": source_key,
            "components": components or [],
            "quantity": quantity,
            "sold_at": sold_at,
            "order_no": order_no,
            "net_product_sales": net_product_sales,
            "other_expense": other_expense,
            "is_inventory_sale": is_inventory_sale,
            "mapping_error": pending_mapping_error,
        })

    if errors:
        return _import_error_response(source_kind, sheet_name, len(rows), len(rows), errors, warnings)

    fx_requirements: set[tuple[str, date]] = set()
    for item in candidates:
        currency = (
            _text_value(item["row"].get("货币")).upper()
            if source_kind == "order_detail"
            else "CNY"
        )
        item["currency"] = currency
        item["exchange_rate"] = None
        if item["sold_at"] is not None and currency:
            fx_requirements.add((currency, item["sold_at"].date()))
    try:
        exchange_rates = await get_cny_rates(fx_requirements)
    except ExchangeRateError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    for item in candidates:
        if item["sold_at"] is not None and item["currency"]:
            item["exchange_rate"] = exchange_rates.get(
                (item["currency"], item["sold_at"].date())
            )

    if source_kind == "order_detail":
        standalone_observations: dict[tuple[int, str], list[dict]] = {}
        for item in candidates:
            if not item["is_inventory_sale"] or len(item["components"]) != 1:
                continue
            component = item["components"][0]
            if component["multiplier"] != 1:
                continue
            price = _standalone_price_from_row(item["row"], item["quantity"])
            if price is None:
                continue
            currency = _lookup_key(item["row"].get("货币"))
            key = (component["product"].id, currency)
            standalone_observations.setdefault(key, []).append({
                "price": price,
                "sold_at": item["sold_at"],
                "row_number": int(item["row"]["_row_number"]),
            })

        for item in candidates:
            if not item["is_inventory_sale"] or len(item["components"]) <= 1:
                continue
            currency = _lookup_key(item["row"].get("货币"))
            missing_price_names = []
            for component in item["components"]:
                price = _select_standalone_price(
                    standalone_observations.get((component["product"].id, currency), []),
                    item["sold_at"],
                )
                if price is None:
                    missing_price_names.append(component["component_name"])
                    continue
                component["price_allocation_weight"] = price * component["multiplier"]
            if missing_price_names:
                item["is_inventory_sale"] = False
                item["mapping_error"] = (
                    "组合商品缺少单独售价：" + "、".join(missing_price_names)
                )

    candidate_order_nos = [item["order_no"] for item in candidates if item["is_inventory_sale"]]
    existing_order_nos = set(
        (
            await db.execute(select(Sale.order_no).where(Sale.order_no.in_(candidate_order_nos)))
        ).scalars()
    ) if candidate_order_nos else set()
    duplicate_candidates = [item for item in candidates if item["is_inventory_sale"] and item["order_no"] in existing_order_nos]
    if duplicate_candidates:
        for item in duplicate_candidates:
            errors.append(SalesImportError(
                row=int(item["row"]["_row_number"]),
                message="销售订单已存在，可能是重复导入",
                identifier=item["order_no"],
            ))
        return _import_error_response(source_kind, sheet_name, len(rows), len(rows), errors, warnings)

    required_product_ids: set[int] = set()
    for item in candidates:
        if item["is_inventory_sale"]:
            for component in item["components"]:
                required_product_ids.add(component["product"].id)
    if required_product_ids:
        stock_rows = await db.execute(
            select(InventoryBatch.product_id, func.sum(InventoryBatch.remaining_quantity))
            .where(InventoryBatch.product_id.in_(required_product_ids))
            .group_by(InventoryBatch.product_id)
        )
        stock_map = {product_id: int(quantity or 0) for product_id, quantity in stock_rows.all()}
        _allocate_available_stock(candidates, stock_map)

    pending_confirmation = sum(1 for item in candidates if item["mapping_error"])
    import_batch = SalesImportBatch(
        user_id=user.id,
        store_id=import_store.id if import_store else None,
        file_name=(file.filename or "未命名导入文件")[:255],
        file_type=source_kind,
        sheet_name=sheet_name,
        total_rows=len(rows),
        settlement_rows=len(candidates),
        sales_created=0,
        skipped_duplicates=skipped_duplicates,
        pending_confirmation=pending_confirmation,
        status="completed",
    )
    db.add(import_batch)
    await db.flush()

    sales_created = 0
    try:
        for item in candidates:
            sales_by_product = {}
            if item["is_inventory_sale"]:
                component_shares = _component_allocation_shares(
                    item["components"], item["quantity"]
                )
                for component, component_share in zip(item["components"], component_shares):
                    component_quantity = item["quantity"] * component["multiplier"]
                    component_sales = item["net_product_sales"] * component_share
                    component_other_expense = item["other_expense"] * component_share
                    try:
                        sale, _ = await fifo_sell(
                            db=db,
                            product_id=component["product"].id,
                            order_no=item["order_no"],
                            quantity=component_quantity,
                            selling_price=component_sales / component_quantity,
                            platform_fee=component_other_expense,
                            sold_at=item["sold_at"],
                            user_id=user.id,
                            store_id=import_store.id if import_store else None,
                            commit=False,
                        )
                    except InsufficientStockError as exc:
                        raise HTTPException(status_code=400, detail=str(exc))
                    sales_by_product[component["product"].id] = sale
                    sales_created += 1

            settlement_product_id = None
            settlement_sale_id = None
            if len(item["components"]) == 1:
                settlement_product_id = item["components"][0]["product"].id
                settlement_sale_id = next(iter(sales_by_product.values())).id if sales_by_product else None
            settlement_entry = _settlement_entry(
                item["row"], item["source_key"], file.filename or "", sheet_name,
                settlement_product_id,
                settlement_sale_id,
                import_store.id if import_store else None,
                import_batch.id,
                item["exchange_rate"],
                item["mapping_error"],
            )
            db.add(settlement_entry)
            await db.flush()
            for component in item["components"]:
                db.add(SettlementEntryAllocation(
                    settlement_entry_id=settlement_entry.id,
                    product_id=component["product"].id,
                    sale_id=sales_by_product.get(component["product"].id).id
                    if component["product"].id in sales_by_product else None,
                    component_name=component["component_name"],
                    multiplier=component["multiplier"],
                    quantity=item["quantity"] * component["multiplier"],
                ))
        import_batch.sales_created = sales_created
        await db.commit()
    except HTTPException:
        await db.rollback()
        raise
    except IntegrityError:
        await db.rollback()
        raise HTTPException(status_code=409, detail="导入数据与已有记录冲突，请刷新后重试")

    return SalesImportResponse(
        success=True,
        file_type=source_kind,
        sheet_name=sheet_name,
        total_rows=len(rows),
        settlement_rows=len(candidates),
        sales_created=sales_created,
        skipped_duplicates=skipped_duplicates,
        failed_rows=len(errors),
        pending_confirmation=pending_confirmation,
        import_batch_id=import_batch.id,
        errors=[],
        warnings=warnings,
    )


@router.get("/imports", response_model=list[SalesImportBatchResponse])
async def list_sales_imports(
    limit: int = Query(default=20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    stmt = (
        select(SalesImportBatch, User.username, Store.name)
        .outerjoin(User, User.id == SalesImportBatch.user_id)
        .outerjoin(Store, Store.id == SalesImportBatch.store_id)
        .order_by(SalesImportBatch.imported_at.desc(), SalesImportBatch.id.desc())
        .limit(limit)
    )
    if user.role != "admin":
        stmt = stmt.where(SalesImportBatch.user_id == user.id)
    rows = (await db.execute(stmt)).all()
    return [
        SalesImportBatchResponse(
            id=batch.id,
            file_name=batch.file_name,
            file_type=batch.file_type,
            sheet_name=batch.sheet_name,
            imported_by=username,
            store_name=store_name,
            total_rows=batch.total_rows,
            settlement_rows=batch.settlement_rows,
            sales_created=batch.sales_created,
            skipped_duplicates=batch.skipped_duplicates,
            pending_confirmation=batch.pending_confirmation,
            status=batch.status,
            imported_at=batch.imported_at,
            rolled_back_at=batch.rolled_back_at,
            can_rollback=(
                batch.status == "completed"
                and (user.role == "admin" or batch.user_id == user.id)
            ),
        )
        for batch, username, store_name in rows
    ]


@router.post("/imports/{import_batch_id}/rollback", response_model=SalesImportRollbackResponse)
async def rollback_sales_import(
    import_batch_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    import_batch = await db.get(SalesImportBatch, import_batch_id)
    if not import_batch:
        raise HTTPException(status_code=404, detail="导入记录不存在")
    if user.role != "admin" and import_batch.user_id != user.id:
        raise HTTPException(status_code=403, detail="只能撤销自己导入的文件")
    if import_batch.status != "completed":
        raise HTTPException(status_code=409, detail="该导入已经撤销")

    entry_ids = list(
        (
            await db.execute(
                select(SettlementEntry.id).where(
                    SettlementEntry.import_batch_id == import_batch.id
                )
            )
        ).scalars()
    )
    sale_ids: set[int] = set()
    if entry_ids:
        sale_ids.update(
            sale_id
            for sale_id in (
                await db.execute(
                    select(SettlementEntryAllocation.sale_id).where(
                        SettlementEntryAllocation.settlement_entry_id.in_(entry_ids),
                        SettlementEntryAllocation.sale_id.is_not(None),
                    )
                )
            ).scalars()
            if sale_id is not None
        )
        sale_ids.update(
            sale_id
            for sale_id in (
                await db.execute(
                    select(SettlementEntry.sale_id).where(
                        SettlementEntry.id.in_(entry_ids),
                        SettlementEntry.sale_id.is_not(None),
                    )
                )
            ).scalars()
            if sale_id is not None
        )

    restored_quantity = 0
    if sale_ids:
        restore_rows = (
            await db.execute(
                select(
                    SaleCostDetail.batch_id,
                    func.sum(SaleCostDetail.quantity),
                )
                .where(SaleCostDetail.sale_id.in_(sale_ids))
                .group_by(SaleCostDetail.batch_id)
            )
        ).all()
        inventory_batches = {
            batch.id: batch
            for batch in (
                await db.execute(
                    select(InventoryBatch)
                    .where(InventoryBatch.id.in_([row.batch_id for row in restore_rows]))
                    .with_for_update()
                )
            ).scalars()
        }
        for batch_id, quantity in restore_rows:
            inventory_batch = inventory_batches.get(batch_id)
            quantity = int(quantity or 0)
            if not inventory_batch:
                raise HTTPException(status_code=409, detail=f"库存批次 {batch_id} 不存在，无法撤销")
            if inventory_batch.remaining_quantity + quantity > inventory_batch.quantity:
                raise HTTPException(status_code=409, detail=f"库存批次 {inventory_batch.batch_no} 的恢复数量异常")
        for batch_id, quantity in restore_rows:
            quantity = int(quantity or 0)
            inventory_batches[batch_id].remaining_quantity += quantity
            restored_quantity += quantity

    if entry_ids:
        await db.execute(
            delete(SettlementEntryAllocation).where(
                SettlementEntryAllocation.settlement_entry_id.in_(entry_ids)
            )
        )
        await db.execute(delete(SettlementEntry).where(SettlementEntry.id.in_(entry_ids)))
    if sale_ids:
        await db.execute(delete(SaleCostDetail).where(SaleCostDetail.sale_id.in_(sale_ids)))
        await db.execute(delete(Sale).where(Sale.id.in_(sale_ids)))

    import_batch.status = "rolled_back"
    import_batch.rolled_back_at = datetime.now()
    import_batch.rolled_back_by_user_id = user.id
    await db.commit()
    return SalesImportRollbackResponse(
        success=True,
        import_batch_id=import_batch.id,
        restored_quantity=restored_quantity,
        sales_removed=len(sale_ids),
        settlement_rows_removed=len(entry_ids),
    )


def _build_sale_response(
    sale: Sale,
    cost_details: list[dict],
    currency: str = "CNY",
    store_name: str | None = None,
    exchange_rate_to_cny: Decimal | None = Decimal("1"),
    exchange_rate_date: date | None = None,
    exchange_rate_source: str | None = "CNY",
    source_net_product_sales: Decimal | None = None,
    source_other_expense: Decimal | None = None,
) -> SaleResponse:
    currency = currency.upper()
    if currency == "CNY":
        exchange_rate_to_cny = exchange_rate_to_cny or Decimal("1")
        exchange_rate_date = exchange_rate_date or sale.sold_at.date()
        exchange_rate_source = exchange_rate_source or "CNY"
    selling_price_cny = None
    sales_revenue_cny = None
    platform_fee_cny = None
    profit_cny = None
    if exchange_rate_to_cny is not None:
        source_net_product_sales = (
            sale.selling_price * sale.quantity
            if source_net_product_sales is None
            else source_net_product_sales
        )
        source_other_expense = (
            sale.platform_fee if source_other_expense is None else source_other_expense
        )
        sales_revenue_cny = source_net_product_sales * exchange_rate_to_cny
        selling_price_cny = (
            sales_revenue_cny / sale.quantity if sale.quantity else Decimal("0")
        )
        platform_fee_cny = source_other_expense * exchange_rate_to_cny
        profit_cny = sales_revenue_cny - sale.total_cost - platform_fee_cny
    return SaleResponse(
        id=sale.id,
        product_id=sale.product_id,
        user_id=sale.user_id,
        order_no=sale.order_no,
        quantity=sale.quantity,
        selling_price=sale.selling_price,
        total_cost=sale.total_cost,
        platform_fee=sale.platform_fee,
        profit=sale.profit,
        currency=currency,
        exchange_rate_to_cny=exchange_rate_to_cny,
        exchange_rate_date=exchange_rate_date,
        exchange_rate_source=exchange_rate_source,
        selling_price_cny=selling_price_cny,
        sales_revenue_cny=sales_revenue_cny,
        platform_fee_cny=platform_fee_cny,
        profit_cny=profit_cny,
        store_name=store_name,
        sold_at=sale.sold_at,
        created_at=sale.created_at,
        cost_details=[
            CostDetailResponse(
                batch_id=d["batch_id"],
                batch_no=d["batch_no"],
                quantity=d["quantity"],
                unit_cost=d["unit_cost"],
            )
            for d in cost_details
        ],
    )


async def _sale_import_context(
    db: AsyncSession, sales: list[Sale]
) -> tuple[dict[int, dict], dict[int, str]]:
    """Resolve imported currency, fixed historical rate, and store."""
    if not sales:
        return {}, {}

    sale_by_id = {sale.id: sale for sale in sales}
    persisted_store_ids = {sale.store_id for sale in sales if sale.store_id is not None}
    store_by_sale: dict[int, str] = {}
    if persisted_store_ids:
        persisted_stores = dict(
            (
                await db.execute(
                    select(Store.id, Store.name).where(Store.id.in_(persisted_store_ids))
                )
            ).all()
        )
        store_by_sale = {
            sale.id: persisted_stores[sale.store_id]
            for sale in sales
            if sale.store_id in persisted_stores
        }
    context_rows = (
        await db.execute(
            select(
                SettlementEntryAllocation.settlement_entry_id,
                SettlementEntryAllocation.sale_id,
                SettlementEntryAllocation.product_id,
                SettlementEntryAllocation.quantity,
                SettlementEntry.platform_sku_id,
                SettlementEntry.currency,
                SettlementEntry.exchange_rate_to_cny,
                SettlementEntry.exchange_rate_date,
                SettlementEntry.exchange_rate_source,
                SettlementEntry.settlement_total,
                SettlementEntry.net_product_sales,
            )
            .join(
                SettlementEntry,
                SettlementEntry.id == SettlementEntryAllocation.settlement_entry_id,
            )
            .where(SettlementEntryAllocation.sale_id.in_(sale_by_id))
        )
    ).all()

    # A bundle creates one Sale per component. Reconstruct each component's
    # share from the stored Sale values so old imports use the same source
    # amounts as new imports without changing their persisted cost history.
    rows_by_entry: dict[int, list] = {}
    for row in context_rows:
        if row.sale_id is not None:
            rows_by_entry.setdefault(row.settlement_entry_id, []).append(row)

    source_amounts_by_sale: dict[int, dict[str, Decimal]] = {}
    for entry_rows in rows_by_entry.values():
        first = entry_rows[0]
        sale_rows: dict[int, list] = {}
        for row in entry_rows:
            sale_rows.setdefault(row.sale_id, []).append(row)
        sale_weights = {
            sale_id: sale_by_id[sale_id].selling_price * sale_by_id[sale_id].quantity
            for sale_id in sale_rows
        }
        total_weight = sum(sale_weights.values(), Decimal("0"))
        if total_weight > 0:
            shares = {
                sale_id: weight / total_weight
                for sale_id, weight in sale_weights.items()
            }
        else:
            allocation_quantities = {
                sale_id: sum(max(row.quantity, 0) for row in rows)
                for sale_id, rows in sale_rows.items()
            }
            total_quantity = sum(allocation_quantities.values())
            if total_quantity > 0:
                shares = {
                    sale_id: Decimal(quantity) / Decimal(total_quantity)
                    for sale_id, quantity in allocation_quantities.items()
                }
            else:
                equal_share = Decimal("1") / Decimal(len(sale_rows))
                shares = {sale_id: equal_share for sale_id in sale_rows}

        net_product_sales = first.net_product_sales or Decimal("0")
        settlement_total = first.settlement_total or Decimal("0")
        other_expense = _other_expense_from_row(net_product_sales, settlement_total)
        for sale_id, share in shares.items():
            amount = net_product_sales * share
            expense = other_expense * share
            context = source_amounts_by_sale.setdefault(
                sale_id,
                {
                    "source_net_product_sales": Decimal("0"),
                    "source_other_expense": Decimal("0"),
                },
            )
            context["source_net_product_sales"] += amount
            context["source_other_expense"] += expense

    import_context_by_sale = {}
    for row in context_rows:
        if row.sale_id is None or row.sale_id in import_context_by_sale:
            continue
        import_context_by_sale[row.sale_id] = {
            "currency": (row.currency or "CNY").upper(),
            "exchange_rate_to_cny": row.exchange_rate_to_cny,
            "exchange_rate_date": row.exchange_rate_date,
            "exchange_rate_source": row.exchange_rate_source,
            **source_amounts_by_sale.get(row.sale_id, {}),
        }

    platform_sku_ids = {row.platform_sku_id for row in context_rows if row.platform_sku_id}
    product_ids = {row.product_id for row in context_rows if row.product_id is not None}
    if not platform_sku_ids or not product_ids:
        return import_context_by_sale, store_by_sale

    mappings = (
        await db.execute(
            select(StoreProduct, Store)
            .join(Store, Store.id == StoreProduct.store_id)
            .where(
                StoreProduct.product_id.in_(product_ids),
                StoreProduct.is_active.is_(True),
            )
        )
    ).all()
    inferred_stores_by_sale: dict[int, set[str]] = {}
    for row in context_rows:
        if row.sale_id is None or not row.platform_sku_id:
            continue
        sale_date = sale_by_id[row.sale_id].sold_at.date()
        lookup = _lookup_key(row.platform_sku_id)
        for mapping, store in mappings:
            if mapping.product_id != row.product_id or not _mapping_is_effective(mapping, sale_date):
                continue
            if lookup not in {
                _lookup_key(mapping.platform_sku_id),
                _lookup_key(mapping.store_sku),
            }:
                continue
            inferred_stores_by_sale.setdefault(row.sale_id, set()).add(store.name)
    for sale_id, values in inferred_stores_by_sale.items():
        if sale_id not in store_by_sale and len(values) == 1:
            store_by_sale[sale_id] = next(iter(values))
    return import_context_by_sale, store_by_sale


@router.post("", response_model=SaleResponse)
async def create_sale(
    data: SaleCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireOperator),
):
    """记录销售 - admin和operator可操作"""
    try:
        sale, cost_details = await fifo_sell(
            db=db,
            product_id=data.product_id,
            order_no=data.order_no,
            quantity=data.quantity,
            selling_price=data.selling_price,
            platform_fee=data.platform_fee,
            sold_at=data.sold_at,
            user_id=user.id,
        )
    except InsufficientStockError as e:
        raise HTTPException(status_code=400, detail=str(e))

    return _build_sale_response(sale, cost_details)


@router.get("/pending-confirmations", response_model=PendingConfirmationPageResponse)
async def list_pending_confirmations(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """列出订单详情导入中尚未匹配本地商品的结算行。"""
    base_filter = SettlementEntry.mapping_status == "pending_confirmation"
    total = (await db.execute(
        select(func.count()).select_from(SettlementEntry).where(base_filter)
    )).scalar_one()
    page = min(page, max(1, (total + page_size - 1) // page_size))
    result = await db.execute(
        select(SettlementEntry).options(selectinload(SettlementEntry.store))
        .where(base_filter)
        .order_by(SettlementEntry.imported_at.desc(), SettlementEntry.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    items = [
        PendingConfirmationResponse(
            id=entry.id,
            source_file=entry.source_file,
            source_sheet=entry.source_sheet,
            source_row_number=entry.source_row_number,
            order_id=entry.order_id,
            platform_sku_id=entry.platform_sku_id,
            mapping_status=entry.mapping_status,
            mapping_error=entry.mapping_error,
            store_name=entry.store.name if entry.store else None,
            settlement_date=entry.settlement_date,
            order_created_at=entry.order_created_at,
            quantity=entry.quantity,
            product_name=entry.product_name,
            sku_name=entry.sku_name,
            settlement_total=entry.settlement_total,
            imported_at=entry.imported_at,
        )
        for entry in result.scalars().all()
    ]
    return PendingConfirmationPageResponse(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
    )


def _group_pending_confirmations(entries):
    grouped = {}
    for entry in entries:
        sku_id = (entry.platform_sku_id or "").strip()
        reason = (entry.mapping_error or "待确认").strip()
        key = (sku_id, reason)
        item = grouped.setdefault(key, {
            "platform_sku_id": sku_id,
            "mapping_error": reason,
            "quantity": 0,
            "record_count": 0,
            "product_names": set(),
            "sku_names": set(),
            "store_names": set(),
            "source_files": set(),
        })
        item["quantity"] += entry.quantity or 0
        item["record_count"] += 1
        if entry.product_name:
            item["product_names"].add(entry.product_name.strip())
        if entry.sku_name:
            item["sku_names"].add(entry.sku_name.strip())
        if entry.store and entry.store.name:
            item["store_names"].add(entry.store.name.strip())
        if entry.source_file:
            item["source_files"].add(entry.source_file.strip())

    rows = []
    for item in grouped.values():
        rows.append({
            **{key: value for key, value in item.items() if not isinstance(value, set)},
            "product_names": "、".join(sorted(item["product_names"])),
            "sku_names": "、".join(sorted(item["sku_names"])),
            "store_names": "、".join(sorted(item["store_names"])),
            "source_files": "、".join(sorted(item["source_files"])),
        })
    return sorted(rows, key=lambda row: (row["platform_sku_id"], row["mapping_error"]))


@router.get("/pending-confirmations/export")
async def export_pending_confirmations(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """导出全部待确认记录，并按 SKU ID 和待确认原因合并。"""
    result = await db.execute(
        select(SettlementEntry).options(selectinload(SettlementEntry.store))
        .where(SettlementEntry.mapping_status == "pending_confirmation")
    )
    rows = _group_pending_confirmations(result.scalars().all())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待确认汇总"
    headers = ["SKU ID", "待确认原因", "合计数量", "记录数", "商品名称", "SKU 名称", "涉及店铺", "来源文件"]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            row["platform_sku_id"], row["mapping_error"], row["quantity"], row["record_count"],
            row["product_names"], row["sku_names"], row["store_names"], row["source_files"],
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = (20, 42, 12, 10, 32, 32, 24, 40)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"pending_confirmations_{date.today():%Y%m%d}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _sale_filter_values(
    product_id: int | None,
    keyword: str | None,
    month: str | None,
    user: User,
):
    conditions = []
    normalized_keyword = keyword.strip() if keyword and keyword.strip() else None
    if product_id:
        conditions.append(Sale.product_id == product_id)
    if normalized_keyword:
        conditions.append(
            (Product.name.ilike(f"%{normalized_keyword}%"))
            | (Product.sku.ilike(f"%{normalized_keyword}%"))
        )
    if month:
        try:
            month_start = datetime.strptime(month, "%Y-%m")
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="月份格式必须为 YYYY-MM") from exc
        next_month = datetime(
            month_start.year + (1 if month_start.month == 12 else 0),
            1 if month_start.month == 12 else month_start.month + 1,
            1,
        )
        conditions.extend((Sale.sold_at >= month_start, Sale.sold_at < next_month))
    if user.role == "operator":
        conditions.append(Sale.user_id == user.id)
    return conditions, normalized_keyword


async def _filtered_sales(
    db: AsyncSession,
    product_id: int | None,
    keyword: str | None,
    month: str | None,
    user: User,
) -> list[Sale]:
    conditions, normalized_keyword = _sale_filter_values(product_id, keyword, month, user)
    stmt = select(Sale)
    if normalized_keyword:
        stmt = stmt.join(Product)
    stmt = stmt.where(*conditions).order_by(Sale.sold_at.desc(), Sale.id.desc())
    return list((await db.execute(stmt)).scalars().all())


def _monthly_sales_summary(
    sales: list[Sale],
    import_context_by_sale: dict[int, dict],
) -> list[MonthlySalesSummaryResponse]:
    grouped: dict[str, dict] = {}
    for sale in sales:
        month = sale.sold_at.strftime("%Y-%m")
        group = grouped.setdefault(
            month,
            {
                "sales_count": 0,
                "sold_quantity": 0,
                "sales_revenue_cny": Decimal("0"),
                "sales_cost_cny": Decimal("0"),
                "platform_fee_cny": Decimal("0"),
                "gross_profit_cny": Decimal("0"),
                "fx_missing_count": 0,
            },
        )
        group["sales_count"] += 1
        group["sold_quantity"] += sale.quantity
        group["sales_cost_cny"] += sale.total_cost

        context = import_context_by_sale.get(sale.id, {})
        currency = (context.get("currency") or "CNY").upper()
        rate = context.get("exchange_rate_to_cny", Decimal("1"))
        if currency == "CNY":
            rate = rate or Decimal("1")
        if rate is None:
            group["fx_missing_count"] += 1
            continue
        source_net_product_sales = context.get(
            "source_net_product_sales", sale.selling_price * sale.quantity
        )
        source_other_expense = context.get("source_other_expense", sale.platform_fee)
        revenue_cny = source_net_product_sales * rate
        platform_fee_cny = source_other_expense * rate
        group["sales_revenue_cny"] += revenue_cny
        group["platform_fee_cny"] += platform_fee_cny
        group["gross_profit_cny"] += revenue_cny - sale.total_cost - platform_fee_cny

    summary = []
    for month in sorted(grouped, reverse=True):
        values = grouped[month]
        if values["fx_missing_count"]:
            values["sales_revenue_cny"] = None
            values["platform_fee_cny"] = None
            values["gross_profit_cny"] = None
        summary.append(MonthlySalesSummaryResponse(month=month, **values))
    return summary


@router.get("/monthly-summary", response_model=list[MonthlySalesSummaryResponse])
async def monthly_sales_summary(
    product_id: int | None = None,
    keyword: str | None = Query(default=None, max_length=255),
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """按月份汇总销售记录，金额统一返回人民币。"""
    sales = await _filtered_sales(db, product_id, keyword, month, user)
    import_context_by_sale, _ = await _sale_import_context(db, sales)
    return _monthly_sales_summary(sales, import_context_by_sale)


@router.get("", response_model=list[SaleResponse] | SalePageResponse)
async def list_sales(
    product_id: int | None = None,
    keyword: str | None = Query(default=None, max_length=255),
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    page: int | None = Query(default=None, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    """列出销售，支持按月份、商品名称或 SKU 筛选。"""
    conditions, normalized_keyword = _sale_filter_values(product_id, keyword, month, user)
    stmt = select(Sale).options(selectinload(Sale.cost_details))
    count_stmt = select(func.count()).select_from(Sale)
    if normalized_keyword:
        stmt = stmt.join(Product)
        count_stmt = count_stmt.join(Product)
    stmt = stmt.where(*conditions).order_by(Sale.sold_at.desc(), Sale.id.desc())
    count_stmt = count_stmt.where(*conditions)
    if page is not None:
        total = (await db.execute(count_stmt)).scalar_one()
        page = min(page, max(1, (total + page_size - 1) // page_size))
        stmt = stmt.offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(stmt)
    sales = result.scalars().all()
    import_context_by_sale, store_by_sale = await _sale_import_context(db, sales)

    batch_ids = {detail.batch_id for sale in sales for detail in sale.cost_details}
    batch_map = {}
    if batch_ids:
        batches_result = await db.execute(
            select(InventoryBatch.id, InventoryBatch.batch_no).where(InventoryBatch.id.in_(batch_ids))
        )
        batch_map = dict(batches_result.all())

    response = []
    for sale in sales:
        details = []
        for cd in sale.cost_details:
            details.append({
                "batch_id": cd.batch_id,
                "batch_no": batch_map.get(cd.batch_id, "未知"),
                "quantity": cd.quantity,
                "unit_cost": cd.unit_cost,
            })
        import_context = import_context_by_sale.get(sale.id, {})
        response.append(_build_sale_response(
            sale,
            details,
            currency=import_context.get("currency", "CNY"),
            store_name=store_by_sale.get(sale.id),
            exchange_rate_to_cny=import_context.get("exchange_rate_to_cny", Decimal("1")),
            exchange_rate_date=import_context.get("exchange_rate_date"),
            exchange_rate_source=import_context.get("exchange_rate_source", "CNY"),
            source_net_product_sales=import_context.get("source_net_product_sales"),
            source_other_expense=import_context.get("source_other_expense"),
        ))
    if page is None:
        return response
    return SalePageResponse(items=response, total=total, page=page, page_size=page_size)


@router.get("/{sale_id}", response_model=SaleResponse)
async def get_sale(
    sale_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(RequireAnyRole),
):
    stmt = select(Sale).where(Sale.id == sale_id).options(selectinload(Sale.cost_details))
    result = await db.execute(stmt)
    sale = result.scalar_one_or_none()
    if not sale:
        raise HTTPException(status_code=404, detail="销售记录不存在")
    if user.role == "operator" and sale.user_id != user.id:
        raise HTTPException(status_code=403, detail="无权查看此记录")

    details = []
    for cd in sale.cost_details:
        batch = await db.get(InventoryBatch, cd.batch_id)
        details.append({
            "batch_id": cd.batch_id,
            "batch_no": batch.batch_no if batch else "未知",
            "quantity": cd.quantity,
            "unit_cost": cd.unit_cost,
        })
    import_context_by_sale, store_by_sale = await _sale_import_context(db, [sale])
    import_context = import_context_by_sale.get(sale.id, {})
    return _build_sale_response(
        sale,
        details,
        currency=import_context.get("currency", "CNY"),
        store_name=store_by_sale.get(sale.id),
        exchange_rate_to_cny=import_context.get("exchange_rate_to_cny", Decimal("1")),
        exchange_rate_date=import_context.get("exchange_rate_date"),
        exchange_rate_source=import_context.get("exchange_rate_source", "CNY"),
        source_net_product_sales=import_context.get("source_net_product_sales"),
        source_other_expense=import_context.get("source_other_expense"),
    )
